# START_MODULE_CONTRACT
# PURPOSE: Парсер формы 0.4 «О проекте в цифрах» (xlsx, 3 листа). Извлекает шапку соглашения, инфо потока, целевые показатели, справочник единиц, подписи в forma_0_4_main.json.
# INPUTS: путь к xlsx, output_dir для записи forma_0_4_main.json.
# OUTPUTS: dict-структура + файл forma_0_4_main.json в output_dir.
# KEYWORDS: forma_0_4, openpyxl, xlsx-parser, three-sheets, indicators, allowed-units.
# LINKS: src/doc_type_validators/forma_0_4.py (потребитель данных), doc_configs/forma_0_4/.
# RATIONALE: 3 листа — основная форма, справочник формул, выпадающие списки допустимых единиц. Логика перенесена из прод-движка audit_engine/forma_0_4 as-is (источник правды). Один проход data_only=True (формулы не нужны — 0.4 проверяет значения, не наличие формул).
# END_MODULE_CONTRACT

from __future__ import annotations

# START_IMPORTS
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional

import openpyxl
# END_IMPORTS


# START_HELPERS
def _v(ws: Any, coord: str) -> Any:
    """Достаёт значение из ячейки. Строку strip-ает, пустую → None."""
    v = ws[coord].value
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _to_iso_date(v: Any) -> Optional[str]:
    """Возвращает ISO дату YYYY-MM-DD или None."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(s, fmt).date().isoformat()
            except ValueError:
                pass
    return None


def _to_float(v: Any) -> Optional[float]:
    """Безопасное преобразование к float (None если невозможно)."""
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None
# END_HELPERS


# START_PARSE_FORMA_0_4
def parse_forma_0_4(xlsx_path: Path, output_dir: Path) -> Dict[str, Any]:
    """Парсит xlsx формы 0.4 (3 листа), кладёт forma_0_4_main.json, возвращает структуру."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws1 = wb[wb.sheetnames[0]]  # О проекте в цифрах

    # --- Шапка соглашения (G2:J4) ---
    appendix_label = _v(ws1, "G2")
    agreement_no = _v(ws1, "G3")
    agreement_date_text = _v(ws1, "G4")

    # --- Общая информация о потоке (строка 8) ---
    flow_info = {
        "company": _v(ws1, "B8"),
        "region": _v(ws1, "C8"),
        "flow_name": _v(ws1, "D8"),
        "share_in_revenue": _to_float(_v(ws1, "E8")),
        "directions": _v(ws1, "F8"),  # F8:G8 merged → берём F8
        "project_start_date_raw": ws1["H8"].value,
        "project_start_date_iso": _to_iso_date(ws1["H8"].value),
    }

    # --- Целевые показатели (строки 10..13) ---
    period_start = _to_iso_date(ws1["F10"].value)
    period_end = _to_iso_date(ws1["G10"].value)

    indicators = []
    for row in (11, 12, 13):
        indicators.append({
            "row": row,
            "name": _v(ws1, f"B{row}"),     # B/C/D merged
            "unit": _v(ws1, f"E{row}"),
            "value_start": _to_float(_v(ws1, f"F{row}")),
            "value_end": _to_float(_v(ws1, f"G{row}")),
        })

    # --- Лист 2: формулы (B2:B7) ---
    formulas_sheet = None
    for sn in wb.sheetnames:
        if "Формул" in sn or "формул" in sn:
            formulas_sheet = wb[sn]
            break
    formulas_filled = []
    if formulas_sheet:
        for r in range(2, 8):
            v = _v(formulas_sheet, f"B{r}")
            formulas_filled.append({"row": r, "filled": bool(v)})

    # --- Лист 3: допустимые единицы (A=время, B=выработка, C=запасы) ---
    units_sheet = None
    for sn in wb.sheetnames:
        if "выпадающ" in sn.lower() or "саисок" in sn.lower() or "список" in sn.lower():
            units_sheet = wb[sn]
            break
    units_dict = {"time": [], "production": [], "stock": []}
    if units_sheet:
        for r in range(2, units_sheet.max_row + 1):
            a = _v(units_sheet, f"A{r}")
            b = _v(units_sheet, f"B{r}")
            c = _v(units_sheet, f"C{r}")
            if a:
                units_dict["time"].append(a)
            if b:
                units_dict["production"].append(b)
            if c:
                units_dict["stock"].append(c)

    # --- Подписи (строки 39..46) ---
    signatures = {
        "consent": _v(ws1, "D39"),
        "signature_date_raw": ws1["G41"].value,
        "signature_date_iso": _to_iso_date(ws1["G41"].value),
        "signer_fio": _v(ws1, "B44"),
        "stamp_marker": _v(ws1, "C45"),
        "doc_date_text": _v(ws1, "B46"),
    }
    if signatures["signature_date_raw"] is not None and not isinstance(signatures["signature_date_raw"], str):
        signatures["signature_date_raw"] = str(signatures["signature_date_raw"])

    # --- Сборка результата ---
    result = {
        "meta": {
            "workbook": xlsx_path.name,
            "sheet": ws1.title,
            "n_sheets": len(wb.sheetnames),
        },
        "header": {
            "appendix_label": appendix_label,
            "agreement_no": agreement_no,
            "agreement_date_text": agreement_date_text,
        },
        "flow_info": flow_info,
        "period_start_iso": period_start,
        "period_end_iso": period_end,
        "period_start_raw": str(ws1["F10"].value) if ws1["F10"].value is not None else None,
        "period_end_raw": str(ws1["G10"].value) if ws1["G10"].value is not None else None,
        "indicators": indicators,
        "formulas_filled": formulas_filled,
        "allowed_units": units_dict,
        "signatures": signatures,
    }

    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / "forma_0_4_main.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2, default=str)

    return result
# END_PARSE_FORMA_0_4
