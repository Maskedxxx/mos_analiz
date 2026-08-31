# START_MODULE_CONTRACT
# PURPOSE: Парсер «Лист присутствия» (тренинги модуль 1 / модуль 2). Извлекает шапку (тренинг, тренер, адрес, дата) и таблицу участников из xlsx в lp_main.json.
# INPUTS: путь к xlsx, output_dir для записи lp_main.json.
# OUTPUTS: dict-структура + файл lp_main.json в output_dir.
# KEYWORDS: list_prisutstviya, openpyxl, xlsx-parser, participants, sheet-picker.
# LINKS: src/doc_type_validators/list_prisutstviya.py (потребитель), doc_configs/list_prisutstviya_modul_{1,2}/.
# RATIONALE: Один общий парсер на оба модуля (различие модулей — только в config-правилах, не в структуре листа). Логика перенесена из прод-движка audit_engine/list_prisutstviya as-is (источник правды). Выбор листа: предпочитаем «Пример ЛП»/«Лист присутствия», пропускаем «Шаблон ЛП».
# END_MODULE_CONTRACT

from __future__ import annotations

# START_IMPORTS
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional

import openpyxl
# END_IMPORTS


# START_SHEET_CONSTANTS
PREFERRED_SHEET_PRIORITY = ["Пример ЛП", "Лист присутствия"]
SKIP_SHEET_NAMES = {"Шаблон ЛП"}  # пропускаем шаблоны
# END_SHEET_CONSTANTS


# START_HELPERS
def _cell(ws: Any, coord: str) -> Optional[Any]:
    v = ws[coord].value
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _normalize_date(v: Any) -> Optional[str]:
    """Возвращает дату в формате YYYY-MM-DD или None если непонятно."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(s, fmt).date().isoformat()
            except ValueError:
                continue
    return None


# Колонки таблицы участников (№, ФИО, должность, организация, ИНН, регион, e-mail, телефон, подпись)
_TABLE_COLUMNS = "ABCDEFGHI"


def _pick_sheet(wb: openpyxl.Workbook):
    """Выбирает лист с реальными данными: предпочитаем «Пример ЛП»/«Лист присутствия»,
    игнорируем «Шаблон ЛП». Иначе первый непустой (max_row > 5), иначе первый."""
    for name in PREFERRED_SHEET_PRIORITY:
        if name in wb.sheetnames:
            return wb[name]
    for name in wb.sheetnames:
        if name in SKIP_SHEET_NAMES:
            continue
        ws = wb[name]
        if ws.max_row > 5:
            return ws
    return wb[wb.sheetnames[0]]
# END_HELPERS


# START_PARSE_LIST_PRISUTSTVIYA
def parse_list_prisutstviya(xlsx_path: Path, output_dir: Path) -> Dict[str, Any]:
    """Парсит xlsx листа присутствия, кладёт lp_main.json в output_dir, возвращает структуру."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = _pick_sheet(wb)

    # Шапка
    training_name = _cell(ws, "C1")
    trainer_fio = _cell(ws, "C2")
    address = _cell(ws, "C3")
    date_raw = ws["I5"].value
    date_iso = _normalize_date(date_raw)

    # Участники: таблица начинается со строки 6 и идёт подряд — сколько бы строк в бланк ни добавили
    # (стандартная форма на 10 человек, но её регулярно расширяют: встречались листы на 17).
    # Конец таблицы — первая ПОЛНОСТЬЮ пустая строка: ниже неё в части бланков лежит служебный блок
    # с подсказками («ПРОВЕРЬ: …»), который нельзя принимать за участников.
    participants = []
    for row_idx in range(6, ws.max_row + 1):
        row_values = [_cell(ws, f"{col}{row_idx}") for col in _TABLE_COLUMNS]
        if not any(v is not None for v in row_values):
            break
        b = _cell(ws, f"B{row_idx}")
        c = _cell(ws, f"C{row_idx}")
        d = _cell(ws, f"D{row_idx}")
        e = _cell(ws, f"E{row_idx}")
        f = _cell(ws, f"F{row_idx}")
        g = _cell(ws, f"G{row_idx}")
        h = _cell(ws, f"H{row_idx}")
        # Строка участника считается заполненной если есть ФИО ИЛИ организация ИЛИ email.
        if b or d or g:
            participants.append({
                "row": row_idx,
                "fio": b,
                "position": c,
                "org": d,
                "inn": str(e).strip() if e is not None else None,
                "region": f,
                "email": g,
                "phone": str(h).strip() if h is not None else None,
            })

    # Изображения (для проверки логотипа)
    images_count = len(getattr(ws, "_images", []) or [])

    result: Dict[str, Any] = {
        "meta": {
            "workbook": xlsx_path.name,
            "sheet": ws.title,
        },
        "header": {
            "training_name": training_name,
            "trainer_fio": trainer_fio,
            "address": address,
            "date_cell": str(date_raw) if date_raw is not None else None,
            "date_iso": date_iso,
        },
        "participants": participants,
        "images_count": images_count,
    }

    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / "lp_main.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2, default=str)

    return result
# END_PARSE_LIST_PRISUTSTVIYA
