# START_MODULE_CONTRACT
# PURPOSE: Парсеры doc_type «Карточка проекта» (xlsx). Три независимые функции, каждая парсит свой лист книги: основную карточку, методику расчёта показателей и справочник единиц измерения из выпадающего списка.
# INPUTS: Путь к .xlsx файлу карточки, путь к output_dir для сохранения промежуточных JSON.
# OUTPUTS: Три структурированных словаря (каждая функция — свой).
# KEYWORDS: kartochka-proekta, xlsx, structured-parsing, openpyxl.
# LINKS: main.py::KARTOCHKA_PARSER_MODULE_DISPATCH, main.py::kartochka_proekta_run_validations, doc_configs/kartochka_proekta/.
# RATIONALE: Три логически независимых парсера одного и того же файла. Поля читаются по подписям на листе («Клиенты процесса:», «Утверждаю», «Наименование показателя»), а не по фиксированным адресам: раскладка формы у заказчика сдвинута (B2→C6, E11→F15, M11→N15, на листе методики B/F→C/G), и жёсткие адреса давали пустой результат по всем полям сразу. Прежние адреса остались запасным вариантом для файлов старого образца. Каждый отвечает за свой лист и сохраняет JSON на диск для последующих валидаторов (в валидатор попадает именно распаршенный JSON).
# END_MODULE_CONTRACT

from __future__ import annotations

# START_IMPORTS
import json
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, TypedDict

import openpyxl
# END_IMPORTS


# START_CONTRACTS
# PURPOSE: TypedDict-контракты возвращаемых значений трёх парсеров. Показывают структуру верхнего уровня; вложенные объекты оставлены как `Dict[str, Any]` — подробности видны в docstring каждого парсера.
# INPUTS: —
# OUTPUTS: Контракты для IDE-автодополнения, mypy и документации.
# KEYWORDS: typeddict, contract, kartochka-proekta.
class DropdownDocument(TypedDict):
    """
    Результат `parse_dropdown` — справочник допустимых единиц измерения.

    Поля:
        categories: Имя категории (из заголовка колонки A1/B1/C1/D1) → список единиц из
            соответствующей колонки начиная со строки 2.
    """

    categories: Dict[str, List[str]]


class KartochkaMainDocument(TypedDict):
    """
    Результат `parse_kartochka_main` — основной лист «Карточка проекта».

    Поля:
        meta: `{workbook, sheet, max_row}` — имя книги, имя листа, кол-во строк.
        header: `{org_name, project_name, signee_position, signee_name, signee_date,
            has_utverzhday}` — шапка документа.
        section1: `{clients, perimeter, owner, boundaries, leader, team}` — секция 1.
        section2: `{key_risk, justification}` — секция 2.
        indicator_dates: `{base_date, target_date, ideal_date}` — даты замеров показателей.
        indicators: Список показателей, каждый — `{number, name, unit, base_value,
            target_value, ideal_value}`.
        events: Список событий, каждое — `{name, start_date, end_date}`.
    """

    meta: Dict[str, Any]
    header: Dict[str, Any]
    section1: Dict[str, Any]
    section2: Dict[str, Any]
    indicator_dates: Dict[str, Any]
    indicators: List[Dict[str, Any]]
    events: List[Dict[str, Any]]


class MetodikaDocument(TypedDict):
    """
    Результат `parse_metodika` — лист «Методика расчёта».

    Поля:
        meta: `{sheet}` — имя найденного листа; `None` если лист не найден.
        header: `{org_name, project_name}` — может быть пустым dict, если лист не найден.
        indicators: Список показателей, каждый — `{name, unit, calc_method, data_source,
            row_start}`.
    """

    meta: Dict[str, Any]
    header: Dict[str, Any]
    indicators: List[Dict[str, Any]]
# END_CONTRACTS


# START_SHARED_HELPERS
# PURPOSE: Общие утилиты для работы с ячейками и поиска листов, используемые всеми тремя парсерами.
# INPUTS: Объекты openpyxl (Workbook, Worksheet) и координаты.
# OUTPUTS: Stripped-строки, значения ячеек, листы или None.
# KEYWORDS: cell, sheet, helper.
def _find_sheet_by_keyword(wb: Any, keyword: str, exclude: Optional[str] = None) -> Any:
    """
    Назначение:
        Ищет лист в книге по подстроке в имени, опционально исключая шаблонные листы.

    Вход:
        wb: Workbook openpyxl.
        keyword: Подстрока для сравнения (регистр не важен).
        exclude: Если задан — листы, в имени которых есть эта подстрока (регистр не важен),
            пропускаются. Обычно используется для исключения `ШАБЛОН`.

    Выход:
        Worksheet openpyxl или `None`, если лист не найден.

    Логика:
        Линейный проход по `wb.sheetnames`. Первый совпавший — возвращается.
    """
    for name in wb.sheetnames:
        if keyword.lower() in name.lower():
            if exclude and exclude.upper() in name.upper():
                continue
            return wb[name]
    return None


def _cell_str(ws: Any, coord: str) -> str:
    """
    Назначение:
        Возвращает текст ячейки по excel-координате (например, `B2`).

    Вход:
        ws: Лист openpyxl.
        coord: Координата ячейки.

    Выход:
        Stripped-строка или пустая строка, если ячейка пустая.
    """
    val = ws[coord].value
    if val is None:
        return ""
    return str(val).strip()


def _cell_value(ws: Any, coord: str) -> Any:
    """
    Назначение:
        Возвращает значение ячейки «как есть» — без конвертации в строку.

    Вход:
        ws: Лист openpyxl.
        coord: Координата ячейки.

    Выход:
        Значение ячейки (число, дата, строка или `None`).

    Логика:
        Нужна в тех местах, где важно сохранить тип (даты, числа для пост-проверок).
    """
    return ws[coord].value


def _format_date(val: Any) -> Optional[str]:
    """
    Назначение:
        Преобразует дату/строку в ISO-строку `YYYY-MM-DD`.

    Вход:
        val: `datetime`, строка или `None`.

    Выход:
        ISO-строка, исходная строка (stripped) или `None`.

    Логика:
        1. `None` → `None`.
        2. `datetime` → strftime.
        3. Иначе — stripped str; пустая строка → `None`.
    """
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    s = str(val).strip()
    if not s:
        return None
    return s


def _extract_fio(raw: str) -> str:
    """
    Назначение:
        Вытаскивает ФИО из строки вида `____________ И.И. Иванов`.

    Вход:
        raw: Исходная строка.

    Выход:
        Та же строка без групп подчёркиваний, stripped.

    Логика:
        Удаляет подряд идущие `_` через regex.
    """
    return re.sub(r"_+", "", raw).strip()


def _is_field_label(text: str, label: str) -> bool:
    """
    Назначение:
        Проверяет, начинается ли текст с указанной метки поля (регистр не важен).

    Вход:
        text: Текст для проверки.
        label: Ожидаемая метка (например, `единицы измерения:`).

    Выход:
        True/False.
    """
    return text.lower().startswith(label.lower())
# --- Поиск полей по подписям (якорям) ---
# PURPOSE: Раскладка карточки у заказчика сдвинута относительно прежней (B2→C6, E11→F15,
# M11→N15; на листе методики B/F→C/G), поэтому поля ищутся по подписям на листе, а прежние
# фиксированные адреса остаются запасным вариантом для файлов старого образца.
_MAX_SCAN_ROWS = 80
_MAX_SCAN_COLS = 30
_SIGNEE_FIO_RE = re.compile(
    r"^(?:[А-ЯЁ]\.\s*[А-ЯЁ]\.\s*[А-ЯЁ][а-яё\-]+|[А-ЯЁ][а-яё\-]+\s+[А-ЯЁ]\.\s*[А-ЯЁ]\.)$"
)
_SIGNEE_DATE_RE = re.compile(
    r"\d{1,2}\s*[\"»']?\s*(?:январ|феврал|март|апрел|ма[йя]|июн|июл|август|сентябр|октябр|ноябр|декабр)"
    r"|\d{1,2}\.\d{1,2}\.\d{2,4}",
    re.IGNORECASE,
)
_METODIKA_LABELS = ("единицы измерения", "способ расчет", "источник данных", "методика расчет", "карточка проекта")


def _norm(value: Any) -> str:
    """Текст ячейки в нижнем регистре с одиночными пробелами — для сравнения с подписью поля."""
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip().lower()


def _text(value: Any) -> str:
    """Значение ячейки как stripped-строка; `None` и пустые значения → пустая строка."""
    if value is None:
        return ""
    return str(value).strip()


def _iter_cells(ws: Any):
    """Непустые ячейки листа в пределах области сканирования (строки и столбцы ограничены)."""
    for row in ws.iter_rows(
        min_row=1,
        max_row=min(ws.max_row, _MAX_SCAN_ROWS),
        max_col=min(ws.max_column, _MAX_SCAN_COLS),
    ):
        for cell in row:
            if cell.value not in (None, ""):
                yield cell


def _find_label(ws: Any, *labels: str, contains: bool = False) -> Any:
    """
    Назначение:
        Находит ячейку с подписью поля («Клиенты процесса:», «Утверждаю», «Начало»).

    Вход:
        ws: лист openpyxl.
        labels: варианты подписи (регистр и лишние пробелы не важны).
        contains: искать вхождение подписи в текст, а не начало строки.

    Выход:
        Ячейка openpyxl или `None`, если подпись не найдена.
    """
    wanted = [_norm(label) for label in labels]
    for cell in _iter_cells(ws):
        text = _norm(cell.value)
        for label in wanted:
            if (label in text) if contains else text.startswith(label):
                return cell
    return None


def _value_right(ws: Any, cell: Any, max_offset: int = 12) -> Any:
    """Значение первой непустой ячейки справа от подписи в той же строке; `None`, если нет."""
    if cell is None:
        return None
    for col in range(cell.column + 1, min(cell.column + max_offset, ws.max_column) + 1):
        value = ws.cell(row=cell.row, column=col).value
        if value not in (None, ""):
            return value
    return None


def _value_below(ws: Any, cell: Any, max_offset: int = 6) -> Any:
    """Значение первой непустой ячейки ниже подписи в том же столбце; `None`, если нет."""
    if cell is None:
        return None
    for row in range(cell.row + 1, min(cell.row + max_offset, ws.max_row) + 1):
        value = ws.cell(row=row, column=cell.column).value
        if value not in (None, ""):
            return value
    return None


def _value_above(ws: Any, cell: Any, max_offset: int = 6) -> Any:
    """Значение ближайшей непустой ячейки выше подписи в том же столбце; `None`, если нет."""
    if cell is None:
        return None
    for row in range(cell.row - 1, max(cell.row - max_offset, 0), -1):
        value = ws.cell(row=row, column=cell.column).value
        if value not in (None, ""):
            return value
    return None


_SECTION_TITLE_RE = re.compile(r"^\s*\d+\s*[.)]")


def _looks_like_date(value: Any) -> bool:
    """Похоже ли значение на дату: `datetime`, «"30" августа 2026 г.», «30.08.2026», «2026-08-30»."""
    if isinstance(value, datetime):
        return True
    text = _text(value)
    if not text:
        return False
    return bool(
        _SIGNEE_DATE_RE.search(text)
        or re.match(r"^\d{4}-\d{2}-\d{2}", text)
        or re.match(r"^\d{1,2}[./]\d{1,2}[./]\d{2,4}$", text)
    )


def _date_below(ws: Any, cell: Any) -> Optional[str]:
    """Дата под заголовком столбца («База», «Цель», «Идеал»); не дата или пусто → `None`."""
    if cell is None:
        return None
    for row in range(cell.row + 1, min(cell.row + 3, ws.max_row) + 1):
        value = ws.cell(row=row, column=cell.column).value
        if value in (None, ""):
            continue
        return _format_date(value) if _looks_like_date(value) else None
    return None


def _value_under_title(ws: Any, cell: Any) -> str:
    """
    Значение под подписью-заголовком («Карточка проекта:» → название потока).

    Берётся ближайшая непустая ячейка в пределах двух строк; заголовок раздела
    («1. Вовлеченные лица…») и подписи полей значением не считаются — иначе при пустом
    названии потока в карточку попадал бы текст соседнего раздела.
    """
    value = _text(_value_below(ws, cell, 2))
    if not value or _SECTION_TITLE_RE.match(value) or _norm(value).startswith(_METODIKA_LABELS):
        return ""
    return value


def _tail_after_label(value: Any, label: str) -> str:
    """Остаток строки после подписи: «Руководитель проекта: Иванов И.И.» → «Иванов И.И.»."""
    return re.sub(rf"^\s*{label}\s*:?\s*", "", _text(value), flags=re.IGNORECASE).strip()


def _parse_signee_block(ws: Any, anchor: Any) -> tuple:
    """
    Назначение:
        Разбирает блок под грифом «Утверждаю»: должность, ФИО и дату утверждения.

    Вход:
        ws: лист openpyxl; anchor: ячейка с грифом «Утверждаю».

    Выход:
        Кортеж `(должность, ФИО, дата)`; ненайденное — пустая строка.

    Логика:
        Обходит прямоугольник под грифом (до 7 строк вниз, столбцы вокруг грифа) сверху вниз.
        Дата распознаётся по типу `datetime` или шаблону («"30" августа 2026 г.», «30.08.2026»),
        ФИО — по шаблону инициалов; первая оставшаяся непустая строка считается должностью.
        Линии подписи из подчёркиваний отбрасываются.
    """
    position = name = date_value = ""
    for row in range(anchor.row + 1, min(anchor.row + 8, ws.max_row) + 1):
        for col in range(max(anchor.column - 2, 1), min(anchor.column + 10, ws.max_column) + 1):
            raw = ws.cell(row=row, column=col).value
            if raw in (None, ""):
                continue
            if isinstance(raw, datetime):
                date_value = date_value or _format_date(raw) or ""
                continue
            # Ячейка может содержать несколько строк (подпись и дата) — разбираем построчно.
            for line in str(raw).splitlines():
                clean = re.sub(r"_+", "", line).strip()
                if not clean:
                    continue
                if _SIGNEE_DATE_RE.search(clean):
                    date_value = date_value or clean
                elif _SIGNEE_FIO_RE.match(clean):
                    name = name or clean
                elif not position:
                    position = clean
    return position, name, date_value
# END_SHARED_HELPERS


# START_PARSE_DROPDOWN
# PURPOSE: Парсит лист «выпадающий список» — справочник допустимых единиц измерения по категориям.
# INPUTS: Путь к xlsx, путь к output_dir.
# OUTPUTS: Dict `{categories: {header: [unit, ...], ...}}`. Сохраняет также `dropdown_units.json` в output_dir.
# KEYWORDS: dropdown, units, categories.
def parse_dropdown(xlsx_path: Path, output_dir: Path) -> DropdownDocument:
    """
    Назначение:
        Парсит лист с выпадающим списком единиц измерения.

    Вход:
        xlsx_path: Путь к xlsx карточки проекта.
        output_dir: Куда сохранить `dropdown_units.json`.

    Выход:
        Dict `{categories: Dict[str, List[str]]}`. Если лист не найден — `{categories: {}}`.

    Логика:
        1. Создаёт output_dir если нужно.
        2. Ищет лист по подстроке `выпадающий список`. Нет → сохраняет пустой JSON.
        3. Первая строка (A1..D1) — заголовки категорий; значения строк 2+ —
           единицы измерения соответствующей категории.
        4. Собирает словарь `header → [unit, ...]`, пустые значения пропускает.
        5. Сохраняет результат в `dropdown_units.json` и возвращает его.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = _find_sheet_by_keyword(wb, "выпадающий список")

    if ws is None:
        wb.close()
        result: Dict[str, Any] = {"categories": {}}
        _save_json(result, output_dir / "dropdown_units.json")
        return result

    categories: Dict[str, List[str]] = {}
    columns = ["A", "B", "C", "D"]
    # Заголовки категорий — первая строка в каждой из 4 колонок.
    headers = [(_cell_str(ws, f"{col}1") or None) for col in columns]
    for col, header in zip(columns, headers):
        if header is None:
            continue
        units: List[str] = []
        for row in range(2, ws.max_row + 1):
            val = ws[f"{col}{row}"].value
            if val is not None:
                unit = str(val).strip()
                if unit:
                    units.append(unit)
        categories[header] = units

    wb.close()
    result = {"categories": categories}
    _save_json(result, output_dir / "dropdown_units.json")
    return result
# END_PARSE_DROPDOWN


# START_PARSE_KARTOCHKA_MAIN
# PURPOSE: Парсит основной лист «Карточка проекта»: шапку, секции 1 и 2, показатели, события.
# INPUTS: Путь к xlsx, путь к output_dir.
# OUTPUTS: Структурированный Dict с meta/header/section1/section2/indicators/indicator_dates/events.
# KEYWORDS: kartochka-main, header, indicators, events, anchors.
def _parse_main_header(ws: Any) -> Dict[str, Any]:
    """
    Назначение:
        Шапка карточки: организация, название потока, реквизиты подписанта.

    Вход:
        ws: лист «Карточка проекта».

    Выход:
        Dict `{org_name, project_name, signee_position, signee_name, signee_date, has_utverzhday}`.

    Логика:
        Организация и поток — над и под подписью «Карточка проекта:»; подписант — блок под
        грифом «Утверждаю». Если подписи не найдены, читаются прежние адреса (B2/B4/K3/K4/K7/K8).
    """
    title = _find_label(ws, "карточка проекта")
    org_name = _text(_value_above(ws, title)) if title is not None else ""
    project_name = _value_under_title(ws, title) if title is not None else ""
    if not org_name:
        org_name = _cell_str(ws, "B2")
    if not project_name:
        project_name = _cell_str(ws, "B4")

    utverzhday = _find_label(ws, "утверждаю", contains=True)
    position = name = date_value = ""
    if utverzhday is not None:
        position, name, date_value = _parse_signee_block(ws, utverzhday)
    if not (position or name or date_value):
        position = _cell_str(ws, "K4")
        name = _extract_fio(_cell_str(ws, "K7"))
        date_value = _cell_str(ws, "K8")

    return {
        "org_name": org_name,
        "project_name": project_name,
        "signee_position": position,
        "signee_name": _extract_fio(name),
        "signee_date": date_value,
        "has_utverzhday": utverzhday is not None or "утверждаю" in _cell_str(ws, "K3").lower(),
    }


def _parse_main_section1(ws: Any) -> Dict[str, Any]:
    """
    Назначение:
        Секция 1 «Вовлечённые лица и рамки проекта».

    Вход:
        ws: лист «Карточка проекта».

    Выход:
        Dict `{clients, perimeter, owner, boundaries, leader, team}`; ненайденное — пустая строка.

    Логика:
        Клиенты/периметр/владелец/границы — значение справа от подписи. Руководитель и команда
        записаны одной строкой вместе с подписью, поэтому берётся остаток ячейки, иначе — соседняя
        справа. Запасной вариант — прежние адреса E11..E14, C15, C16.
    """
    simple = (
        ("clients", ("клиенты процесса",), "E11"),
        ("perimeter", ("периметр проекта", "периметр процесса"), "E12"),
        ("owner", ("владелец процесса",), "E13"),
        ("boundaries", ("границы процесса",), "E14"),
    )
    section: Dict[str, Any] = {}
    for key, labels, legacy_addr in simple:
        cell = _find_label(ws, *labels)
        value = _text(_value_right(ws, cell)) if cell is not None else ""
        section[key] = value or _cell_str(ws, legacy_addr)

    for key, label, legacy_addr in (
        ("leader", "руководитель проекта", "C15"),
        ("team", "команда проекта", "C16"),
    ):
        cell = _find_label(ws, label)
        value = ""
        if cell is not None:
            value = _tail_after_label(cell.value, label) or _text(_value_right(ws, cell))
        section[key] = value or _tail_after_label(_cell_str(ws, legacy_addr), label)
    return section


def _parse_main_section2(ws: Any) -> Dict[str, Any]:
    """
    Назначение:
        Секция 2 «Обоснование выбора»: ключевой риск и обоснование.

    Вход:
        ws: лист «Карточка проекта».

    Выход:
        Dict `{key_risk, justification}`.

    Логика:
        Значения справа от подписей «Ключевой риск:» и «Обоснование:». Заголовок секции
        «2. Обоснование выбора» подписью не считается (начинается с номера). Запас — M11/M13.
    """
    risk_cell = _find_label(ws, "ключевой риск")
    just_cell = _find_label(ws, "обоснование")
    return {
        "key_risk": (_text(_value_right(ws, risk_cell)) if risk_cell is not None else "") or _cell_str(ws, "M11"),
        "justification": (_text(_value_right(ws, just_cell)) if just_cell is not None else "") or _cell_str(ws, "M13"),
    }


def _parse_main_indicators(ws: Any) -> tuple:
    """
    Назначение:
        Таблица показателей секции 3 и даты замеров база/цель/идеал.

    Вход:
        ws: лист «Карточка проекта».

    Выход:
        Кортеж `(indicator_dates, indicators)`.

    Логика:
        Столбцы определяются по заголовкам таблицы («Наименование показателя», «ед. изм.»,
        «База», «Цель», «Идеал»), даты замеров — строка под заголовками База/Цель/Идеал,
        строки показателей — ниже заголовка, пока встречается наименование. Если заголовки
        не найдены, читаются прежние адреса (F21/G21/H21 и строки 22..39 колонок C..H).
    """
    name_cell = _find_label(ws, "наименование показателя")
    unit_cell = _find_label(ws, "ед. изм", "ед.изм", "единица измерения", "единицы измерения")
    base_cell = _find_label(ws, "база")
    target_cell = _find_label(ws, "цель")
    ideal_cell = _find_label(ws, "идеал")

    dates = {
        "base_date": _date_below(ws, base_cell),
        "target_date": _date_below(ws, target_cell),
        "ideal_date": _date_below(ws, ideal_cell),
    }
    if not any(dates.values()):
        legacy = {addr: _cell_value(ws, addr) for addr in ("F21", "G21", "H21")}
        dates = {
            "base_date": _format_date(legacy["F21"]) if _looks_like_date(legacy["F21"]) else None,
            "target_date": _format_date(legacy["G21"]) if _looks_like_date(legacy["G21"]) else None,
            "ideal_date": _format_date(legacy["H21"]) if _looks_like_date(legacy["H21"]) else None,
        }

    indicators: List[Dict[str, Any]] = []
    if name_cell is not None and base_cell is not None:
        header_row = max(name_cell.row, base_cell.row)
        number_col = name_cell.column - 1
        for row in range(header_row + 1, min(ws.max_row, header_row + 25) + 1):
            name = _text(ws.cell(row=row, column=name_cell.column).value)
            if not name:
                continue
            raw_number = ws.cell(row=row, column=number_col).value if number_col >= 1 else None
            try:
                number = int(raw_number)
            except (TypeError, ValueError):
                number = len(indicators) + 1
            unit = _text(ws.cell(row=row, column=unit_cell.column).value) if unit_cell is not None else ""
            indicators.append({
                "number": number,
                "name": name,
                "unit": unit or None,
                "base_value": ws.cell(row=row, column=base_cell.column).value,
                "target_value": ws.cell(row=row, column=target_cell.column).value if target_cell is not None else None,
                "ideal_value": ws.cell(row=row, column=ideal_cell.column).value if ideal_cell is not None else None,
            })

    if not indicators:
        # Прежняя раскладка: номер в C, наименование в D, единица в E, значения в F/G/H.
        for row in range(22, 40):
            c_val = _cell_value(ws, f"C{row}")
            if c_val is None:
                break
            try:
                number = int(c_val)
            except (ValueError, TypeError):
                break
            d_val = _cell_str(ws, f"D{row}")
            e_val = _cell_str(ws, f"E{row}")
            f_val = _cell_value(ws, f"F{row}")
            if not d_val and not e_val and f_val is None:
                continue
            indicators.append({
                "number": number,
                "name": d_val if d_val else None,
                "unit": e_val if e_val else None,
                "base_value": f_val,
                "target_value": _cell_value(ws, f"G{row}"),
                "ideal_value": _cell_value(ws, f"H{row}"),
            })
    return dates, indicators


def _parse_main_events(ws: Any) -> List[Dict[str, Any]]:
    """
    Назначение:
        Список ключевых событий секции 4 с датами начала и окончания.

    Вход:
        ws: лист «Карточка проекта».

    Выход:
        Список `{name, start_date, end_date}`.

    Логика:
        Столбцы дат — по заголовкам «Начало» и «Окончание», столбец названия — от заголовка
        секции «Ключевые события» до столбца даты начала. Если заголовков нет — прежняя
        раскладка (названия в K, даты в Q и S, строки 20..39).
    """
    start_cell = _find_label(ws, "начало")
    end_cell = _find_label(ws, "окончание")
    title_cell = _find_label(ws, "ключевые события", contains=True)

    events: List[Dict[str, Any]] = []
    if start_cell is not None and title_cell is not None:
        header_row = max(start_cell.row, title_cell.row)
        for row in range(header_row + 1, min(ws.max_row, header_row + 30) + 1):
            name = ""
            for col in range(title_cell.column, max(start_cell.column, title_cell.column + 1)):
                name = _text(ws.cell(row=row, column=col).value)
                if name:
                    break
            start_value = ws.cell(row=row, column=start_cell.column).value
            end_value = ws.cell(row=row, column=end_cell.column).value if end_cell is not None else None
            if not name and start_value is None:
                continue
            events.append({
                "name": name,
                "start_date": _format_date(start_value),
                "end_date": _format_date(end_value),
            })

    if not events:
        for row in range(20, 40):
            k_val = _cell_str(ws, f"K{row}")
            q_val = _cell_value(ws, f"Q{row}")
            if not k_val and q_val is None:
                continue
            if k_val and "ключевые события" in k_val.lower():
                continue
            events.append({
                "name": k_val.strip() if k_val else "",
                "start_date": _format_date(q_val),
                "end_date": _format_date(_cell_value(ws, f"S{row}")),
            })
    return events


def parse_kartochka_main(xlsx_path: Path, output_dir: Path) -> KartochkaMainDocument:
    """
    Назначение:
        Парсит основной лист карточки проекта в структурированный словарь.

    Вход:
        xlsx_path: Путь к xlsx.
        output_dir: Куда сохранить `kartochka_main.json`.

    Выход:
        Dict с полями `meta`, `header`, `section1`, `section2`, `indicator_dates`,
        `indicators` (список), `events` (список).

    Логика:
        1. Ищет лист «Карточка проекта», исключая шаблоны.
        2. Поля читаются по подписям на листе (см. `_find_label`): раскладка формы у заказчика
           менялась, фиксированные адреса давали пустой результат по всем полям сразу.
           Прежние адреса остаются запасным вариантом для файлов старого образца.
        3. Шапка, секция 1, секция 2, показатели с датами замеров, ключевые события.
        4. Если не распознано ни одно поле и нет ни показателей, ни событий — поднимается
           `ValueError`: пустой результат нельзя выдавать за «всё поля пустые» (иначе отчёт
           состоит из выдуманных замечаний). JSON при этом сохраняется для разбора.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = _find_sheet_by_keyword(wb, "Карточка проекта", exclude="ШАБЛОН")

    # Fallback: если явно не нашли — берём второй лист (если есть) или первый.
    if ws is None:
        if len(wb.sheetnames) > 1:
            ws = wb[wb.sheetnames[1]]
        else:
            ws = wb[wb.sheetnames[0]]

    indicator_dates, indicators = _parse_main_indicators(ws)
    result: Dict[str, Any] = {
        "meta": {"workbook": Path(xlsx_path).name, "sheet": ws.title, "max_row": ws.max_row},
        "header": _parse_main_header(ws),
        "section1": _parse_main_section1(ws),
        "section2": _parse_main_section2(ws),
        "indicators": indicators,
        "indicator_dates": indicator_dates,
        "events": _parse_main_events(ws),
    }

    wb.close()
    _save_json(result, output_dir / "kartochka_main.json")

    recognized = [
        result["header"]["org_name"],
        result["header"]["project_name"],
        *result["section1"].values(),
        *result["section2"].values(),
    ]
    if not any(recognized) and not result["indicators"] and not result["events"]:
        raise ValueError(
            f"Форма карточки проекта не распознана: на листе «{ws.title}» не найдены ни реквизиты "
            "шапки, ни разделы карточки. Проверьте, что загружена карточка проекта 2.4."
        )
    return result
# END_PARSE_KARTOCHKA_MAIN


# START_PARSE_METODIKA
# PURPOSE: Парсит лист «Методика расчёта» — для каждого показателя читает единицы, способ расчёта, источник данных.
# INPUTS: Путь к xlsx, путь к output_dir.
# OUTPUTS: Dict с `meta`, `header`, `indicators`.
# KEYWORDS: metodika, indicators, calc-method, anchors.
def parse_metodika(xlsx_path: Path, output_dir: Path) -> MetodikaDocument:
    """
    Назначение:
        Парсит лист «Методика расчёта» в структуру с показателями и их метаданными.

    Вход:
        xlsx_path: Путь к xlsx.
        output_dir: Куда сохранить `metodika.json`.

    Выход:
        Dict `{meta, header, indicators}`. Если лист не найден — пустой каркас.

    Логика:
        Лист устроен блоками: строка с названием показателя, под ней три строки с подписями
        «Единицы измерения:», «Способ расчета:», «Источник данных:»; значения — правее подписи.
        Столбцы подписей и значений в разных версиях формы разные (B/F против C/G), поэтому
        блоки ищутся по самим подписям, а не по адресам.

        Алгоритм:
        1. Находит все ячейки с подписью «Единицы измерения:».
        2. Для каждой берёт значение справа; из строк ниже — способ расчёта и источник данных.
        3. Название показателя — ближайшая непустая ячейка выше в том же столбце.
        4. Если подписи есть, а разобрать не удалось ни одного показателя — `ValueError`
           (иначе правила «единицы измерения» и «способ расчёта» проходят вхолостую).
        5. Если подписей нет вовсе — читает лист прежним способом (название в B, значения в F).
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = _find_sheet_by_keyword(wb, "Методика расчет", exclude="ШАБЛОН")

    if ws is None:
        wb.close()
        result: Dict[str, Any] = {"meta": {"sheet": None}, "header": {}, "indicators": []}
        _save_json(result, output_dir / "metodika.json")
        return result

    title = _find_label(ws, "карточка проекта")
    result = {
        "meta": {"sheet": ws.title},
        "header": {
            "org_name": (_text(_value_above(ws, title)) if title is not None else "") or _cell_str(ws, "B2"),
            "project_name": (_value_under_title(ws, title) if title is not None else "") or _cell_str(ws, "B4"),
        },
        "indicators": [],
    }

    unit_labels = [cell for cell in _iter_cells(ws) if _norm(cell.value).startswith("единицы измерения")]
    for label_cell in unit_labels:
        unit = _text(_value_right(ws, label_cell))
        calc_method = source = ""
        below_first = ws.cell(row=label_cell.row + 1, column=label_cell.column)
        if _norm(below_first.value).startswith("способ расчет"):
            calc_method = _text(_value_right(ws, below_first))
        below_second = ws.cell(row=label_cell.row + 2, column=label_cell.column)
        if _norm(below_second.value).startswith("источник данных"):
            source = _text(_value_right(ws, below_second))

        # Название показателя — ближайшая непустая ячейка выше, не являющаяся подписью поля.
        name = ""
        name_row = label_cell.row
        for row in range(label_cell.row - 1, max(label_cell.row - 4, 0), -1):
            raw = _text(ws.cell(row=row, column=label_cell.column).value)
            if not raw or _norm(raw).startswith(_METODIKA_LABELS):
                continue
            name = re.sub(r"\s*:\s*$", "", raw).strip()
            name_row = row
            break
        if not name:
            continue
        result["indicators"].append({
            "name": name,
            "unit": unit or None,
            "calc_method": calc_method or None,
            "data_source": source or None,
            "row_start": name_row,
        })

    if unit_labels and not result["indicators"]:
        wb.close()
        _save_json(result, output_dir / "metodika.json")
        raise ValueError(
            f"Лист «{ws.title}» не распознан: подписи «Единицы измерения» найдены, но ни один "
            "показатель не прочитан. Проверьте версию формы карточки проекта 2.4."
        )

    if not unit_labels:
        # Прежняя раскладка: название показателя в колонке B, значения — в колонке F.
        field_labels = ("единицы измерения", "способ расчет", "источник данных", "методика расчет")
        skip_keywords = ("карточка проекта", "методика расчет")
        row = 1
        max_row = ws.max_row
        while row <= max_row:
            b_val = _cell_str(ws, f"B{row}")
            if not b_val:
                row += 1
                continue
            b_lower = b_val.lower()
            if any(b_lower.startswith(lbl) for lbl in field_labels) or any(kw in b_lower for kw in skip_keywords):
                row += 1
                continue
            next_b = _cell_str(ws, f"B{row + 1}") if row + 1 <= max_row else ""
            if _is_field_label(next_b, "единицы измерения"):
                unit = _cell_str(ws, f"F{row + 1}")
                calc_method = _cell_str(ws, f"F{row + 2}") if row + 2 <= max_row else ""
                source = _cell_str(ws, f"F{row + 3}") if row + 3 <= max_row else ""
                result["indicators"].append({
                    "name": re.sub(r":$", "", b_val).strip(),
                    "unit": unit if unit else None,
                    "calc_method": calc_method if calc_method else None,
                    "data_source": source if source else None,
                    "row_start": row,
                })
                row += 4
                continue
            row += 1

    wb.close()
    _save_json(result, output_dir / "metodika.json")
    return result
# END_PARSE_METODIKA


# START_IO_HELPERS
# PURPOSE: Точечные утилиты ввода-вывода, используемые парсерами выше.
# INPUTS: Словарь результата и путь к JSON-файлу.
# OUTPUTS: Файл на диске.
# KEYWORDS: json-save, io.
def _save_json(data: Dict[str, Any], output_path: Path) -> None:
    """
    Назначение:
        Сохраняет словарь в JSON-файл с ensure_ascii=False и отступами.

    Вход:
        data: Словарь для сериализации.
        output_path: Путь к выходному файлу.

    Выход:
        None.

    Логика:
        Используется всеми тремя парсерами для сохранения промежуточных JSON,
        которые потом читают валидаторы.
    """
    with open(output_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
# END_IO_HELPERS
