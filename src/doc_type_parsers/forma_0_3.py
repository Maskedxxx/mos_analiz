# START_MODULE_CONTRACT
# PURPOSE: Парсер формы 0.3 «О предприятии в цифрах» (xlsx). Извлекает значения и формулы для 14 валидаторов forma_0_3 в JSON forma_0_3_main.json.
# INPUTS: путь к xlsx, output_dir для записи forma_0_3_main.json.
# OUTPUTS: dict-структура (см. STRUCTURE ниже) + файл forma_0_3_main.json в output_dir.
# KEYWORDS: forma_0_3, openpyxl, xlsx-parser, values, formulas, two-pass.
# LINKS: src/doc_type_validators/forma_0_3.py (потребитель данных), doc_configs/forma_0_3/.
# RATIONALE: Два прохода openpyxl: data_only=True (вычисленные значения), data_only=False (формулы — для проверки наличия формул прогноза в I/J/K). Логика перенесена из прод-движка as-is (источник правды).
# END_MODULE_CONTRACT

from __future__ import annotations

# START_IMPORTS
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Optional

import openpyxl
# END_IMPORTS


# START_STRUCTURE
# Структура forma_0_3_main.json:
# {
#   "meta": {"workbook", "sheet", "n_sheets"},
#   "headers": {"title" (B1), "section_general" (B6), "section_indicators" (B11)},
#   "general_info": {name B9, inn C9, okved D9, region E9, agreement_date_subject F9,
#                    agreement_date_fck G9, base_year H9, company_info I9,
#                    appendix_label I1, agreement_ref I2},
#   "years": [G12,H12,I12,J12,K12],
#   "performance": {"values": {строка 23}, "formulas": {строка 23}},
#   "index": {"values": {I24,J24,K24}, "formulas": {...}},
#   "target_index": {I25,J25,K25},
#   "indicators_data": {"<col><row>": {value, formula}} строки 13-22 × G:K,
#   "signatures": {consent G27, signer_fio B31, signature_date J28, doc_date B33},
#   "validations_meta": {has_inn_validation, has_okved_validation}
# }
# END_STRUCTURE


# START_HELPERS
def _v(ws: Any, coord: str) -> Any:
    """Достаёт значение из ячейки. Строку strip-ает, пустую → None."""
    v = ws[coord].value
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _to_float(v: Any) -> Optional[float]:
    """Безопасное преобразование к float (None если невозможно)."""
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        return None


def _to_iso_date(v: Any) -> Optional[str]:
    """Возвращает ISO дату YYYY-MM-DD или None."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
            try:
                return datetime.strptime(s, fmt).date().isoformat()
            except ValueError:
                pass
    return None


_RU_MONTHS = {
    "января": 1, "февраля": 2, "марта": 3, "апреля": 4, "мая": 5, "июня": 6,
    "июля": 7, "августа": 8, "сентября": 9, "октября": 10, "ноября": 11, "декабря": 12,
}


def _parse_russian_date(s: str) -> Optional[str]:
    """Парсит '04 марта 2026' / '04 марта 2026 г.' → '2026-03-04'."""
    if not s:
        return None
    parts = s.replace("г.", "").strip().split()
    if len(parts) < 3:
        return None
    try:
        day = int(parts[0])
        month = _RU_MONTHS.get(parts[1].lower())
        year = int(parts[2])
        if month is None:
            return None
        return f"{year:04d}-{month:02d}-{day:02d}"
    except (ValueError, KeyError):
        return None
# END_HELPERS


# START_PARSE_FORMA_0_3
def parse_forma_0_3(xlsx_path: Path, output_dir: Path) -> Dict[str, Any]:
    """Парсит xlsx формы 0.3, кладёт forma_0_3_main.json в output_dir, возвращает структуру."""
    # Проход 1: вычисленные значения; Проход 2: формулы.
    wb_v = openpyxl.load_workbook(xlsx_path, data_only=True)
    wb_f = openpyxl.load_workbook(xlsx_path, data_only=False)

    # Берём первый лист (основной — «О предприятии в цифрах»).
    ws_v = wb_v[wb_v.sheetnames[0]]
    ws_f = wb_f[wb_f.sheetnames[0]]

    def fval(coord: str) -> Optional[str]:
        """Формула из ячейки (или None если нет формулы)."""
        v = ws_f[coord].value
        if isinstance(v, str) and v.startswith("="):
            return v
        return None

    # === Заголовки ===
    headers = {
        "title": _v(ws_v, "B1"),
        "section_general": _v(ws_v, "B6"),
        "section_indicators": _v(ws_v, "B11"),
    }

    # === Общая информация (строка 9) ===
    general_info = {
        "name": _v(ws_v, "B9"),
        "inn": str(_v(ws_v, "C9")) if _v(ws_v, "C9") is not None else None,
        "okved": _v(ws_v, "D9"),
        "region": _v(ws_v, "E9"),
        "agreement_date_subject": _to_iso_date(ws_v["F9"].value),
        "agreement_date_fck": _to_iso_date(ws_v["G9"].value),
        "base_year": _to_float(_v(ws_v, "H9")),
        "company_info": _v(ws_v, "I9"),
        "appendix_label": _v(ws_v, "I1"),
        "agreement_ref": _v(ws_v, "I2"),
    }
    if general_info["base_year"] is not None:
        general_info["base_year"] = int(general_info["base_year"])

    # === Годы в заголовке таблицы показателей (строка 12) ===
    year_cells = ["G12", "H12", "I12", "J12", "K12"]
    years = []
    for c in year_cells:
        y = _to_float(_v(ws_v, c))
        years.append(int(y) if y is not None else None)

    # === Строка 23: производительность труда ===
    perf_cells = ["G23", "H23", "I23", "J23", "K23"]
    performance = {
        "values": {c: _to_float(_v(ws_v, c)) for c in perf_cells},
        "formulas": {c: fval(c) for c in perf_cells},
    }

    # === Строка 24: индекс производительности ===
    idx_cells = ["I24", "J24", "K24"]
    index_data = {
        "values": {c: _to_float(_v(ws_v, c)) for c in idx_cells},
        "formulas": {c: fval(c) for c in idx_cells},
    }

    # === Строка 25: целевые показатели индекса ===
    target_cells = ["I25", "J25", "K25"]
    target_index = {c: _to_float(_v(ws_v, c)) for c in target_cells}

    # === Данные строк 13–22 в столбцах G:K (заполненность и формулы) ===
    indicators_data: Dict[str, Any] = {}
    for row in range(13, 23):
        for col in ["G", "H", "I", "J", "K"]:
            coord = f"{col}{row}"
            indicators_data[coord] = {
                "value": _to_float(_v(ws_v, coord)),
                "formula": fval(coord),
            }

    # === Подписи ===
    signatures = {
        "consent": _v(ws_v, "G27"),
        "signer_fio": _v(ws_v, "B31"),
        "signature_date_raw": ws_v["J28"].value if ws_v["J28"].value is not None else None,
        "signature_date_iso": _to_iso_date(ws_v["J28"].value),
        "doc_date_raw": _v(ws_v, "B33"),
        "doc_date_iso": _to_iso_date(_v(ws_v, "B33")),
    }
    # Если doc_date — строка вида "04 марта 2026", парсим вручную.
    if signatures["doc_date_iso"] is None and isinstance(signatures["doc_date_raw"], str):
        signatures["doc_date_iso"] = _parse_russian_date(signatures["doc_date_raw"])
    # Сериализуем datetime.
    if signatures["signature_date_raw"] is not None and not isinstance(signatures["signature_date_raw"], str):
        signatures["signature_date_raw"] = str(signatures["signature_date_raw"])

    # === Метаданные валидаций (data validations на C9/D9) ===
    has_inn_val = False
    has_okved_val = False
    try:
        for dv in (ws_f.data_validations.dataValidation or []):
            sqref = str(dv.sqref) if dv.sqref else ""
            if "C9" in sqref:
                has_inn_val = True
            if "D9" in sqref:
                has_okved_val = True
    except Exception:
        pass

    result = {
        "meta": {
            "workbook": xlsx_path.name,
            "sheet": ws_v.title,
            "n_sheets": len(wb_v.sheetnames),
        },
        "headers": headers,
        "general_info": general_info,
        "years": years,
        "performance": performance,
        "index": index_data,
        "target_index": target_index,
        "indicators_data": indicators_data,
        "signatures": signatures,
        "validations_meta": {
            "has_inn_validation": has_inn_val,
            "has_okved_validation": has_okved_val,
        },
    }

    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / "forma_0_3_main.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2, default=str)

    return result
# END_PARSE_FORMA_0_3
