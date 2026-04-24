# START_MODULE_CONTRACT
# PURPOSE: Низкоуровневый парсер DOCX. Читает файл и отдаёт полный текст документа целиком плюс метаданные. Не знает ни про страницы, ни про правила, ни про LLM, ни про типы документов.
# INPUTS: Путь к DOCX, опциональный VLM endpoint для OCR картинок в шапке.
# OUTPUTS: Dict `{filename, path, raw_text}`.
# KEYWORDS: docx, python-docx, raw-text, full-document, header-ocr.
# LINKS: src/format_parsers/__init__.py, main.py::_parse_document, tests/test_smoke_parsers.py.
# RATIONALE: Весь DOCX-специфичный код живёт в одном файле. Сервис парсит формат, отдаёт сырой текст — нарезка на смысловые части происходит в верхних слоях.
# END_MODULE_CONTRACT

from __future__ import annotations

# START_IMPORTS
import logging
from pathlib import Path
from typing import Any, List, Optional

from src.format_parsers._types import ParsedDocument
# END_IMPORTS


# START_LOGGER
# PURPOSE: Отдельный логгер DOCX-парсера. Имя совпадает с именем модуля для прозрачной трассировки.
docx_parser_logger = logging.getLogger(__name__)
# END_LOGGER


# START_DOCX_PARSER
# PURPOSE: Прочитать DOCX и отдать полный текст документа вместе с метаданными.
# INPUTS: Путь к DOCX, опциональный VLM endpoint для OCR картинок в шапке.
# OUTPUTS: Dict `{filename, path, raw_text}`.
# KEYWORDS: docx, raw-text, full-document, public-api.
# RATIONALE: Сервис парсит формат — смысловая нарезка документа делается выше по конвейеру (doc_type / LLM).
def parse_docx(
    file_path: str,
    vlm_base_url: Optional[str] = None,
) -> ParsedDocument:
    """
    Назначение:
        Читает DOCX-файл целиком и возвращает полный текст документа без деления на
        страницы или чанки.

    Вход:
        file_path: Путь к DOCX-файлу.
        vlm_base_url: Адрес VLM/OCR сервиса для распознавания картинок в шапке.

    Выход:
        dict: Словарь `{filename, path, raw_text}`.

    Логика:
        1. Открывает DOCX через python-docx.
        2. При необходимости OCR-ит картинки в шапке и ставит текст в самое начало.
        3. Извлекает все текстовые элементы и таблицы в правильном порядке.
        4. Склеивает всё в единый `raw_text` без разбиения по страницам.
    """
    try:
        from docx import Document
    except ImportError as exc:
        raise ModuleNotFoundError("python-docx is required for DOCX parsing") from exc

    file_path_obj = Path(file_path)
    doc = Document(str(file_path_obj))
    docx_parser_logger.info(
        "DOCX: %s, paragraphs=%s, tables=%s",
        file_path_obj.name,
        len(doc.paragraphs),
        len(doc.tables),
    )

    # Дополняем текст шапки OCR-результатом, если в начале документа есть картинки.
    header_ocr_text = _ocr_header_images(doc, vlm_base_url)
    if header_ocr_text:
        docx_parser_logger.info("  Header OCR chars: %s", len(header_ocr_text))

    # Сохраняем фактический порядок элементов документа: абзацы, текстбоксы, таблицы.
    elements = _extract_elements_in_order(doc)
    if header_ocr_text:
        elements.insert(0, header_ocr_text)

    # Склеиваем весь документ в один raw_text. Парсер не знает про страницы и правила —
    # он отдаёт всё содержимое как есть, а решение «что важно» принимают верхние слои.
    raw_text = "\n\n".join(elements)
    docx_parser_logger.info(
        "  Elements: %s, raw_text chars: %s", len(elements), len(raw_text)
    )

    return {
        "filename": file_path_obj.name,
        "path": str(file_path_obj.resolve()),
        "raw_text": raw_text,
    }
# END_DOCX_PARSER


# START_DOCX_HELPERS
# PURPOSE: Приватные хелперы парсинга DOCX, нужны только этому модулю.
# INPUTS: Объект Document python-docx, опциональный VLM endpoint.
# OUTPUTS: Распознанный текст шапки, упорядоченный список элементов и HTML таблицы.
# KEYWORDS: docx-internals, header-ocr, ordered-extraction, table-html.
def _ocr_header_images(doc: Any, vlm_base_url: Optional[str] = None) -> str:
    """
    Назначение:
        Извлекает и OCR-ит изображения из начальных абзацев DOCX.

    Вход:
        doc: Загруженный объект DOCX.
        vlm_base_url: Адрес VLM/OCR сервиса.

    Выход:
        str: Объединённый текст, распознанный из картинок в шапке.

    Логика:
        1. Ищет embedded-картинки в первых абзацах.
        2. Фильтрует слишком маленькие изображения.
        3. Отправляет изображение в OCR/VLM сервис.
        4. Склеивает распознанный текст в одну строку шапки.
    """
    if not vlm_base_url:
        return ""

    import base64
    import io

    import requests
    from PIL import Image

    ns_a = "http://schemas.openxmlformats.org/drawingml/2006/main"
    ns_r = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    ocr_parts: List[str] = []

    for paragraph in doc.paragraphs[:6]:
        # Смотрим только начало документа, потому что шапка обычно живёт именно там.
        blips = paragraph._element.findall(f".//{{{ns_a}}}blip")
        for blip in blips:
            rel_id = blip.get(f"{{{ns_r}}}embed")
            if not rel_id or rel_id not in doc.part.rels:
                continue
            try:
                blob = doc.part.rels[rel_id].target_part.blob
                image = Image.open(io.BytesIO(blob))
                # Отбрасываем мелкие декоративные картинки, чтобы не шуметь в OCR.
                if image.width < 200 or image.height < 50:
                    continue
                buf = io.BytesIO()
                image.save(buf, format="PNG")
                image_b64 = base64.b64encode(buf.getvalue()).decode()
                response = requests.post(
                    f"{vlm_base_url}chat/completions",
                    json={
                        "model": "PaddleOCR-VL-1.5",
                        "messages": [
                            {
                                "role": "user",
                                "content": [
                                    {
                                        "type": "image_url",
                                        "image_url": {
                                            "url": f"data:image/png;base64,{image_b64}",
                                        },
                                    },
                                    {
                                        "type": "text",
                                        "text": "OCR this image. Return all text exactly as written.",
                                    },
                                ],
                            }
                        ],
                        "max_tokens": 256,
                        "temperature": 0.0,
                    },
                    timeout=30,
                )
                text = response.json()["choices"][0]["message"]["content"].strip()
                if text:
                    ocr_parts.append(text)
                    docx_parser_logger.info(
                        "  OCR image (%sx%s): %s",
                        image.width,
                        image.height,
                        text[:80],
                    )
            except Exception as exc:
                docx_parser_logger.warning("  OCR image failed: %s", exc)

    return "\n".join(ocr_parts)


def _extract_elements_in_order(doc: Any) -> List[str]:
    """
    Назначение:
        Извлекает элементы DOCX в реальном порядке их появления.

    Вход:
        doc: Загруженный объект DOCX.

    Выход:
        list[str]: Список текстовых элементов и HTML-таблиц в правильном порядке.

    Логика:
        1. Обходит XML body документа.
        2. Собирает обычные абзацы и тексты из текстбоксов.
        3. Конвертирует таблицы в HTML.
        4. Возвращает единый линейный список элементов.
    """
    from docx.oxml.ns import qn

    elements: List[str] = []
    body = doc.element.body
    table_elements = {table._element: table for table in doc.tables}

    for child in body:
        tag = child.tag
        if tag == qn("w:p"):
            # Собираем текст по run-узлам, чтобы не потерять дробление на уровне Word XML.
            runs_text: List[str] = []
            for run in child.findall(qn("w:r")):
                for text_node in run.findall(qn("w:t")):
                    if text_node.text:
                        runs_text.append(text_node.text)
            full_text = "".join(runs_text).strip()
            if full_text:
                elements.append(full_text)

            ns_w = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            ns_wps = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"
            # Отдельно вытаскиваем текст из текстбоксов, который не попадает в обычные paragraphs.
            for txbx_content in child.findall(f".//{{{ns_wps}}}txbx/{{{ns_w}}}txbxContent"):
                txbx_runs: List[str] = []
                for paragraph in txbx_content.findall(f"{{{ns_w}}}p"):
                    for run in paragraph.findall(f"{{{ns_w}}}r"):
                        for text_node in run.findall(f"{{{ns_w}}}t"):
                            if text_node.text:
                                txbx_runs.append(text_node.text)
                txbx_text = "".join(txbx_runs).strip()
                if txbx_text:
                    elements.append(txbx_text)
        elif tag == qn("w:tbl"):
            table = table_elements.get(child)
            if table:
                html = docx_parser__table_to_html(table)
                if html.strip():
                    elements.append(html)
    return elements


def docx_parser__table_to_html(table: Any) -> str:
    """
    Назначение:
        Преобразует таблицу DOCX в простой HTML.

    Вход:
        table: Объект таблицы из python-docx.

    Выход:
        str: HTML-представление таблицы.

    Логика:
        1. Проходит по строкам и ячейкам таблицы.
        2. Нормализует текст ячеек.
        3. Собирает итоговую HTML-таблицу.
    """
    rows_html: List[str] = []
    for row in table.rows:
        cells_html = []
        for cell in row.cells:
            cell_text = cell.text.strip().replace("\n", " ")
            cells_html.append(f"  <td>{cell_text}</td>")
        rows_html.append("<tr>\n" + "\n".join(cells_html) + "\n</tr>")
    if not rows_html:
        return ""
    return "<table>\n" + "\n".join(rows_html) + "\n</table>"
# END_DOCX_HELPERS
