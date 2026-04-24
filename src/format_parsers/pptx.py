# START_MODULE_CONTRACT
# PURPOSE: Низкоуровневый парсер PPTX. Читает презентацию и отдаёт весь её текст целиком плюс метаданные. Не знает ни про правила, ни про LLM, ни про типы документов.
# INPUTS: Путь к PPTX.
# OUTPUTS: Dict `{filename, path, raw_text}`.
# KEYWORDS: pptx, python-pptx, raw-text, full-document, slides.
# LINKS: src/format_parsers/__init__.py, main.py::_parse_document, tests/test_smoke_parsers.py.
# RATIONALE: Весь PPTX-специфичный код живёт в одном файле. Сервис парсит формат — смысловая нарезка делается выше по конвейеру.
# END_MODULE_CONTRACT

from __future__ import annotations

# START_IMPORTS
import logging
from pathlib import Path
from typing import Any, List

from src.format_parsers._types import ParsedDocument
# END_IMPORTS


# START_LOGGER
# PURPOSE: Отдельный логгер PPTX-парсера. Имя логгера совпадает с именем модуля для прозрачной трассировки.
pptx_parser_logger = logging.getLogger(__name__)
# END_LOGGER


# START_PPTX_PARSER
# PURPOSE: Прочитать PPTX и отдать полный текст презентации вместе с метаданными.
# INPUTS: Путь к PPTX.
# OUTPUTS: Dict `{filename, path, raw_text}`.
# KEYWORDS: pptx, raw-text, full-document, public-api.
# RATIONALE: Сервис парсит формат — смысловая нарезка презентации делается выше по конвейеру (doc_type / LLM).
def parse_pptx(file_path: str) -> ParsedDocument:
    """
    Назначение:
        Читает PPTX-файл целиком и возвращает полный текст презентации без разбивки
        по слайдам.

    Вход:
        file_path: Путь к PPTX-файлу.

    Выход:
        dict: Словарь `{filename, path, raw_text}`.

    Логика:
        1. Открывает презентацию через python-pptx.
        2. Проходит по слайдам и собирает текст каждого слайда в визуальном порядке.
        3. Склеивает все слайды в единый `raw_text` с маркерами границ `[СЛАЙД N]`,
           чтобы не терять порядок при анализе на верхнем уровне.
    """
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise ModuleNotFoundError("python-pptx is required for PPTX parsing") from exc

    file_path_obj = Path(file_path)
    presentation = Presentation(str(file_path_obj))
    total_slides = len(presentation.slides)
    pptx_parser_logger.info("PPTX: %s, slides=%s", file_path_obj.name, total_slides)

    # Собираем текст всей презентации в один буфер. Пустые слайды пропускаем, чтобы не
    # засорять вывод маркерами без содержимого.
    slides_parts: List[str] = []
    for slide_idx, slide in enumerate(presentation.slides, start=1):
        text = _extract_slide_text(slide)
        if not text.strip():
            continue
        # Маркер `[СЛАЙД N]` сохраняет порядок и визуальные границы в сыром тексте.
        slides_parts.append(f"[СЛАЙД {slide_idx}]\n{text}")

    raw_text = "\n\n".join(slides_parts)
    pptx_parser_logger.info(
        "  Slides with text: %s, raw_text chars: %s", len(slides_parts), len(raw_text)
    )

    return {
        "filename": file_path_obj.name,
        "path": str(file_path_obj.resolve()),
        "raw_text": raw_text,
    }
# END_PPTX_PARSER


# START_PPTX_HELPERS
# PURPOSE: Приватные хелперы парсинга PPTX, нужны только этому модулю.
# INPUTS: Объекты Slide/TextFrame/Table python-pptx.
# OUTPUTS: Линейный текст слайда, текст одного текстового блока и HTML таблицы.
# KEYWORDS: pptx-internals, slide-text, visual-order, table-html.
def _extract_slide_text(slide: Any) -> str:
    """
    Назначение:
        Извлекает весь значимый текст со слайда.

    Вход:
        slide: Объект слайда из python-pptx.

    Выход:
        str: Линейный текст слайда с сохранением визуального порядка блоков.

    Логика:
        1. Проходит по shape-элементам слайда.
        2. Извлекает текстовые блоки и таблицы.
        3. Сортирует их по координатам сверху вниз и слева направо.
        4. Склеивает результат в один текст.
    """
    items: List[tuple[int, int, str]] = []
    for shape in slide.shapes:
        top = shape.top if shape.top is not None else 0
        left = shape.left if shape.left is not None else 0
        if shape.has_table:
            html = pptx_parser__table_to_html(shape.table)
            if html.strip():
                items.append((top, left, html))
        elif shape.has_text_frame:
            text = _textframe_to_text(shape.text_frame)
            if text.strip():
                items.append((top, left, text))
    # Сортировка нужна, чтобы текст следовал визуальному порядку на слайде.
    items.sort(key=lambda item: (item[0], item[1]))
    return "\n\n".join((item[2] for item in items))


def _textframe_to_text(text_frame: Any) -> str:
    """
    Назначение:
        Собирает текст из одного текстового блока PPTX.

    Вход:
        text_frame: Объект TextFrame.

    Выход:
        str: Текст блока с сохранением абзацев.

    Логика:
        1. Проходит по абзацам text frame.
        2. Склеивает тексты run-элементов внутри абзаца.
        3. Возвращает блок как многострочный текст.
    """
    paragraphs: List[str] = []
    for paragraph in text_frame.paragraphs:
        text = "".join((run.text for run in paragraph.runs))
        if text.strip():
            paragraphs.append(text.strip())
    return "\n".join(paragraphs)


def pptx_parser__table_to_html(table: Any) -> str:
    """
    Назначение:
        Преобразует таблицу PPTX в простой HTML.

    Вход:
        table: Объект таблицы из python-pptx.

    Выход:
        str: HTML-представление таблицы.

    Логика:
        1. Проходит по строкам и ячейкам таблицы.
        2. Собирает текст ячеек через текстовые блоки.
        3. Возвращает итоговую HTML-таблицу.
    """
    rows_html: List[str] = []
    for row in table.rows:
        cells_html = []
        for cell in row.cells:
            # В ячейке PPTX текст живёт внутри paragraph/run, поэтому собираем его вручную.
            cell_text = " ".join(
                ("".join((run.text for run in paragraph.runs)) for paragraph in cell.text_frame.paragraphs)
            ).strip()
            cells_html.append(f"  <td>{cell_text}</td>")
        rows_html.append("<tr>\n" + "\n".join(cells_html) + "\n</tr>")
    if not rows_html:
        return ""
    return "<table>\n" + "\n".join(rows_html) + "\n</table>"
# END_PPTX_HELPERS
