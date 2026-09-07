# START_MODULE_CONTRACT
# PURPOSE: Special-runner сквозной (кросс-документной) проверки 3 документов проекта через LLM.
#   R1 (комплектность 3 префиксов) — Python-гейт; сама сверка R2..N — один вызов Qwen по размеченному контексту 3 доков.
# INPUTS: argparse-Namespace с полем `target` = ПУТЬ ПАПКИ с 3 файлами (2.4 xlsx, 0.6, 0.5) (+ session_dir, rule_filter, parse_only).
# OUTPUTS: `AuditResult` (violations/rules_checked/session_dir/...); артефакты validation_outputs/validate_N.json + validation_report.xlsx + llm_prompt/response.
# KEYWORDS: crosscheck, multi-doc, llm, qwen, kartochka, protokol, tirazh, special-runner.
# LINKS: src/api/server.py (SPECIAL_ENGINE_RUNNERS), src/audit/models.py (AuditResult),
#   src/format_parsers (parse_pdf/parse_docx), config/parsers.py (PARSERS_CONFIG), config/llm.py (LLM_CONFIG), src/llm/client.py (make_llm_client).
# END_MODULE_CONTRACT

from __future__ import annotations

import json
import re
import time
import traceback
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

from src.audit.models import AuditResult
from src.format_parsers import parse_docx
from src.format_parsers.pdf import parse_pdf
from config.parsers import PARSERS_CONFIG
from config.llm import LLM_CONFIG
from src.llm.client import make_llm_client

# START_CONSTANTS
_PROJECT_ROOT = Path(__file__).resolve().parents[2]
_DOC_CONFIGS_DIR = _PROJECT_ROOT / "doc_configs"
_LOGS_RESULT_DIR = _PROJECT_ROOT / "logs_result"
_DOC_TYPE = "crosscheck_2_4_0_6_0_5"

_SYSTEM_PROMPT = """Ты — эксперт по перекрёстной (сквозной) проверке комплекта документов ОДНОГО проекта.
Тебе дают контекст ТРЁХ документов: «2.4 Карточка проекта» (Excel-лист как таблица), «0.6 Протокол выполнения» (текст), «0.5 Приказ о тираже» (текст).
Задача — проверить СОГЛАСОВАННОСТЬ данных МЕЖДУ этими документами строго по заданным правилам.

Показатели проекта (ключевые): ВПП (время протекания процесса — по смыслу СНИЖАЕТСЯ), Выработка (РАСТЁТ), Запасы/НЗП (СНИЖАЮТСЯ). У показателей бывают значения «До проекта», «Цель», «Факт».

ФОРМАТ ОТВЕТА — СТРОГО JSON-массив, по ОДНОМУ объекту на каждое правило, в том же порядке, что правила. У объекта ровно три поля:
- "rule_index": номер правила (как дан),
- "reasoning": краткое обоснование (1-3 предложения): что сверил между документами, что нашёл,
- "verdict": "ok" если правило соблюдено, "fail" если есть расхождение.
Никакого текста вне JSON-массива. Двойные кавычки внутри строк экранируй или заменяй на «ёлочки».

Если в документах НЕТ данных для проверки правила — verdict "ok" (не выдумывай нарушение), но отметь нехватку данных в reasoning."""
# END_CONSTANTS


# START_HELPERS
def _load_config() -> Dict[str, Any]:
    """Читает doc_configs/<_DOC_TYPE>/config.json (required_docs, llm_base_url?, model?) в dict."""
    return json.loads((_DOC_CONFIGS_DIR / _DOC_TYPE / "config.json").read_text(encoding="utf-8"))


def _load_rules() -> List[Dict[str, Any]]:
    """Читает doc_configs/<_DOC_TYPE>/rules.json и возвращает список из ключа "rules" (пустой список, если ключа нет)."""
    p = _DOC_CONFIGS_DIR / _DOC_TYPE / "rules.json"
    return json.loads(p.read_text(encoding="utf-8")).get("rules", [])


def _norm(s: Any) -> str:
    """Нормализует строку для сравнения: str(s or "") → lower → убрать кавычки («»"'`) → схлопнуть пробелы → strip."""
    t = str(s or "").lower()
    t = re.sub(r"[«»\"'`]", "", t)
    return re.sub(r"\s+", " ", t).strip()


def _classify_files(files: List[Path], required: List[Dict[str, Any]]) -> Dict[str, Optional[Path]]:
    """Сопоставляет роль (kartochka/protokol/tirazh) файлу по коду-префиксу ИЛИ ключевым словам имени."""
    roles: Dict[str, Optional[Path]] = {r["role"]: None for r in required}
    for f in files:
        name = _norm(f.name)
        for r in required:
            if roles[r["role"]] is not None:
                continue
            code_hit = bool(r.get("code")) and r["code"] in f.name
            kw_hit = all(_norm(k) in name for k in r.get("keywords", [])) if r.get("keywords") else False
            if code_hit or kw_hit:
                roles[r["role"]] = f
                break
    return roles


def _dump_xlsx_sheet(path: Path) -> str:
    """Весь целевой лист Карточки как текстовая таблица (без фикс-ячеек → устойчиво к раскладке)."""
    wb = openpyxl.load_workbook(str(path), data_only=True)
    ws = None
    for s in wb.sheetnames:
        if "карточка проекта" in s.lower() and "шаблон" not in s.lower():
            ws = wb[s]
            break
    if ws is None:
        ws = wb[wb.sheetnames[1]] if len(wb.sheetnames) > 1 else wb[wb.sheetnames[0]]
    lines: List[str] = []
    for row in ws.iter_rows(values_only=True):
        cells = [("" if c is None else str(c)).strip() for c in row]
        if any(cells):
            lines.append(" | ".join(cells).strip(" |"))
    return "Лист «%s»:\n%s" % (ws.title, "\n".join(lines))


def _parse_text(path: Path) -> str:
    """pdf/docx → raw_text (переиспользует форматные парсеры сервиса)."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        return parse_pdf(str(path), PARSERS_CONFIG.pdf).get("raw_text", "")
    if ext == ".docx":
        return parse_docx(str(path), vlm=None, header_ocr_cfg=None).get("raw_text", "")
    return ""


def _parse_verdicts(text: str) -> List[Dict[str, Any]]:
    """Достаёт JSON-массив вердиктов из ответа модели."""
    t = text.strip()
    if t.startswith("```"):
        t = re.sub(r"^```[a-z]*\n?", "", t).rstrip("`").strip()
    i, j = t.find("["), t.rfind("]")
    if i != -1 and j != -1 and j > i:
        try:
            data = json.loads(t[i:j + 1])
            if isinstance(data, list):
                return data
        except json.JSONDecodeError:
            pass
    # запасной разбор по объектам
    out = []
    for m in re.findall(r"\{[^{}]*\}", t, flags=re.S):
        try:
            out.append(json.loads(m))
        except json.JSONDecodeError:
            continue
    return out


def _write_report(results: List[Dict[str, Any]], path: Path) -> None:
    """
    Назначение:
        Пишет Excel-отчёт по результатам проверки: один лист «Отчёт», строка заголовка
        (rule_index, rule_title, status, Обоснование, duration_sec) и по строке на каждое правило.

    Вход:
        results: список dict-результатов из _emit ({rule_index, rule_title, status, discrepancy, duration_sec}).
        path: путь к сохраняемому .xlsx (validation_report.xlsx в session_dir).

    Выход:
        None; файл сохраняется на диск. В колонку «Обоснование» идёт discrepancy (по умолчанию ""),
        duration_sec округляется до 3 знаков (по умолчанию 0.0).
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.append(["rule_index", "rule_title", "status", "Обоснование", "duration_sec"])
    for r in results:
        ws.append([r["rule_index"], r["rule_title"], r["status"], r.get("discrepancy", ""), round(r.get("duration_sec", 0.0), 3)])
    wb.save(str(path))
# END_HELPERS


# START_ENTRY
def run_crosscheck_special(args: Any) -> AuditResult:
    """
    Назначение:
        Публичный entrypoint сквозной LLM-проверки. R1 (комплектность) — Python-гейт;
        сверка R2..N — один вызов Qwen по размеченному контексту 3 документов.

    Вход:
        args: SimpleNamespace/Namespace с {target=папка с 3 файлами, session_dir?, rule_filter?, parse_only?}.

    Выход:
        AuditResult.
    """
    start_time = time.time()
    session_dir = (
        Path(args.session_dir) if getattr(args, "session_dir", None)
        else _LOGS_RESULT_DIR / _DOC_TYPE / f"session_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
    )
    session_dir.mkdir(parents=True, exist_ok=True)
    val_dir = session_dir / "validation_outputs"
    val_dir.mkdir(parents=True, exist_ok=True)

    results: List[Dict[str, Any]] = []

    def _emit(idx: Any, title: str, status: str, disc: str) -> None:
        """
        Назначение:
            Фиксирует результат по одному правилу: добавляет dict в замыкание `results`
            и пишет validation_outputs/validate_<idx>.json (без duration_sec).

        Вход:
            idx: индекс правила (кладётся как str(idx) в rule_index и в имя файла).
            title: заголовок правила.
            status: "PASS" / "FAIL" / "ERROR".
            disc: текст обоснования/расхождения (для PASS может быть пустым).

        Выход:
            None; побочные эффекты — results.append и запись JSON-файла (duration_sec всегда 0.0).
        """
        res = {"rule_index": str(idx), "rule_title": title, "status": status, "discrepancy": disc, "duration_sec": 0.0}
        results.append(res)
        (val_dir / f"validate_{idx}.json").write_text(
            json.dumps({k: res[k] for k in ("rule_index", "rule_title", "status", "discrepancy")}, ensure_ascii=False, indent=2),
            encoding="utf-8")

    try:
        cfg = _load_config()
        rules = _load_rules()
        required = cfg.get("required_docs", [])
        folder = Path(args.target)
        files = [f for f in folder.iterdir() if f.is_file()] if folder.is_dir() else [folder]
        roles = _classify_files(files, required)

        # R1 — комплектность (Python-гейт).
        gate_rule = next((r for r in rules if r.get("index") == 1), {"index": 1, "title": "Комплектность и типы документов"})
        missing = [r["role"] for r in required if roles.get(r["role"]) is None]
        if missing:
            _emit(1, gate_rule["title"], "FAIL",
                  "неполный/неверный комплект: не распознаны роли " + ", ".join(missing) +
                  f" (получено файлов: {len(files)}; нужны 2.4 Карточка + 0.6 Протокол + 0.5 Приказ о тираже)")
            _write_report(results, session_dir / "validation_report.xlsx")
            return AuditResult(violations=[r for r in results if r["status"] == "FAIL"], doc_type=_DOC_TYPE,
                               session_dir=session_dir, duration_sec=time.time() - start_time,
                               rules_checked=len(results), target_path=str(folder))
        _emit(1, gate_rule["title"], "PASS", "")

        # Парсинг 3 документов в размеченный контекст.
        kart_dump = _dump_xlsx_sheet(roles["kartochka"])
        prot_text = _parse_text(roles["protokol"])
        tir_text = _parse_text(roles["tirazh"])
        context = (
            "=== ДОКУМЕНТ 2.4 — КАРТОЧКА ПРОЕКТА (Excel, целевой лист как таблица) ===\n" + kart_dump +
            "\n\n=== ДОКУМЕНТ 0.6 — ПРОТОКОЛ ВЫПОЛНЕНИЯ (полный текст) ===\n" + prot_text +
            "\n\n=== ДОКУМЕНТ 0.5 — ПРИКАЗ О ПЕРЕХОДЕ НА ТИРАЖИРОВАНИЕ (полный текст) ===\n" + tir_text
        )
        (session_dir / "context.txt").write_text(context, encoding="utf-8")

        if getattr(args, "parse_only", False):
            return AuditResult(doc_type=_DOC_TYPE, session_dir=session_dir,
                               duration_sec=time.time() - start_time, target_path=str(folder))

        # Правила для LLM (все, кроме R1-гейта).
        llm_rules = [r for r in rules if r.get("index") != 1]
        rules_block = "\n".join("Правило %s: %s — %s" % (r["index"], r["title"], r["check"]) for r in llm_rules)
        user_prompt = (
            "ПРАВИЛА КРОСС-ПРОВЕРКИ:\n" + rules_block +
            "\n\nКОНТЕКСТ ДОКУМЕНТОВ:\n" + context +
            "\n\nВерни JSON-массив вердиктов по всем правилам (по объекту на правило, в порядке правил)."
        )
        (session_dir / "llm_user_prompt.txt").write_text(user_prompt, encoding="utf-8")

        # Вызов Qwen.
        base_url = cfg.get("llm_base_url") or LLM_CONFIG.base_url
        model = cfg.get("model") or LLM_CONFIG.default_model
        client = make_llm_client(base_url)
        resp = client.chat.completions.create(
            model=model,
            messages=[{"role": "system", "content": _SYSTEM_PROMPT}, {"role": "user", "content": user_prompt}],
            temperature=0.0,
            max_tokens=4096,
            extra_body={"chat_template_kwargs": {"enable_thinking": False}},
        )
        raw = resp.choices[0].message.content or ""
        (session_dir / "llm_response.txt").write_text(raw, encoding="utf-8")
        verdicts = _parse_verdicts(raw)
        by_idx = {str(v.get("rule_index")): v for v in verdicts}

        # Раскладываем вердикты по правилам R2..N.
        for r in llm_rules:
            v = by_idx.get(str(r["index"]), {})
            verdict = str(v.get("verdict", "")).lower()
            reasoning = (v.get("reasoning") or "").strip()
            if verdict == "fail":
                _emit(r["index"], r["title"], "FAIL", reasoning or "расхождение (без обоснования модели)")
            elif verdict == "ok":
                _emit(r["index"], r["title"], "PASS", reasoning)
            else:
                _emit(r["index"], r["title"], "ERROR", "модель не вернула вердикт по правилу")

        _write_report(results, session_dir / "validation_report.xlsx")
        return AuditResult(violations=[r for r in results if r["status"] == "FAIL"], doc_type=_DOC_TYPE,
                           session_dir=session_dir, duration_sec=time.time() - start_time,
                           rules_checked=len(results), target_path=str(folder))
    except Exception as e:
        (session_dir / "ERROR.txt").write_text(f"{type(e).__name__}: {e}\n{traceback.format_exc()}", encoding="utf-8")
        raise
# END_ENTRY
