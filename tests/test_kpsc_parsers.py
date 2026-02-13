#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Юнит-тесты для kpsc парсеров.

Проверяем что sheet_finder и все 9 парсеров корректно работают
на всех 5 тестовых xlsx файлах.
"""

import json
import tempfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

# --- Путь к тестовым файлам ---
PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEST_FILES = {
    "biznes_otel": PROJECT_ROOT / "test_docs/kpsc/biznes_otel/КПСЦ_ТС_ООО_Бизнес_Отель_v3.xlsx",
    "mapper": PROJECT_ROOT / "test_docs/kpsc/mapper/КПСЦ Текущее состояние ООО МАППЕР версия 2.xlsx",
    "rotosnab": PROJECT_ROOT / "test_docs/kpsc/rotosnab/КПСЦ_и_Спагетти_ТС_ООО_Ротоснаб_итог.xlsx",
    "ruslet": PROJECT_ROOT / "test_docs/kpsc/ruslet/КПСЦ и Спагетти ТС_ООО РУСЛЕТ_v.04.1.xlsx",
    "sodex": PROJECT_ROOT / "test_docs/kpsc/sodex/КПСЦ Текущее состояние ООО СОДЕКС версия 4 (1).xlsx",
}


# --- sheet_finder ---

class TestSheetFinder:
    """Тесты для find_sheet — нечёткого поиска листов."""

    def _load_wb(self, company):
        return load_workbook(str(TEST_FILES[company]), data_only=True)

    @pytest.mark.parametrize("company", list(TEST_FILES.keys()))
    def test_find_kpsc_sheet(self, company):
        """КПСЦ-лист должен находиться у всех 5 компаний."""
        from audit_engine.kpsc.sheet_finder import find_sheet
        wb = self._load_wb(company)
        ws = find_sheet(wb, keywords=["кпсц"],
                        exclude_keywords=["спагетти", "укрупн", "оцифровк"],
                        prefer_keywords=["тс", "текущ"])
        assert ws is not None, f"{company}: КПСЦ-лист не найден среди {wb.sheetnames}"
        wb.close()

    @pytest.mark.parametrize("company", list(TEST_FILES.keys()))
    def test_find_pa1_sheet(self, company):
        """ПА-1 лист должен находиться у всех 5 компаний."""
        from audit_engine.kpsc.sheet_finder import find_sheet
        wb = self._load_wb(company)
        ws = find_sheet(wb, keywords=["па"],
                        exclude_keywords=["спагетти", "кпсц", "ямадз"])
        assert ws is not None, f"{company}: ПА-1 лист не найден среди {wb.sheetnames}"
        wb.close()

    @pytest.mark.parametrize("company", list(TEST_FILES.keys()))
    def test_find_spaghetti_sheet(self, company):
        """Диаграмма Спагетти должна находиться у всех компаний."""
        from audit_engine.kpsc.sheet_finder import find_sheet
        wb = self._load_wb(company)
        ws = find_sheet(wb, keywords=["спагетти"],
                        exclude_keywords=["пробл", "улучш", "перечень"])
        assert ws is not None, f"{company}: Диаграмма Спагетти не найдена среди {wb.sheetnames}"
        wb.close()

    def test_latin_to_cyrillic_normalization(self):
        """Лист 'Диаграмма Cпагетти' (Latin C) должен находиться."""
        from audit_engine.kpsc.sheet_finder import find_sheet
        # biznes_otel имеет 'Диаграмма Cпагетти' с Latin C
        wb = self._load_wb("biznes_otel")
        ws = find_sheet(wb, keywords=["спагетти"],
                        exclude_keywords=["пробл", "улучш", "перечень"])
        assert ws is not None
        assert "пагетти" in ws.title.lower() or "Cпагетти" in ws.title
        wb.close()

    def test_trailing_space_handling(self):
        """Лист 'ПА1 ' (trailing space) должен находиться."""
        from audit_engine.kpsc.sheet_finder import find_sheet
        wb = self._load_wb("biznes_otel")
        ws = find_sheet(wb, keywords=["па"],
                        exclude_keywords=["спагетти", "кпсц", "ямадз"])
        assert ws is not None
        assert "ПА1" in ws.title
        wb.close()


# --- parse_kpsc_header ---

class TestParseKpscHeader:
    """Тесты для парсера заголовков КПСЦ."""

    @pytest.mark.parametrize("company", list(TEST_FILES.keys()))
    def test_header_no_crash(self, company):
        """Парсер не должен падать ни на одном файле."""
        from audit_engine.kpsc.parser_scripts.parse_kpsc_header import build_payload
        payload = build_payload(TEST_FILES[company])
        assert "meta" in payload
        assert "fields" in payload

    def test_biznes_otel_fields(self):
        """biznes_otel: ключевые поля header должны быть заполнены."""
        from audit_engine.kpsc.parser_scripts.parse_kpsc_header import build_payload
        payload = build_payload(TEST_FILES["biznes_otel"])
        fields = payload["fields"]
        assert fields["title"] is not None, "title = None"
        assert fields["responsible"] is not None, "responsible = None"
        assert fields["date_developed"] is not None, "date_developed = None"
        assert fields["compiled_by"] is not None, "compiled_by = None"

    def test_rotosnab_uses_correct_sheet(self):
        """rotosnab: должен использовать 'КПСЦ ТС' (с реальными данными), а не 'КПСЦ'."""
        from audit_engine.kpsc.parser_scripts.parse_kpsc_header import build_payload
        payload = build_payload(TEST_FILES["rotosnab"])
        assert payload["meta"]["sheet"] == "КПСЦ ТС"
        assert payload["fields"]["responsible"] is not None

    def test_ruslet_fields(self):
        """ruslet: поля должны корректно парситься (D-столбец)."""
        from audit_engine.kpsc.parser_scripts.parse_kpsc_header import build_payload
        payload = build_payload(TEST_FILES["ruslet"])
        fields = payload["fields"]
        assert fields["responsible"] is not None, "responsible = None (ожидалось из D3)"
        assert fields["compiled_by"] is not None, "compiled_by = None"


# --- Полный запуск всех парсеров ---

class TestAllParsersRun:
    """Тесты полного запуска всех 9 парсеров на каждой компании."""

    @pytest.mark.parametrize("company", list(TEST_FILES.keys()))
    def test_all_parsers_no_crash(self, company):
        """Все 9 парсеров должны выполниться без ошибок."""
        from audit_engine.kpsc.run_validations import run_parsers
        with tempfile.TemporaryDirectory() as tmpdir:
            results = run_parsers(TEST_FILES[company], Path(tmpdir))
            errors = {name: r["error"] for name, r in results.items()
                      if r["status"] == "error"}
            assert len(errors) == 0, f"{company}: парсеры с ошибками: {errors}"

    @pytest.mark.parametrize("company", list(TEST_FILES.keys()))
    def test_parser_outputs_created(self, company):
        """Все 9 парсеров должны создать JSON-файлы."""
        from audit_engine.kpsc.run_validations import run_parsers
        with tempfile.TemporaryDirectory() as tmpdir:
            run_parsers(TEST_FILES[company], Path(tmpdir))
            json_files = list(Path(tmpdir).glob("*.json"))
            assert len(json_files) == 9, \
                f"{company}: ожидалось 9 JSON файлов, создано {len(json_files)}: {[f.name for f in json_files]}"
