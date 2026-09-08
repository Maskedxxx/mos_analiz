#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты F15: неполный ответ модели не выдаётся за полный отчёт.

- run_multi_rule_audit возвращает `unchecked` для правил без вердикта и `checked_count`;
- AuditEngine.run помечает непроверенные правила, rules_checked = фактически проверенные;
- если проверено 0 правил — ValueError;
- save_to_excel рисует «НЕ ПРОВЕРЕНО», а не «ОК».
Модель подменяется: возвращает вердикт только по первому правилу (стаб `missing`).
"""
import tempfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook


def test_save_to_excel_marks_unchecked_not_ok():
    from src.audit.excel_reporter import save_to_excel
    rules = [
        {"index": 1, "title": "Правило 1", "layer": "base"},
        {"index": 2, "title": "Правило 2", "layer": "base"},
    ]
    with tempfile.TemporaryDirectory() as tmp:
        out = str(Path(tmp) / "r.xlsx")
        # правило 2 не проверено моделью
        save_to_excel([], out, multi_rules=rules, unchecked=[{"index": 2, "layer": "base"}])
        ws = load_workbook(out).active
        statuses = {row[0]: row[3] for row in ws.iter_rows(min_row=2, values_only=True)}
        assert statuses[1] == "ОК"
        assert statuses[2] == "НЕ ПРОВЕРЕНО"


class _StubMessage:
    def __init__(self, content):
        self.content = content


class _StubChoice:
    def __init__(self, content):
        self.message = _StubMessage(content)


class _StubUsage:
    prompt_tokens = 1
    completion_tokens = 1
    total_tokens = 2


class _StubResponse:
    def __init__(self, content):
        self.choices = [_StubChoice(content)]
        self.usage = _StubUsage()


class _StubCompletions:
    def __init__(self, content):
        self._content = content

    def create(self, **kwargs):
        return _StubResponse(self._content)


class _StubChat:
    def __init__(self, content):
        self.completions = _StubCompletions(content)


class _StubClient:
    """Отдаёт вердикт только по ПЕРВОМУ правилу — имитация неполного ответа модели."""

    def __init__(self, content):
        self.chat = _StubChat(content)


def _install_stub(monkeypatch, first_index):
    import json as _json
    import src.llm.multi_rule as mr
    content = _json.dumps(
        [{"rule_index": first_index, "reasoning": "stub", "verdict": {"status": "ok"}}],
        ensure_ascii=False,
    )
    monkeypatch.setattr(mr, "make_llm_client", lambda base_url=None: _StubClient(content))


def test_run_multi_rule_reports_unchecked(monkeypatch):
    _install_stub(monkeypatch, first_index=1)
    from src.llm.multi_rule import run_multi_rule_audit
    rules = [
        {"index": 1, "title": "Правило 1", "check": "c", "target_sections": []},
        {"index": 2, "title": "Правило 2", "check": "c", "target_sections": []},
        {"index": 3, "title": "Правило 3", "check": "c", "target_sections": []},
    ]
    res = run_multi_rule_audit(
        parsed={"filename": "d.docx", "path": "d.docx", "raw_text": "текст"},
        sections={}, rules=rules, include_scopes=None, filename="d.docx",
        llm_base_url="http://stub/v1/", llm_model="stub",
    )
    assert res["checked_count"] == 1
    assert [r["index"] for r in res["unchecked"]] == [2, 3]


def test_engine_marks_unchecked_and_counts(monkeypatch):
    # akt_nachala: базовые правила имеют индексы 2..9 — берём реальный первый (2),
    # чтобы одно правило проверилось, а остальные попали в unchecked.
    _install_stub(monkeypatch, first_index=2)
    import src.audit.engine as engine_mod
    with tempfile.TemporaryDirectory() as tmp:
        doc = Document()
        doc.add_paragraph("АКТ о начале реализации мероприятий г. Москва 2026")
        target = Path(tmp) / "akt.docx"
        doc.save(str(target))
        session = Path(tmp) / "session"
        result = engine_mod.AuditEngine("akt_nachala").run(str(target), session_dir=str(session))
        # хотя бы одно правило проверено, часть — нет; rules_checked = проверенные, unchecked непуст
        assert result.rules_checked >= 1
        assert len(result.unchecked_rules) >= 1
        assert result.rules_checked < result.rules_checked + len(result.unchecked_rules)
        wb = load_workbook(str(session / "audit_result.xlsx")).active
        statuses = [row[3] for row in wb.iter_rows(min_row=2, values_only=True)]
        assert "НЕ ПРОВЕРЕНО" in statuses


def test_engine_zero_checked_raises(monkeypatch):
    # Модель отдаёт вердикт по несуществующему индексу → проверено 0 правил → ошибка.
    _install_stub(monkeypatch, first_index=9999)
    import src.audit.engine as engine_mod
    with tempfile.TemporaryDirectory() as tmp:
        doc = Document()
        doc.add_paragraph("АКТ о начале реализации мероприятий г. Москва 2026")
        target = Path(tmp) / "akt.docx"
        doc.save(str(target))
        with pytest.raises(ValueError, match="ни одного вердикта"):
            engine_mod.AuditEngine("akt_nachala").run(str(target), session_dir=str(Path(tmp) / "s"))
