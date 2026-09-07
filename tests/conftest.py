#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Общие фикстуры для тестов аудит-движка.

Содержит пути к тестовым файлам и вспомогательные функции.
"""
import os
import sys
from pathlib import Path

import fitz
import pytest
from docx import Document

# Корень проекта
PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

# Учётные данные веб-API для тестов. В коде сервера дефолтов нет (см. src/api/server.py START_AUTH),
# поэтому задаём тестовые значения до первого импорта main. Реальные значения из окружения не перекрываются.
os.environ.setdefault("AUDIT_LOGIN", "test")
os.environ.setdefault("AUDIT_PASSWORD", "test-password")
os.environ.setdefault("AUDIT_TOKEN", "test-token")

TEST_DOCS = PROJECT_ROOT / "test_docs"
DOC_CONFIGS = PROJECT_ROOT / "doc_configs"
GENERATED_FIXTURES_DIR = PROJECT_ROOT / ".pytest_cache" / "generated_fixtures"


# === Пути к тестовым файлам kpsc ===
KPSC_FILES = {
    "biznes_otel": TEST_DOCS / "kpsc" / "biznes_otel" / "КПСЦ_ТС_ООО_Бизнес_Отель_v3.xlsx",
    "mapper": TEST_DOCS / "kpsc" / "mapper" / "КПСЦ Текущее состояние ООО МАППЕР версия 2.xlsx",
    "rotosnab": TEST_DOCS / "kpsc" / "rotosnab" / "КПСЦ_и_Спагетти_ТС_ООО_Ротоснаб_итог.xlsx",
    "ruslet": TEST_DOCS / "kpsc" / "ruslet" / "КПСЦ и Спагетти ТС_ООО РУСЛЕТ_v.04.1.xlsx",
    "sodex": TEST_DOCS / "kpsc" / "sodex" / "КПСЦ Текущее состояние ООО СОДЕКС версия 4 (1).xlsx",
}


# === Пути к тестовым файлам prikaz_comp_ppu ===
PRIKAZ_COMP_PPU_FILES = {
    "biznes_otel": TEST_DOCS / "prikaz_comp_ppu" / "biznes_otel" / "3_ПРИКАЗ_о_конкурсах_ППУ_ООО_Бизнес_отель.docx",
    "mapper": TEST_DOCS / "prikaz_comp_ppu" / "mapper" / "ПРИКАЗ о конкурсах ППУ Маппер.docx",
    "rotosnab": TEST_DOCS / "prikaz_comp_ppu" / "rotosnab" / "ПРИКАЗ о конкурсах ППУ.docx",
    "ruslet": TEST_DOCS / "prikaz_comp_ppu" / "ruslet" / "Приказ о проведении конкурсов ППУ_от 14.10.2025.docx",
    "sodex": TEST_DOCS / "prikaz_comp_ppu" / "sodex" / "ПРИКАЗ о конкурсах ППУ ООО Содекс.docx",
}

LOCAL_SMOKE_FIXTURES = {
    "akt_nachala": DOC_CONFIGS / "akt_nachala" / "template" / "template.docx",
    "cheklist_eu": DOC_CONFIGS / "cheklist_eu" / "template" / "template.docx",
    "polozhenie_comp_ppu": DOC_CONFIGS / "polozhenie_comp_ppu" / "template" / "template.docx",
    "polozhenie_po": DOC_CONFIGS / "polozhenie_po" / "template" / "template.docx",
    "polozhenie_ppu": DOC_CONFIGS / "polozhenie_ppu" / "template" / "template.docx",
    "presentation_eu": DOC_CONFIGS / "presentation_eu" / "template" / "template.pptx",
    "prikaz_comp_ppu": DOC_CONFIGS / "prikaz_comp_ppu" / "template" / "template.docx",
    "prikaz_formirovanie_po": DOC_CONFIGS / "prikaz_formirovanie_po" / "template" / "template.docx",
    "prikaz_ic": DOC_CONFIGS / "prikaz_ic" / "template" / "template.docx",
    "prikaz_ic_el": DOC_CONFIGS / "prikaz_ic_el" / "template" / "template.docx",
    "prikaz_ic_potoka": DOC_CONFIGS / "prikaz_ic_potoka" / "template" / "template.docx",
    "prikaz_ic_potoka_el": DOC_CONFIGS / "prikaz_ic_potoka_el" / "template" / "template.docx",
    "prikaz_ppu": DOC_CONFIGS / "prikaz_ppu" / "template" / "template.docx",
    "prikaz_vyhod": DOC_CONFIGS / "prikaz_vyhod" / "template" / "template.docx",
}


def _docx_to_pdf(src: Path, dest: Path) -> Path:
    """Build a simple local PDF fixture from a DOCX template for smoke tests."""
    doc = Document(src)
    lines = []
    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        if text:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))

    pdf = fitz.open()
    page = pdf.new_page()
    rect = page.rect
    y = 36
    for line in lines[:500]:
        chunks = [line[i:i + 90] for i in range(0, len(line), 90)] or [""]
        for chunk in chunks:
            if y > rect.height - 36:
                page = pdf.new_page()
                y = 36
            page.insert_text((36, y), chunk, fontsize=10)
            y += 12
    dest.parent.mkdir(parents=True, exist_ok=True)
    pdf.save(dest)
    pdf.close()
    return dest


@pytest.fixture
def kpsc_files():
    """Словарь путей к тестовым XLSX файлам kpsc."""
    return KPSC_FILES


@pytest.fixture
def prikaz_comp_ppu_files():
    """Словарь путей к тестовым DOCX файлам prikaz_comp_ppu."""
    return PRIKAZ_COMP_PPU_FILES


@pytest.fixture(scope="session")
def generated_pdf_fixtures():
    GENERATED_FIXTURES_DIR.mkdir(parents=True, exist_ok=True)
    source = DOC_CONFIGS / "akt_nachala" / "template" / "template.docx"
    pdf_path = _docx_to_pdf(source, GENERATED_FIXTURES_DIR / "akt_nachala_template.pdf")
    return {
        "prikaz_pa": pdf_path,
        "otchet_rezultatov": pdf_path,
    }
