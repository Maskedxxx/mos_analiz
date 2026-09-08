#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты guard'а пустого/нераспознанного документа (аудит устойчивости, фикс F1).

Проверяют, что:
- `check_document_text` отличает текст документа от маркеров границ и ошибок OCR;
- `AuditEngine.run` на пустом docx/pptx падает с ValueError ДО вызова модели;
- `parse_pdf` при сбое OCR на всех страницах бросает ValueError, при сбое на части —
  отдаёт `warnings`, а маркер ошибки в `raw_text` не попадает;
- `save_to_excel` при непустых `warnings` добавляет лист «Предупреждения».
Модель, OCR и layout не нужны: PDF-экстрактор подменяется monkeypatch.
"""
import tempfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from conftest import PROJECT_ROOT


def test_check_document_text_accepts_real_text():
    from src.audit.engine import check_document_text
    check_document_text("[СТРАНИЦА 1]\nАКТ о начале реализации мероприятий")


@pytest.mark.parametrize("raw_text", ["", "   \n\n", "[СТРАНИЦА 1]\n\n[СТРАНИЦА 2]", "[СЛАЙД 1]"])
def test_check_document_text_rejects_empty_or_markers_only(raw_text):
    from src.audit.engine import check_document_text
    with pytest.raises(ValueError, match="не содержит распознаваемого текста"):
        check_document_text(raw_text)


def test_check_document_text_reports_first_ocr_error():
    from src.audit.engine import check_document_text
    raw = "[СТРАНИЦА 1]\n[ОШИБКА Paddle OCR: Layout Detection API недоступен: http://127.0.0.1:2]"
    with pytest.raises(ValueError, match="не удалось распознать.*Layout Detection API недоступен"):
        check_document_text(raw)


def test_engine_rejects_blank_docx_before_llm(monkeypatch):
    """Пустой docx: ValueError из engine.run, run_multi_rule_audit не вызывается."""
    import src.audit.engine as engine_mod

    def _must_not_be_called(*args, **kwargs):
        raise AssertionError("модель вызвана для пустого документа")

    monkeypatch.setattr(engine_mod, "run_multi_rule_audit", _must_not_be_called)
    with tempfile.TemporaryDirectory() as tmpdir:
        blank = Path(tmpdir) / "blank.docx"
        Document().save(str(blank))
        engine = engine_mod.AuditEngine("akt_nachala")
        with pytest.raises(ValueError, match="не содержит распознаваемого текста"):
            engine.run(str(blank), session_dir=str(Path(tmpdir) / "session"))


def test_engine_rejects_image_only_pptx_before_llm(monkeypatch):
    """pptx из одной картинки: ValueError из engine.run, модель не вызывается."""
    import io

    from PIL import Image
    from pptx.util import Inches

    import src.audit.engine as engine_mod

    def _must_not_be_called(*args, **kwargs):
        raise AssertionError("модель вызвана для пустого документа")

    monkeypatch.setattr(engine_mod, "run_multi_rule_audit", _must_not_be_called)
    with tempfile.TemporaryDirectory() as tmpdir:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        buf = io.BytesIO()
        Image.new("RGB", (200, 100), "white").save(buf, "PNG")
        buf.seek(0)
        slide.shapes.add_picture(buf, Inches(1), Inches(1))
        pptx_path = Path(tmpdir) / "image_only.pptx"
        prs.save(str(pptx_path))
        engine = engine_mod.AuditEngine("presentation_eu")
        with pytest.raises(ValueError, match="не содержит распознаваемого текста"):
            engine.run(str(pptx_path), session_dir=str(Path(tmpdir) / "session"))


def _fake_pdf(tmpdir: str) -> str:
    """Файл-заглушка: parse_pdf проверяет только существование, страницы читает подменённый экстрактор."""
    p = Path(tmpdir) / "scan.pdf"
    p.write_bytes(b"%PDF-1.4 fake")
    return str(p)


def test_parse_pdf_all_pages_failed_raises(monkeypatch):
    from config.parsers import PARSERS_CONFIG
    import src.format_parsers.pdf.parse as parse_mod

    monkeypatch.setattr(parse_mod.PaddleExtractor, "detect_blank_pages", lambda self, path: (2, set()))
    monkeypatch.setattr(
        parse_mod.PaddleExtractor, "extract_pages",
        lambda self, path, idx: {1: "[ОШИБКА Paddle OCR: Layout Detection API недоступен: http://127.0.0.1:2]",
                                 2: "[ОШИБКА Paddle OCR: Layout Detection API недоступен: http://127.0.0.1:2]"},
    )
    with tempfile.TemporaryDirectory() as tmpdir:
        with pytest.raises(ValueError, match="ни одной страницы PDF \\(2 из 2\\).*Layout Detection API недоступен"):
            parse_mod.parse_pdf(_fake_pdf(tmpdir), PARSERS_CONFIG.pdf)


def test_parse_pdf_partial_failure_gives_warnings(monkeypatch):
    from config.parsers import PARSERS_CONFIG
    import src.format_parsers.pdf.parse as parse_mod

    monkeypatch.setattr(parse_mod.PaddleExtractor, "detect_blank_pages", lambda self, path: (2, set()))
    monkeypatch.setattr(
        parse_mod.PaddleExtractor, "extract_pages",
        lambda self, path, idx: {1: "АКТ о начале реализации", 2: "[ОШИБКА Paddle OCR: timeout]"},
    )
    with tempfile.TemporaryDirectory() as tmpdir:
        parsed = parse_mod.parse_pdf(_fake_pdf(tmpdir), PARSERS_CONFIG.pdf)
    assert "[ОШИБКА" not in parsed["raw_text"]
    assert "АКТ о начале реализации" in parsed["raw_text"]
    assert parsed["warnings"] == ["Страница 2 не распознана и в проверку не вошла: timeout"]


def test_parse_pdf_success_has_no_warnings_key(monkeypatch):
    from config.parsers import PARSERS_CONFIG
    import src.format_parsers.pdf.parse as parse_mod

    monkeypatch.setattr(parse_mod.PaddleExtractor, "detect_blank_pages", lambda self, path: (1, set()))
    monkeypatch.setattr(parse_mod.PaddleExtractor, "extract_pages", lambda self, path, idx: {1: "Текст"})
    with tempfile.TemporaryDirectory() as tmpdir:
        parsed = parse_mod.parse_pdf(_fake_pdf(tmpdir), PARSERS_CONFIG.pdf)
    assert "warnings" not in parsed
    assert parsed["raw_text"] == "[СТРАНИЦА 1]\nТекст"


def test_save_to_excel_warnings_sheet():
    from src.audit.excel_reporter import save_to_excel
    rules = [{"index": 1, "title": "Правило 1", "layer": "base"}]
    with tempfile.TemporaryDirectory() as tmpdir:
        out = str(Path(tmpdir) / "r.xlsx")
        save_to_excel([], out, multi_rules=rules, warnings=["Страница 2 не распознана и в проверку не вошла: timeout"])
        wb = load_workbook(out)
        assert wb.sheetnames == ["Результаты проверки", "Предупреждения"]
        assert wb["Предупреждения"]["A2"].value.startswith("Страница 2 не распознана")
        save_to_excel([], out, multi_rules=rules)
        assert load_workbook(out).sheetnames == ["Результаты проверки"]
