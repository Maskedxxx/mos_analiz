#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Тесты F17: обезвреживание Excel-ячеек (формулы → текст, усечение с пометкой).
"""
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from src.audit.excel_reporter import save_to_excel, _EXCEL_CELL_LIMIT, _TRUNC_NOTE


def _cells(path):
    ws = load_workbook(path).active
    return {c.coordinate: c for row in ws.iter_rows(min_row=2) for c in row}


def test_formula_strings_written_as_text():
    rules = [{"index": 1, "title": "П1", "layer": "base"}]
    viol = [{
        "rule_index": 1, "layer": "base", "правило": "П1",
        "Целевой документ": "=1+1", "Различие": '=HYPERLINK("http://x","c")', "Обоснование": "@at",
    }]
    with tempfile.TemporaryDirectory() as tmp:
        out = str(Path(tmp) / "r.xlsx")
        save_to_excel(viol, out, multi_rules=rules)
        cells = _cells(out)
        # Целевой документ (кол.6/F), Различие (7/G), Обоснование (8/H)
        assert cells["F2"].value == "=1+1" and cells["F2"].data_type == "s"
        assert cells["G2"].data_type == "s" and cells["G2"].value.startswith("=HYPERLINK")
        assert cells["H2"].value == "@at" and cells["H2"].data_type == "s"


def test_leading_plus_minus_are_text():
    rules = [{"index": 1, "title": "П1", "layer": "base"}]
    viol = [{"rule_index": 1, "layer": "base", "правило": "П1",
             "Целевой документ": "-1+1", "Различие": "+плюс", "Обоснование": ""}]
    with tempfile.TemporaryDirectory() as tmp:
        out = str(Path(tmp) / "r.xlsx")
        save_to_excel(viol, out, multi_rules=rules)
        cells = _cells(out)
        assert cells["F2"].data_type == "s" and cells["G2"].data_type == "s"


def test_long_text_truncated_with_note():
    rules = [{"index": 1, "title": "П1", "layer": "base"}]
    viol = [{"rule_index": 1, "layer": "base", "правило": "П1",
             "Целевой документ": "x" * 40000, "Различие": "", "Обоснование": ""}]
    with tempfile.TemporaryDirectory() as tmp:
        out = str(Path(tmp) / "r.xlsx")
        save_to_excel(viol, out, multi_rules=rules)
        v = _cells(out)["F2"].value
        assert len(v) <= _EXCEL_CELL_LIMIT + len(_TRUNC_NOTE)
        assert v.endswith(_TRUNC_NOTE)


def test_plain_text_unchanged():
    rules = [{"index": 1, "title": "П1", "layer": "base"}]
    viol = [{"rule_index": 1, "layer": "base", "правило": "П1",
             "Целевой документ": "обычный текст 🚨 «кавычки»", "Различие": "", "Обоснование": ""}]
    with tempfile.TemporaryDirectory() as tmp:
        out = str(Path(tmp) / "r.xlsx")
        save_to_excel(viol, out, multi_rules=rules)
        assert _cells(out)["F2"].value == "обычный текст 🚨 «кавычки»"
