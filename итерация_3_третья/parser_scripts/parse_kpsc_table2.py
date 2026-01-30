#!/usr/bin/env python3
"""
Парсер таблицы "2. Определение времени цикла изготовления деталей (ТЦ)" на листе КПСЦ.
Определяет границы по заголовку "2.", собирает прямоугольник от первой непустой
колонки заголовка до колонки X (включая X даже если пустая), учитывает merged-ячейки.

Использование:
  python parse_kpsc_table2.py -i path/to/file.xlsx [-s КПСЦ] [-o out.json]
"""

import argparse
import json
from pathlib import Path
from typing import Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


SECTION_PHRASE = "определение времени цикла изготовления деталей"
FORCED_RIGHT_COL = 24  # column X


def find_section_anchor(ws, phrase: str) -> Optional[Tuple[int, int]]:
    phrase_low = phrase.lower()
    for row in ws.iter_rows():
        for cell in row:
            val = cell.value
            if isinstance(val, str) and phrase_low in val.lower():
                if val.strip().startswith("2"):
                    return cell.row, cell.column
    return None


def find_header_row(ws, anchor_row: int) -> int:
    for r in range(anchor_row + 1, anchor_row + 6):
        if any(c.value not in (None, "") for c in ws[r]):
            return r
    return anchor_row + 1


def find_bottom_row(ws, header_row: int, section_col: int) -> int:
    r = header_row + 1
    while r <= ws.max_row:
        cell_val = ws.cell(row=r, column=section_col).value
        if isinstance(cell_val, str) and cell_val.strip().startswith("3."):
            return r - 1
        row_vals = [c.value for c in ws[r]]
        if all(v in (None, "") for v in row_vals) and r > header_row + 2:
            return r - 1
        r += 1
    return ws.max_row


def compute_col_bounds(ws, header_row: int) -> Tuple[int, int]:
    left = None
    for c in ws[header_row]:
        if c.value not in (None, ""):
            left = c.column
            break
    if left is None:
        left = 1
    right = max(left, FORCED_RIGHT_COL)
    return left, right


def build_merged_lookup(ws):
    lookup = {}
    for merge in ws.merged_cells.ranges:
        coord = merge.coord
        for r in range(merge.min_row, merge.max_row + 1):
            for c in range(merge.min_col, merge.max_col + 1):
                lookup[(r, c)] = coord
    return lookup


def extract_table(ws, top_row: int, bottom_row: int, left_col: int, right_col: int):
    merged_lookup = build_merged_lookup(ws)
    rows = []
    for r in range(top_row, bottom_row + 1):
        row_cells = []
        for c in range(left_col, right_col + 1):
            cell = ws.cell(row=r, column=c)
            coord = f"{get_column_letter(c)}{r}"
            merge_range = merged_lookup.get((r, c))
            if merge_range:
                min_col_m, min_row_m, max_col_m, max_row_m = range_boundaries(merge_range)
                anchor = (r == min_row_m and c == min_col_m)
                value = ws.cell(row=min_row_m, column=min_col_m).value
            else:
                anchor = True
                value = cell.value
            row_cells.append(
                {
                    "coord": coord,
                    "col": c,
                    "row": r,
                    "value": value,
                    "merge_range": merge_range,
                    "merge_anchor": anchor if merge_range else False,
                }
            )
        rows.append({"row": r, "cells": row_cells})
    return rows


def build_payload(xlsx_path: Path, sheet: str):
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet]

    anchor = find_section_anchor(ws, SECTION_PHRASE)
    if not anchor:
        raise SystemExit("Не найден заголовок '2. Определение времени цикла изготовления деталей (ТЦ)'")
    anchor_row, anchor_col = anchor

    header_row = find_header_row(ws, anchor_row)
    bottom_row = find_bottom_row(ws, header_row, anchor_col)
    left_col, right_col = compute_col_bounds(ws, header_row)

    table_rows = extract_table(ws, header_row, bottom_row, left_col, right_col)

    payload = {
        "meta": {
            "workbook": str(xlsx_path),
            "sheet": sheet,
            "section_title_cell": f"{get_column_letter(anchor_col)}{anchor_row}",
        },
        "bounds": {
            "top_row": header_row,
            "bottom_row": bottom_row,
            "left_col": left_col,
            "right_col": right_col,
            "left_letter": get_column_letter(left_col),
            "right_letter": get_column_letter(right_col),
            "height": bottom_row - header_row + 1,
            "width": right_col - left_col + 1,
        },
        "rows": table_rows,
    }
    return payload


def main():
    parser = argparse.ArgumentParser(description="Парсер таблицы '2. Определение времени цикла изготовления деталей (ТЦ)'")
    parser.add_argument("-i", "--input", required=True, help="XLSX файл")
    parser.add_argument("-s", "--sheet", default="КПСЦ", help="Лист (default: КПСЦ)")
    parser.add_argument("-o", "--output", default=None, help="JSON файл вывода (stdout если не указан)")
    args = parser.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
