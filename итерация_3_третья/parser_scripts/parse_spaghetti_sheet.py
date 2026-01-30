#!/usr/bin/env python3
"""
Парсер листа "Диаграмма Спагетти" (без самой диаграммы).
Снимает:
  - верхние текстовые/служебные строки до таблицы;
  - таблицу путей перемещений начиная со строки с заголовком ("Путь", ...),
    с учётом merged-ячееек.
Диаграмма-графика игнорируется.
"""

import argparse
import json
from pathlib import Path
from typing import Optional, Tuple, Dict, Any, List

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


HEADER_PHRASE = "путь"


def find_header_row(ws) -> Optional[int]:
    for r in range(1, ws.max_row + 1):
        row_vals = [c.value for c in ws[r]]
        if any(isinstance(v, str) and HEADER_PHRASE in v.lower() for v in row_vals):
            return r
    return None


def compute_col_bounds(ws, header_row: int) -> Tuple[int, int]:
    left = None
    right = 0
    # левая граница — первый непустой столбец заголовка
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v not in (None, ""):
            left = c
            break
    if left is None:
        left = 1
    # правая граница — максимальный столбец с данными начиная с строки заголовка
    for r in range(header_row, ws.max_row + 1):
        for c in range(left, ws.max_column + 1):
            if ws.cell(row=r, column=c).value not in (None, ""):
                right = max(right, c)
    if right < left:
        right = left
    return left, right


def find_bottom_row(ws, header_row: int, left_col: int, right_col: int) -> int:
    r = header_row + 1
    while r <= ws.max_row:
        if all(ws.cell(row=r, column=c).value in (None, "") for c in range(left_col, right_col + 1)):
            return r - 1
        r += 1
    return ws.max_row


def build_merged_lookup(ws):
    lookup = {}
    for m in ws.merged_cells.ranges:
        coord = m.coord
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                lookup[(r, c)] = coord
    return lookup


def extract_table(ws, top_row: int, bottom_row: int, left_col: int, right_col: int):
    merged_lookup = build_merged_lookup(ws)
    rows = []
    for r in range(top_row, bottom_row + 1):
        row_cells = []
        for c in range(left_col, right_col + 1):
            coord = f"{get_column_letter(c)}{r}"
            merge_range = merged_lookup.get((r, c))
            if merge_range:
                min_col_m, min_row_m, max_col_m, max_row_m = range_boundaries(merge_range)
                anchor = (r == min_row_m and c == min_col_m)
                value = ws.cell(row=min_row_m, column=min_col_m).value
            else:
                anchor = True
                value = ws.cell(row=r, column=c).value
            row_cells.append(
                {
                    "coord": coord,
                    "row": r,
                    "col": c,
                    "value": value,
                    "merge_range": merge_range,
                    "merge_anchor": anchor if merge_range else False,
                }
            )
        rows.append({"row": r, "cells": row_cells})
    return rows


def extract_pre_table(ws, header_row: int):
    cells = []
    for r in range(1, header_row):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                cells.append(
                    {
                        "coord": f"{get_column_letter(c)}{r}",
                        "row": r,
                        "col": c,
                        "value": v,
                    }
                )
    return cells


def build_payload(xlsx_path: Path, sheet: str):
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet]

    header_row = find_header_row(ws)
    if not header_row:
        raise SystemExit("Не найден заголовок таблицы (\"Путь\") на листе 'Диаграмма Спагетти'")

    left_col, right_col = compute_col_bounds(ws, header_row)
    bottom_row = find_bottom_row(ws, header_row, left_col, right_col)

    pre_table = extract_pre_table(ws, header_row)
    table_rows = extract_table(ws, header_row, bottom_row, left_col, right_col)

    return {
        "meta": {"workbook": str(xlsx_path), "sheet": sheet, "header_row": header_row},
        "bounds": {
            "top_row": header_row,
            "bottom_row": bottom_row,
            "left_col": left_col,
            "right_col": right_col,
            "left_letter": get_column_letter(left_col),
            "right_letter": get_column_letter(right_col),
            "height": bottom_row - header_row + 1,
            "width": right_col - left_col + 1,
        },
        "pre_table_cells": pre_table,
        "rows": table_rows,
    }


def main():
    ap = argparse.ArgumentParser(description="Парсер листа 'Диаграмма Спагетти' (без диаграммы)")
    ap.add_argument("-i", "--input", required=True)
    ap.add_argument("-s", "--sheet", default="Диаграмма Спагетти")
    ap.add_argument("-o", "--output")
    args = ap.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
