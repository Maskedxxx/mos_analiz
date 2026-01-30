#!/usr/bin/env python3
"""
Парсер диаграмм на листе "ПА-1" (участок после стартовой таблицы).
Считывает встроенные графики, извлекает:
  - тип диаграммы и её заголовок;
  - координаты привязки (anchor) в листе;
  - список категорий (подписи оси X) с диапазоном;
  - для каждой серии: имя, диапазон значений и сами значения (data_only).

Ориентирован на текущую структуру: один столбчатый график, построенный
на данных таблицы выше (B1:V1 и B3:V12), но работает с произвольным числом
диаграмм на листе.
"""

import argparse
import json
import zipfile
from pathlib import Path
from typing import Dict, List, Any, Optional, Tuple

from openpyxl import load_workbook
from openpyxl.chart._chart import ChartBase
from openpyxl.utils import range_boundaries


def _sheet_name_from_range(rng: str) -> Tuple[str, str]:
    """Возвращает (sheet_name, range_part) из строки вида 'Лист'!$A$1:$B$2."""
    if "!" in rng:
        sheet, r = rng.split("!", 1)
        sheet = sheet.strip("'")
        return sheet, r
    return "", rng


def _values_from_range(wb, sheet_name: str, rng: str):
    ws = wb[sheet_name]
    min_col, min_row, max_col, max_row = range_boundaries(rng)
    values = []
    for r in range(min_row, max_row + 1):
        row = []
        for c in range(min_col, max_col + 1):
            row.append(ws.cell(r, c).value)
        values.append(row)
    # Приведение: если диапазон в одну строку или в один столбец — плоский список
    if min_row == max_row or min_col == max_col:
        return [v[0] if min_col == max_col else v for v in values] if min_row != max_row else values[0]
    return values


def _chart_title(chart: ChartBase) -> Optional[str]:
    t = chart.title
    if t is None:
        return None
    # openpyxl хранит как RichText
    try:
        if t.tx and t.tx.rich and t.tx.rich.p:
            return "".join(r.t for r in t.tx.rich.p[0].r)
    except Exception:
        pass
    return None


def extract_chart_payload(chart: ChartBase, wb_data, wb_formulas) -> Dict[str, Any]:
    payload: Dict[str, Any] = {
        "type": type(chart).__name__,
        "title": _chart_title(chart),
    }

    # anchor
    if getattr(chart, "anchor", None) and getattr(chart.anchor, "_from", None):
        a_from = chart.anchor._from
        a_to = chart.anchor.to
        payload["anchor"] = {
            "from": {"col": a_from.col + 1, "row": a_from.row + 1},
            "to": {"col": a_to.col + 1, "row": a_to.row + 1},
        }

    series_list = []
    for s in chart.series:
        # название серии
        name = None
        if s.title:
            if getattr(s.title, "v", None):
                name = s.title.v
            elif getattr(s.title, "strRef", None) and s.title.strRef.strCache and s.title.strRef.strCache.pt:
                name = s.title.strRef.strCache.pt[0].v

        # значения
        val_ref = getattr(getattr(s, "val", None), "numRef", None)
        val_range = val_ref.f if val_ref else None
        values = None
        if val_range:
            sheet_name, rng = _sheet_name_from_range(val_range)
            values = _values_from_range(wb_data, sheet_name, rng)

        # категории
        cat_ref_obj = getattr(getattr(s, "cat", None), "strRef", None)
        cat_range = cat_ref_obj.f if cat_ref_obj else None
        categories = None
        if cat_range:
            sheet_name, rng = _sheet_name_from_range(cat_range)
            categories = _values_from_range(wb_data, sheet_name, rng)

        series_list.append(
            {
                "name": name,
                "values_range": val_range,
                "values": values,
                "categories_range": cat_range,
                "categories": categories,
            }
        )

    payload["series"] = series_list
    return payload


def build_payload(xlsx_path: Path, sheet: str):
    # два открытия: значения и формулы (если потом потребуется)
    wb_data = load_workbook(xlsx_path, data_only=True)
    wb_formulas = load_workbook(xlsx_path, data_only=False)
    ws = wb_formulas[sheet]
    charts = getattr(ws, "_charts", [])

    charts_payload = [extract_chart_payload(ch, wb_data, wb_formulas) for ch in charts]

    # Текстовые блоки (textbox) из drawing: openpyxl их теряет, читаем напрямую из drawingX.xml
    text_boxes: List[Dict[str, Any]] = []
    try:
        with zipfile.ZipFile(xlsx_path) as zf:
            import xml.etree.ElementTree as ET

            ns_wb = {
                "n": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
                "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
            }
            wb_root = ET.fromstring(zf.read("xl/workbook.xml"))
            rid = None
            for sh in wb_root.findall("n:sheets/n:sheet", ns_wb):
                if sh.attrib.get("name") == sheet:
                    rid = sh.attrib[f"{{{ns_wb['r']}}}id"]
                    break
            if rid:
                wb_rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
                sheet_target = None
                for rel in wb_rels:
                    if rel.attrib.get("Id") == rid:
                        sheet_target = rel.attrib["Target"]
                        break
                if sheet_target:
                    sheet_rels_path = "xl/" + sheet_target.replace("worksheets/", "worksheets/_rels/") + ".rels"
                    if sheet_rels_path in zf.namelist():
                        sheet_rels = ET.fromstring(zf.read(sheet_rels_path))
                        drawing_target = None
                        for rel in sheet_rels:
                            if rel.attrib.get("Type") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing":
                                drawing_target = rel.attrib["Target"]
                                break
                        if drawing_target:
                            drawing_target = drawing_target.lstrip("../")
                            drawing_path = "xl/" + drawing_target
                            if drawing_path in zf.namelist():
                                ns = {
                                    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
                                    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
                                }
                                drw_root = ET.fromstring(zf.read(drawing_path))
                                for two in drw_root.findall("xdr:twoCellAnchor", ns):
                                    text_parts = [t.text or "" for t in two.findall(".//a:t", ns)]
                                    if text_parts:
                                        frm = two.find("xdr:from", ns)
                                        to = two.find("xdr:to", ns)
                                        text_boxes.append(
                                            {
                                                "text": "".join(text_parts).strip(),
                                                "anchor": {
                                                    "from": {
                                                        "col": int(frm.find("xdr:col", ns).text) + 1,
                                                        "row": int(frm.find("xdr:row", ns).text) + 1,
                                                    },
                                                    "to": {
                                                        "col": int(to.find("xdr:col", ns).text) + 1,
                                                        "row": int(to.find("xdr:row", ns).text) + 1,
                                                    },
                                                },
                                            }
                                        )
    except Exception:
        pass

    return {
        "meta": {"workbook": str(xlsx_path), "sheet": sheet},
        "charts": charts_payload,
        "text_boxes": text_boxes,
    }


def main():
    ap = argparse.ArgumentParser(description="Парсер диаграмм на листе 'ПА-1' (после таблицы)")
    ap.add_argument("-i", "--input", required=True, help="Путь к XLSX")
    ap.add_argument("-s", "--sheet", default="ПА-1")
    ap.add_argument("-o", "--output", help="JSON вывод (stdout если не указан)")
    args = ap.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
