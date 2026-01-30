from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple, Union

import pandas as pd
from docx import Document
from openai import OpenAI
from docx.text.paragraph import Paragraph as DocxParagraph


@dataclass(frozen=True)
class RuleSpec:
    index: int
    compare: str  # template | target_only
    scope: str    # HEADER | BODY | META | HEADER+META | ALL ...
    title: str
    instructions: List[str]


SYSTEM_PROMPT_ORDER = """---- SYSTEM ----
Ты — строгий аудитор документов.

Ты проверяешь ОДНО правило за один запрос.

Тебе даётся:
- фрагменты целевого документа (TARGET_HEADER, TARGET_BODY, TARGET_META);
- фрагменты шаблона (TEMPLATE_HEADER, TEMPLATE_BODY);
- параметры правила: RULE_INDEX, COMPARE, SCOPE, RULE_TITLE;
- текстовые инструкции правила (RULE_INSTRUCTIONS).

COMPARE:
- template — можно использовать соответствующие блоки шаблона и целевого документа:
  - для заголовка: HEADER target ↔ HEADER template;
  - для текста: BODY target ↔ BODY template.
- target_only — используй только целевой документ, шаблон игнорируй полностью.

SCOPE:
- HEADER — используй только TARGET_HEADER (и TEMPLATE_HEADER, если COMPARE=template).
- BODY — используй только TARGET_BODY (и TEMPLATE_BODY, если COMPARE=template).
- META — используй только TARGET_META.
- HEADER+META — используй TARGET_HEADER и TARGET_META.
- ALL — можно использовать все доступные блоки.

Правило области видимости:
- Ты обязан учитывать ТОЛЬКО те блоки, которые разрешены SCOPE и COMPARE для ЭТОГО правила.
- Остальные блоки контекста игнорируй, даже если там есть подходящий текст.

Формат ответа:
- Верни ТОЛЬКО JSON-массив.
- Если нарушений нет по этому правилу — верни [].
- Если есть одно или несколько нарушений по этому правилу — верни массив объектов вида:
  {
    "rule_index": <номер правила из RULE_INDEX>,
    "rule_title": "<RULE_TITLE>",
    "Целевой документ": "<что есть или 'отсутствует'>",
    "Различие": "<что не так / что должно быть>"
  }

Не добавляй никакой текст вне JSON.
---- SYSTEM ----
"""


SYSTEM_PROMPT_REG = """---- SYSTEM ----
Ты — строгий аудитор документов.

Ты проверяешь ОДНО правило за один запрос.

Дано:
- фрагменты целевого документа: TARGET_HEADER, TARGET_BODY, TARGET_SECTIONS, TARGET_CATEGORIES, TARGET_META;
- фрагменты шаблона: TEMPLATE_HEADER, TEMPLATE_BODY, TEMPLATE_SECTIONS, TEMPLATE_CATEGORIES;
- параметры правила: RULE_INDEX, COMPARE, SCOPE, RULE_TITLE;
- текстовые инструкции правила: RULE_INSTRUCTIONS.

COMPARE:
- template — можно использовать соответствующие блоки шаблона и целевого документа (HEADER↔HEADER, BODY↔BODY, SECTIONS↔SECTIONS, CATEGORIES↔CATEGORIES).
- target_only — используй только целевой документ, шаблон игнорируй полностью.

SCOPE (какие блоки можно использовать в ЭТОМ правиле):
- HEADER, BODY, SECTIONS, CATEGORIES, META, ALL (можно всё из target; если COMPARE=template — параллельные блоки шаблона).
- Учитывай ТОЛЬКО разрешённые блоки. Остальные игнорируй даже если там есть совпадающий текст.

Плейсхолдеры в шаблоне:
- Любой фрагмент вида <...> — плейсхолдер.
- Перед сравнением заменяй ВСЕ плейсхолдеры в строках на единый маркер ("<*>") и затем сравнивай.
- Отличия только из-за конкретных значений вместо плейсхолдеров НЕ считать нарушением.

Формат ответа:
- Верни ТОЛЬКО JSON-массив.
- Если нарушений нет — верни []
- Нарушения по правилу оформляй объектами:
  {
    "rule_index": <номер правила>,
    "rule_title": "<RULE_TITLE>",
    "Целевой документ": "<что есть или 'отсутствует'>",
    "Различие": "<что не так / что должно быть>"
  }
---- SYSTEM ----
"""


SYSTEM_PROMPT_PACKAGE = """Ты — строгий аудитор документов.

Ты проверяешь ОДНО правило за один запрос.

Формат ответа:
- Верни ТОЛЬКО JSON-массив.
- Если нарушений нет — верни [].
- Каждое нарушение оформляй объектом:
  {
    "rule_index": <номер правила>,
    "rule_title": "<название правила>",
    "Целевой документ": "<что есть в документе или 'отсутствует'>",
    "Различие": "<в чем несоответствие / что должно быть>"
  }

Плейсхолдеры вида <...> считать произвольным текстом;
отличия только из-за их заполнения НЕ считать нарушением.
"""


STOP_TOKENS = ("лист ознакомления",)


def _clean_text(text: str) -> str:
    text = text.strip()
    hint_patterns = (" — ", " —", "— ", " – ", " –", "– ", " - ", " -", "- ")
    for pat in hint_patterns:
        idx = text.find(pat)
        if idx != -1:
            text = text[:idx]
            break
    return " ".join(text.split())


def parse_order_doc(path: Path) -> Dict[str, str]:
    doc = Document(path)
    paragraphs: List[str] = []
    for para in doc.paragraphs:
        raw = para.text
        stripped = raw.strip()
        if not stripped:
            continue
        if any(stripped.lower().startswith(tok) for tok in STOP_TOKENS):
            break
        paragraphs.append(stripped)

    start_markers = ("с целью обеспечения", "с целью", "в целях обеспечения")
    end_markers = (
        "приказа оставляю за собой",
        "контроль исполнения настоящего приказа оставляю за собой",
        "контроль исполнения приказа оставляю за собой",
    )

    header: List[str] = []
    body: List[str] = []
    footer: List[str] = []

    state = "header"
    for para in paragraphs:
        cleaned = _clean_text(para)
        if not cleaned:
            continue
        low = cleaned.lower()

        if state == "header":
            if any(m in low for m in start_markers):
                state = "body"
                body.append(cleaned)
            else:
                header.append(cleaned)
        elif state == "body":
            body.append(cleaned)
            if any(m in low for m in end_markers):
                state = "footer"
        else:
            footer.append(cleaned)

    if not body:
        return {"Имя файла": path.name, "чанк_шапка": "\n".join(header), "чанк_текст": "", "чанк_мета": ""}
    return {"Имя файла": path.name, "чанк_шапка": "\n".join(header), "чанк_текст": "\n".join(body), "чанк_мета": "\n".join(footer)}


def _clean(text: str) -> str:
    return " ".join(text.replace("\xa0", " ").split())


def _normalize_heading(text: str) -> str:
    low = text.lower().strip()
    low = re.sub(r"^[\d\s\.]+", "", low)
    low = low.rstrip(":")
    return " ".join(low.split())


def parse_reg_comp_ppu(path: Path) -> Dict[str, object]:
    doc = Document(path)
    paragraphs: List[str] = []

    if doc.sections:
        hdr = doc.sections[0].header
        for para in hdr.paragraphs:
            t = _clean(para.text)
            if t:
                paragraphs.append(t)

    for para in doc.paragraphs:
        t = _clean(para.text)
        if t:
            paragraphs.append(t)

    section_titles = (
        "общие положения",
        "термины, сокращения, определения",
        "распределение функций и ответственности между участниками процесса по подаче предложений",
        "порядок организации и проведения конкурсов предложений по улучшениям",
    )

    header: List[str] = []
    sections: List[Dict[str, str]] = []
    current_title: Optional[str] = None
    current_body: List[str] = []

    def flush() -> None:
        if current_title is not None:
            sections.append({"заголовок": current_title, "текст": "\n".join(current_body).strip()})

    for text in paragraphs:
        norm = _normalize_heading(text)
        if "таблица 2" in norm or norm.startswith("номинации") or "номинации для проектов по улучшениям" in norm:
            flush()
            current_title = None
            current_body = []
            break

        if any(norm.startswith(prefix) for prefix in section_titles):
            flush()
            current_title = text
            current_body = []
            continue

        if current_title is None:
            header.append(text)
        else:
            current_body.append(text)

    flush()
    return {"Имя файла": path.name, "чанк_шапка": "\n".join(header).strip(), "чанк_разделы": sections}


def parse_reg_ppu(path: Path) -> Dict[str, object]:
    doc = Document(path)
    paragraphs: List[str] = []

    if doc.sections:
        hdr = doc.sections[0].header
        for para in hdr.paragraphs:
            t = _clean(para.text)
            if t:
                paragraphs.append(t)

    for para in doc.paragraphs:
        t = _clean(para.text)
        if t:
            paragraphs.append(t)

    section_titles = (
        "общие положения",
        "термины, сокращения, определения",
        "распределение функций и ответственности между участниками процесса по подаче предложений",
        "категории предложений",
    )
    stop_prefixes = ("форма №", "утверждаю", "а к т", "акт о", "порядок премирования")

    header: List[str] = []
    sections: List[Dict[str, str]] = []
    current_title: Optional[str] = None
    current_body: List[str] = []

    def flush() -> None:
        if current_title is not None:
            sections.append({"заголовок": current_title, "текст": "\n".join(current_body).strip()})

    for text in paragraphs:
        norm = _normalize_heading(text)
        if any(norm.startswith(pref) for pref in stop_prefixes):
            flush()
            current_title = None
            current_body = []
            break

        if any(norm.startswith(prefix) for prefix in section_titles):
            flush()
            current_title = text
            current_body = []
            continue

        if current_title is None:
            header.append(text)
        else:
            current_body.append(text)

    flush()

    categories_table: List[Dict[str, str]] = []
    if doc.tables:
        table = doc.tables[0]
        if table.rows:
            headers = [_clean(cell.text) for cell in table.rows[0].cells]
            for row in table.rows[1:]:
                values = [_clean(cell.text) for cell in row.cells]
                if all(not v for v in values):
                    continue
                categories_table.append({headers[i]: values[i] for i in range(min(len(headers), len(values)))})

    return {
        "Имя файла": path.name,
        "чанк_шапка": "\n".join(header).strip(),
        "чанк_разделы": sections,
        "таблица_категорий": categories_table,
    }


def _is_numbered(paragraph: DocxParagraph) -> bool:
    ppr = paragraph._p.pPr if paragraph._p is not None else None
    if ppr is None or ppr.numPr is None:
        return False
    return ppr.numPr.numId is not None


def _is_bold(paragraph: DocxParagraph) -> bool:
    return any(run.bold for run in paragraph.runs if run.text and run.text.strip())


def parse_order_tz(path: Path) -> Dict[str, object]:
    document = Document(path)

    title: str | None = None
    rules: List[Dict[str, Any]] = []
    current: Dict[str, Any] | None = None

    for paragraph in document.paragraphs:
        raw_text = paragraph.text
        if not raw_text or not raw_text.strip():
            continue
        normalized = " ".join(raw_text.split())
        bold = _is_bold(paragraph)
        numbered = _is_numbered(paragraph)

        if title is None and bold and not numbered:
            title = normalized
            continue

        if numbered:
            if current is not None:
                rules.append(current)
            idx = len(rules) + 1
            rule_title = normalized if bold else f"Правило {idx}"
            current = {"index": idx, "title": rule_title, "content": []}
            continue

        if current is not None:
            current["content"].append(normalized)

    if current is not None:
        rules.append(current)

    return {"Имя файла": path.name, "Заголовок документа": title or path.stem, "правила": rules}


def parse_reg_comp_ppu_tz(path: Path) -> Dict[str, object]:
    document = Document(path)
    title: Optional[str] = None
    rules: List[Dict[str, Any]] = []
    current: Optional[Dict[str, Any]] = None

    paragraphs = [p.text for p in document.paragraphs if p.text and p.text.strip()]
    rule_start_prefixes = ("проверка", "сверить")

    for para in paragraphs:
        text = _clean(para)
        if not text:
            continue
        if title is None:
            title = text
            continue
        norm = text.lower()
        if any(norm.startswith(prefix) for prefix in rule_start_prefixes):
            if current is not None:
                rules.append(current)
            current = {"index": len(rules) + 1, "title": text, "content": []}
            continue
        if current is not None:
            current["content"].append(text)

    if current is not None:
        rules.append(current)

    return {"Имя файла": path.name, "Заголовок документа": title or path.stem, "правила": rules}


def parse_reg_ppu_tz(path: Path) -> Dict[str, object]:
    document = Document(path)
    paragraphs = [_clean(p.text) for p in document.paragraphs if p.text and p.text.strip()]
    title = paragraphs[0] if paragraphs else path.name
    blocks: List[Dict[str, Any]] = []
    current: Optional[Dict[str, Any]] = None

    def is_rule_start(text: str) -> bool:
        low = text.lower().strip()
        return low.startswith(("проверка", "сверить"))

    for text in paragraphs[1:]:
        if is_rule_start(text):
            if current:
                blocks.append(current)
            current = {"index": len(blocks) + 1, "title": text, "content": []}
        else:
            if current is None:
                continue
            current["content"].append(text)

    if current:
        blocks.append(current)

    return {"document_title": title, "blocks": blocks}


def _strip_code_fence(text: str) -> str:
    t = text.strip()
    if t.startswith("```"):
        if t.endswith("```"):
            t = t[3:-3]
        else:
            t = t[3:]
    return t.strip()


def _try_parse_json(text: str) -> Union[List, Dict, None]:
    cleaned = _strip_code_fence(text)
    try:
        return json.loads(cleaned)
    except Exception:
        return None


def _save_xlsx(data: Any, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(data, list):
        if data and isinstance(data[0], dict):
            df = pd.DataFrame(data)
        else:
            df = pd.DataFrame({"data": data})
    elif isinstance(data, dict):
        df = pd.DataFrame([data])
    else:
        df = pd.DataFrame({"raw": [data]})
    df.to_excel(path, index=False)


def _join_sections(sections: List[Dict[str, str]]) -> str:
    parts: List[str] = []
    for sec in sections:
        title = sec.get("заголовок", "")
        text = sec.get("текст", "")
        merged = (title + "\n" + text).strip()
        if merged:
            parts.append(merged)
    return "\n\n".join(parts)


def _format_categories(table: List[Dict[str, str]]) -> str:
    lines: List[str] = []
    for row in table:
        row_items = [f"{k}: {v}" for k, v in row.items()]
        lines.append(" | ".join(row_items))
    return "\n".join(lines)


def _call_llm(messages: List[Dict[str, str]], model: str, temperature: float) -> str:
    api_key = os.environ.get("OPENAI_API_KEY")
    if not api_key:
        raise RuntimeError("OPENAI_API_KEY is not set")
    client = OpenAI(api_key=api_key)
    resp = client.chat.completions.create(model=model, messages=messages, temperature=temperature)
    return resp.choices[0].message.content or ""


def build_order_messages(spec: RuleSpec, target_doc: Dict[str, str], template_doc: Dict[str, str]) -> List[Dict[str, str]]:
    target_header = target_doc.get("чанк_шапка", "")
    target_body = target_doc.get("чанк_текст", "")
    target_meta = target_doc.get("чанк_мета", "")
    template_header = template_doc.get("чанк_шапка", "")
    template_body = template_doc.get("чанк_текст", "")

    scope_raw = spec.scope or "ALL"
    scopes = {part.strip().upper() for part in scope_raw.split("+") if part.strip()}

    def has_scope(name: str) -> bool:
        name = name.upper()
        return "ALL" in scopes or name in scopes

    user_parts: List[str] = []
    user_parts.append(f"RULE_INDEX: {spec.index}")
    user_parts.append(f"COMPARE: {spec.compare}")
    user_parts.append(f"SCOPE: {spec.scope}")
    user_parts.append(f"RULE_TITLE: {spec.title}")
    user_parts.append("RULE_INSTRUCTIONS:")
    for line in spec.instructions:
        user_parts.append(f"- {line}")
    user_parts.append("")
    user_parts.append("CONTEXT:")

    if has_scope("HEADER"):
        user_parts.append("[TARGET_HEADER]")
        user_parts.append(target_header)
        user_parts.append("[/TARGET_HEADER]")
    if has_scope("BODY"):
        user_parts.append("[TARGET_BODY]")
        user_parts.append(target_body)
        user_parts.append("[/TARGET_BODY]")
    if has_scope("META"):
        user_parts.append("[TARGET_META]")
        user_parts.append(target_meta)
        user_parts.append("[/TARGET_META]")

    if spec.compare == "template":
        if has_scope("HEADER"):
            user_parts.append("[TEMPLATE_HEADER]")
            user_parts.append(template_header)
            user_parts.append("[/TEMPLATE_HEADER]")
        if has_scope("BODY"):
            user_parts.append("[TEMPLATE_BODY]")
            user_parts.append(template_body)
            user_parts.append("[/TEMPLATE_BODY]")

    if spec.compare == "target_only":
        joined = "\n".join(user_parts)
        if "[TEMPLATE_HEADER]" in joined or "[TEMPLATE_BODY]" in joined:
            raise RuntimeError("BUG: compare=target_only, but TEMPLATE_* leaked into prompt")

    return [{"role": "system", "content": SYSTEM_PROMPT_ORDER}, {"role": "user", "content": "\n".join(user_parts)}]


def build_reg_messages(spec: RuleSpec, target_doc: Dict[str, Any], template_doc: Dict[str, Any]) -> List[Dict[str, str]]:
    target_header = target_doc.get("чанк_шапка", "")
    target_sections = target_doc.get("чанк_разделы", [])
    target_body = target_doc.get("чанк_текст", _join_sections(target_sections))
    target_categories = target_doc.get("таблица_категорий", [])
    target_meta = target_doc.get("чанк_мета", "")

    template_header = template_doc.get("чанк_шапка", "")
    template_sections = template_doc.get("чанк_разделы", [])
    template_body = template_doc.get("чанк_текст", _join_sections(template_sections))
    template_categories = template_doc.get("таблица_категорий", [])

    scope_raw = spec.scope or "ALL"
    scopes = {part.strip().upper() for part in scope_raw.split("+") if part.strip()}

    def has_scope(name: str) -> bool:
        name = name.upper()
        return "ALL" in scopes or name in scopes

    user_parts: List[str] = []
    user_parts.append(f"RULE_INDEX: {spec.index}")
    user_parts.append(f"COMPARE: {spec.compare}")
    user_parts.append(f"SCOPE: {spec.scope}")
    user_parts.append(f"RULE_TITLE: {spec.title}")
    user_parts.append("RULE_INSTRUCTIONS:")
    for line in spec.instructions:
        user_parts.append(f"- {line}")
    user_parts.append("")
    user_parts.append("CONTEXT:")

    if has_scope("HEADER"):
        user_parts.append("[TARGET_HEADER]")
        user_parts.append(target_header)
        user_parts.append("[/TARGET_HEADER]")
    if has_scope("BODY"):
        user_parts.append("[TARGET_BODY]")
        user_parts.append(target_body)
        user_parts.append("[/TARGET_BODY]")
    if has_scope("SECTIONS"):
        user_parts.append("[TARGET_SECTIONS]")
        for sec in target_sections:
            title = sec.get("заголовок", "")
            text = sec.get("текст", "")
            user_parts.append(f'<section title="{title}">{text}</section>')
        user_parts.append("[/TARGET_SECTIONS]")
    if has_scope("CATEGORIES"):
        user_parts.append("[TARGET_CATEGORIES]")
        user_parts.append(_format_categories(target_categories))
        user_parts.append("[/TARGET_CATEGORIES]")
    if has_scope("META"):
        user_parts.append("[TARGET_META]")
        user_parts.append(target_meta)
        user_parts.append("[/TARGET_META]")

    if spec.compare == "template":
        if has_scope("HEADER"):
            user_parts.append("[TEMPLATE_HEADER]")
            user_parts.append(template_header)
            user_parts.append("[/TEMPLATE_HEADER]")
        if has_scope("BODY"):
            user_parts.append("[TEMPLATE_BODY]")
            user_parts.append(template_body)
            user_parts.append("[/TEMPLATE_BODY]")
        if has_scope("SECTIONS"):
            user_parts.append("[TEMPLATE_SECTIONS]")
            for sec in template_sections:
                title = sec.get("заголовок", "")
                text = sec.get("текст", "")
                user_parts.append(f'<section title="{title}">{text}</section>')
            user_parts.append("[/TEMPLATE_SECTIONS]")
        if has_scope("CATEGORIES"):
            user_parts.append("[TEMPLATE_CATEGORIES]")
            user_parts.append(_format_categories(template_categories))
            user_parts.append("[/TEMPLATE_CATEGORIES]")

    if spec.compare == "target_only":
        joined = "\n".join(user_parts)
        if any(tag in joined for tag in ("[TEMPLATE_HEADER]", "[TEMPLATE_BODY]", "[TEMPLATE_SECTIONS]", "[TEMPLATE_CATEGORIES]")):
            raise RuntimeError("BUG: compare=target_only, but TEMPLATE_* leaked into prompt")

    return [{"role": "system", "content": SYSTEM_PROMPT_REG}, {"role": "user", "content": "\n".join(user_parts)}]


def _build_rule_specs(tz_data: Dict[str, Any], compare_template_rules: set[int], scope_by_idx: Dict[int, str]) -> List[RuleSpec]:
    rules_list = tz_data.get("правила") or tz_data.get("blocks") or []
    specs: List[RuleSpec] = []
    for raw in rules_list:
        try:
            idx = int(raw.get("index", 0))
        except Exception:
            continue
        compare = "template" if idx in compare_template_rules else "target_only"
        scope = scope_by_idx.get(idx, "ALL")
        specs.append(
            RuleSpec(
                index=idx,
                compare=compare,
                scope=scope,
                title=raw.get("title", ""),
                instructions=raw.get("content", []) or [],
            )
        )
    return specs


def audit_order(
    target_path: Path,
    template_path: Path,
    tz_path: Path,
    model: str,
    temperature: float,
    ruleset: Dict[str, Any],
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    target_doc = parse_order_doc(target_path)
    template_doc = parse_order_doc(template_path)
    tz_data = parse_order_tz(tz_path)

    compare_template_rules = set(ruleset["compare_template_rules"])
    scope_by_idx = {int(k): v for k, v in ruleset["scope_by_idx"].items()}
    rule_specs = _build_rule_specs(tz_data, compare_template_rules, scope_by_idx)

    violations: List[Dict[str, Any]] = []
    for spec in rule_specs:
        messages = build_order_messages(spec, target_doc, template_doc)
        raw = _call_llm(messages, model=model, temperature=temperature)
        parsed = _try_parse_json(raw)
        if parsed is None:
            continue
        if isinstance(parsed, dict):
            parsed_list = [parsed]
        elif isinstance(parsed, list):
            parsed_list = [x for x in parsed if isinstance(x, dict)]
        else:
            parsed_list = []
        for obj in parsed_list:
            obj.setdefault("rule_index", spec.index)
            obj.setdefault("rule_title", spec.title)
        violations.extend(parsed_list)

    debug = {"target_doc": target_doc, "template_doc": template_doc, "tz_data": tz_data}
    return violations, debug


def audit_reg(
    target_path: Path,
    template_path: Path,
    tz_path: Path,
    model: str,
    temperature: float,
    ruleset: Dict[str, Any],
    reg_type: str,
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    target_doc = parse_reg_ppu(target_path) if reg_type == "reg_ppu" else parse_reg_comp_ppu(target_path)
    template_doc = parse_reg_ppu(template_path) if reg_type == "reg_ppu" else parse_reg_comp_ppu(template_path)
    tz_data = parse_reg_ppu_tz(tz_path) if reg_type == "reg_ppu" else parse_reg_comp_ppu_tz(tz_path)

    compare_template_rules = set(ruleset["compare_template_rules"])
    scope_by_idx = {int(k): v for k, v in ruleset["scope_by_idx"].items()}
    rule_specs = _build_rule_specs(tz_data, compare_template_rules, scope_by_idx)

    violations: List[Dict[str, Any]] = []
    for spec in rule_specs:
        messages = build_reg_messages(spec, target_doc, template_doc)
        raw = _call_llm(messages, model=model, temperature=temperature)
        parsed = _try_parse_json(raw)
        if parsed is None:
            continue
        if isinstance(parsed, dict):
            parsed_list = [parsed]
        elif isinstance(parsed, list):
            parsed_list = [x for x in parsed if isinstance(x, dict)]
        else:
            parsed_list = []
        for obj in parsed_list:
            obj.setdefault("rule_index", spec.index)
            obj.setdefault("rule_title", spec.title)
        violations.extend(parsed_list)

    debug = {"target_doc": target_doc, "template_doc": template_doc, "tz_data": tz_data}
    return violations, debug


def _find_line_starting(text_block: str, prefix: str) -> str:
    for line in (text_block or "").splitlines():
        stripped = line.strip()
        if stripped.startswith(prefix):
            return stripped
    return ""


def _extract_point2(order: Dict[str, Any]) -> str:
    body = order.get("чанк_текст", "") or ""
    return _find_line_starting(body, "2.")


def _extract_point31(reg: Dict[str, Any]) -> str:
    sections = reg.get("чанк_разделы", []) or []
    for sec in sections:
        title = (sec.get("заголовок") or "").strip()
        if title.startswith("3"):
            line = _find_line_starting(sec.get("текст", "") or "", "3.1")
            if line:
                return line
    for sec in sections:
        line = _find_line_starting(sec.get("текст", "") or "", "3.1")
        if line:
            return line
    return ""


def audit_package(
    order_comp_path: Path,
    reg_comp_path: Path,
    order_ppu_path: Path,
    reg_ppu_path: Path,
    model: str,
    temperature: float,
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    docs = {
        "order_comp_ppu": parse_order_doc(order_comp_path),
        "reg_comp_ppu": parse_reg_comp_ppu(reg_comp_path),
        "order_ppu": parse_order_doc(order_ppu_path),
        "reg_ppu": parse_reg_ppu(reg_ppu_path),
    }

    rules = [
        {"index": 1, "title": "Сверка номеров и дат приказов и положений (оба комплекта документов)", "kind": "headers"},
        {"index": 2, "title": "Сверка должности и ФИО (оба комплекта документов)", "kind": "fio_post"},
    ]

    results: List[Dict[str, Any]] = []
    for rule in rules:
        user_parts: List[str] = []
        user_parts.append(f"RULE_INDEX: {rule['index']}")
        user_parts.append(f"RULE_TITLE: {rule['title']}")
        user_parts.append("RULE_INSTRUCTIONS:")

        if rule["kind"] == "headers":
            user_parts.append("- Норма: для КАЖДОЙ пары документов номер и дата приказа совпадают между приказом и положением.")
            user_parts.append("- Пара 1: Приказ о конкурсах ППУ ↔ Положение о конкурсах ППУ.")
            user_parts.append("- Пара 2: Приказ о ППУ ↔ Положение о ППУ.")
            user_parts.append("- Если норма по обеим парам: верни []. Если есть нарушения — верни массив объектов (по одному объекту на каждую проблемную пару).")

            user_parts.append("CONTEXT:")
            user_parts.append("[ORDER_COMP_HEADER]")
            user_parts.append(docs["order_comp_ppu"].get("чанк_шапка", ""))
            user_parts.append("[/ORDER_COMP_HEADER]")
            user_parts.append("[REG_COMP_HEADER]")
            user_parts.append(docs["reg_comp_ppu"].get("чанк_шапка", ""))
            user_parts.append("[/REG_COMP_HEADER]")
            user_parts.append("[ORDER_PPU_HEADER]")
            user_parts.append(docs["order_ppu"].get("чанк_шапка", ""))
            user_parts.append("[/ORDER_PPU_HEADER]")
            user_parts.append("[REG_PPU_HEADER]")
            user_parts.append(docs["reg_ppu"].get("чанк_шапка", ""))
            user_parts.append("[/REG_PPU_HEADER]")
        else:
            user_parts.append("- Норма: должность и ФИО из пункта 2 каждого приказа совпадают с должностью и ФИО из пункта 3.1 соответствующего положения.")
            user_parts.append("- Если норма по обеим парам: верни []. Если есть нарушения — верни массив объектов (по одному объекту на каждую проблемную пару).")

            user_parts.append("CONTEXT:")
            user_parts.append("[ORDER_COMP_POINT2]")
            user_parts.append(_extract_point2(docs["order_comp_ppu"]))
            user_parts.append("[/ORDER_COMP_POINT2]")
            user_parts.append("[REG_COMP_POINT31]")
            user_parts.append(_extract_point31(docs["reg_comp_ppu"]))
            user_parts.append("[/REG_COMP_POINT31]")
            user_parts.append("[ORDER_PPU_POINT2]")
            user_parts.append(_extract_point2(docs["order_ppu"]))
            user_parts.append("[/ORDER_PPU_POINT2]")
            user_parts.append("[REG_PPU_POINT31]")
            user_parts.append(_extract_point31(docs["reg_ppu"]))
            user_parts.append("[/REG_PPU_POINT31]")

        messages = [{"role": "system", "content": SYSTEM_PROMPT_PACKAGE}, {"role": "user", "content": "\n".join(user_parts)}]
        raw = _call_llm(messages, model=model, temperature=temperature)
        parsed = _try_parse_json(raw)
        if parsed is None:
            continue
        if isinstance(parsed, dict):
            parsed_list = [parsed]
        elif isinstance(parsed, list):
            parsed_list = [x for x in parsed if isinstance(x, dict)]
        else:
            parsed_list = []
        results.extend(parsed_list)

    return results, {"docs": docs}


def write_outputs(run_dir: Path, violations: List[Dict[str, Any]], debug: Dict[str, Any]) -> Dict[str, str]:
    run_dir.mkdir(parents=True, exist_ok=True)

    (run_dir / "violations.json").write_text(json.dumps(violations, ensure_ascii=False, indent=2), encoding="utf-8")
    (run_dir / "debug_parsed.json").write_text(json.dumps(debug, ensure_ascii=False, indent=2, default=str), encoding="utf-8")

    report_path = run_dir / "report.xlsx"
    _save_xlsx(violations, report_path)
    return {"violations_json": "violations.json", "report_xlsx": "report.xlsx", "debug_json": "debug_parsed.json"}

