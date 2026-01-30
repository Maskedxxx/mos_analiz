#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Парсер для документа "График обхода ОМ" (.xlsx).

Извлекает 2 чанка с листа "График" или "Шаблон":
- шапка: наименование компании + "Приложение №1 к Приказу № XX от XX.XX.XXXX"
- фио_должности: список всех ФИО и должностей из таблицы графика

Использование:
    from parser_grafik_xlsx import parse_grafik_obhod
    chunks = parse_grafik_obhod("path/to/document.xlsx")
"""

import re
from pathlib import Path
from openpyxl import load_workbook


def parse_grafik_obhod(xlsx_path: str) -> dict:
    """
    Парсит документ "График обхода ОМ".

    Args:
        xlsx_path: путь к .xlsx файлу

    Returns:
        dict с чанками: шапка, фио_должности
    """
    wb = load_workbook(xlsx_path, data_only=True)

    # Ищем нужный лист: "График" или "Шаблон"
    sheet_name = None
    for name in wb.sheetnames:
        if name.lower() in ["график", "шаблон"]:
            sheet_name = name
            break

    if not sheet_name:
        # Берём первый лист
        sheet_name = wb.sheetnames[0]

    ws = wb[sheet_name]

    # Инициализация чанков
    chunks = {
        "имя_файла": Path(xlsx_path).name,
        "лист": sheet_name,
        "шапка": "",
        "фио_должности": ""
    }

    # Собираем все данные
    company_name = ""
    prikaz_ref = ""
    employees = []

    # Проходим по всем ячейкам первых 50 строк
    for row_num in range(1, min(51, ws.max_row + 1)):
        row_values = []
        for col_num in range(1, min(35, ws.max_column + 1)):
            cell = ws.cell(row=row_num, column=col_num)
            if cell.value:
                row_values.append(str(cell.value).strip())

        # Ищем наименование компании (ООО/ЗАО/АО/ПАО)
        for val in row_values:
            if re.search(r'\b(ООО|ЗАО|АО|ПАО|ИП)\b', val) and not company_name:
                company_name = val

        # Ищем "Приложение №1 к Приказу..."
        for val in row_values:
            if "приложение" in val.lower() and "приказ" in val.lower():
                prikaz_ref = val.replace("\n", " ").strip()

        # Ищем строки с ФИО и должностью
        # Обычно: первая ячейка = ФИО, вторая = должность
        if len(row_values) >= 2:
            first_val = row_values[0]
            second_val = row_values[1]

            # Проверяем что это не заголовок "ФИО" и не "График обхода..."
            if first_val.lower() not in ["фио", "график", ""] and \
               "график обхода" not in first_val.lower() and \
               "приложение" not in first_val.lower():

                # Проверяем паттерн ФИО (Фамилия И.О. или Фамилия Имя Отчество)
                if re.search(r'[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.?[А-ЯЁа-яё]?\.?', first_val) or \
                   re.search(r'[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+', first_val):
                    # Второй столбец должен быть должностью
                    if second_val and second_val.lower() != "должность":
                        employees.append(f"{first_val} — {second_val}")

    # Формируем чанки
    chunks["шапка"] = f"{company_name}\n{prikaz_ref}".strip()

    # Убираем дубликаты из списка сотрудников
    unique_employees = list(dict.fromkeys(employees))
    chunks["фио_должности"] = "\n".join(unique_employees)

    return chunks


# === Тест при запуске напрямую ===
if __name__ == "__main__":
    import sys

    # По умолчанию тестируем на файле ООО Пример
    project_dir = Path(__file__).parent.parent
    default_path = project_dir / "documents" / "grafik_obhod_ok.xlsx"

    if len(sys.argv) > 1:
        test_path = sys.argv[1]
    else:
        test_path = str(default_path)

    print(f"📊 Парсинг: {Path(test_path).name}")
    print("=" * 60)

    chunks = parse_grafik_obhod(test_path)

    for key, value in chunks.items():
        print(f"\n[{key}]")
        print("-" * 40)
        print(value if value else "(пусто)")
