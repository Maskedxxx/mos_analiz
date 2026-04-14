#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
DOCX-парсер: извлечение текста и таблиц из DOCX через python-docx.

Альтернатива Paddle OCR для формата DOCX — читает текст и таблицы
напрямую из XML, без конвертации в изображения.

Возвращает Dict[str, Any] в формате legacy-парсера (как PaddleParser/VisionParser).

Ограничение: DOCX не имеет фиксированных страниц (пагинация зависит от рендерера).
Поле "pages" в chunks_vision.json интерпретируется как приблизительные доли документа:
  - pages: [1] → первая часть документа
  - pages: [-1] → последняя часть документа
  - pages: "all" или полный список → весь документ
"""

import json
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from docx import Document

logger = logging.getLogger(__name__)


def parse_docx(
    file_path: str,
    chunks_vision_path: str,
    chunk_filter: Optional[str] = None,
    vlm_base_url: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Парсинг DOCX-файла через python-docx.

    Args:
        file_path: путь к DOCX-файлу
        chunks_vision_path: путь к chunks_vision.json (берём name + pages)
        chunk_filter: имя чанка (если указан — собираем только его)

    Returns:
        Dict[str, Any] — {"имя_файла": ..., "путь": ..., "чанк1": текст, ...}
    """
    file_path = Path(file_path)

    # --- Загрузка конфигурации чанков ---
    with open(chunks_vision_path, "r", encoding="utf-8") as f:
        chunks_config = json.load(f)
    chunks = chunks_config.get("chunks", [])

    # --- Открытие документа ---
    doc = Document(str(file_path))
    logger.info(f"DOCX: {file_path.name}, абзацев: {len(doc.paragraphs)}, таблиц: {len(doc.tables)}")

    # --- Извлечение текста из embedded images (шапка-картинка) через OCR ---
    header_ocr_text = _ocr_header_images(doc, vlm_base_url)
    if header_ocr_text:
        logger.info(f"  OCR шапки-картинки: {len(header_ocr_text)} chars")

    # --- Извлечение всех элементов в порядке документа ---
    elements = _extract_elements_in_order(doc)
    if header_ocr_text:
        elements.insert(0, header_ocr_text)
    total_elements = len(elements)
    logger.info(f"  Извлечено элементов: {total_elements}")

    # --- Определяем общее число "страниц" для маппинга ---
    # Берём максимальный номер страницы из всех чанков
    max_page = 1
    for chunk in chunks:
        pages = chunk.get("pages", [])
        if pages == "all":
            continue
        if isinstance(pages, str):
            # Диапазон "4-7" → [4, 5, 6, 7]
            parts = pages.split("-")
            try:
                pages = list(range(int(parts[0]), int(parts[1]) + 1))
            except (ValueError, IndexError):
                continue
        for p in pages:
            if not isinstance(p, int):
                continue
            abs_p = abs(p)
            if abs_p > max_page:
                max_page = abs_p
    total_pages = max(max_page, 1)

    # --- Разбиваем элементы на "страницы" (равные доли) ---
    pages_content = _split_into_pages(elements, total_pages)

    # --- Сборка чанков ---
    result: Dict[str, Any] = {
        "имя_файла": file_path.name,
        "путь": str(file_path.resolve()),
    }

    for chunk in chunks:
        name = chunk["name"]
        if chunk_filter and name != chunk_filter:
            continue

        pages = chunk.get("pages", [])

        # "all" или полный список → весь документ
        if pages == "all":
            result[name] = "\n\n".join(elements)
            continue

        # Нормализуем строковые диапазоны "4-7" → [4, 5, 6, 7]
        if isinstance(pages, str):
            parts_str = pages.split("-")
            try:
                pages = list(range(int(parts_str[0]), int(parts_str[1]) + 1))
            except (ValueError, IndexError):
                pages = []

        # Резолвим номера страниц
        resolved = []
        for p in pages:
            if not isinstance(p, int):
                continue
            if p < 0:
                resolved.append(total_pages + p + 1)
            else:
                resolved.append(p)

        # Собираем текст из указанных "страниц"
        parts = []
        for page_num in resolved:
            if page_num < 1 or page_num > total_pages:
                logger.warning(f"  Чанк '{name}': страница {page_num} вне диапазона (1-{total_pages})")
                continue
            page_text = pages_content.get(page_num, "")
            if page_text.strip():
                parts.append(page_text)

        result[name] = "\n\n".join(parts)

    return result


def _ocr_header_images(doc: Document, vlm_base_url: Optional[str] = None) -> str:
    """
    Извлекает embedded images из первых параграфов документа и OCR-ит их.

    Многие DOCX используют растровые бланки (логотип + реквизиты) вместо текста.
    python-docx не видит текст в таких изображениях — нужен OCR.

    Returns:
        Объединённый OCR-текст заголовочных картинок, или "" если нет картинок/VLM.
    """
    if not vlm_base_url:
        return ""

    import base64
    import io
    import requests
    from PIL import Image

    ns_a = 'http://schemas.openxmlformats.org/drawingml/2006/main'
    ns_r = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'

    ocr_parts = []
    # Проверяем только первые 6 параграфов (шапка обычно в начале)
    for p in doc.paragraphs[:6]:
        blips = p._element.findall(f'.//{{{ns_a}}}blip')
        for blip in blips:
            rId = blip.get(f'{{{ns_r}}}embed')
            if not rId or rId not in doc.part.rels:
                continue
            try:
                blob = doc.part.rels[rId].target_part.blob
                img = Image.open(io.BytesIO(blob))
                # Пропускаем мелкие картинки (иконки, декор)
                if img.width < 200 or img.height < 50:
                    continue
                buf = io.BytesIO()
                img.save(buf, format='PNG')
                b64 = base64.b64encode(buf.getvalue()).decode()
                r = requests.post(
                    f"{vlm_base_url}chat/completions",
                    json={
                        "model": "PaddleOCR-VL-1.5",
                        "messages": [{"role": "user", "content": [
                            {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}},
                            {"type": "text", "text": "OCR this image. Return all text exactly as written."}
                        ]}],
                        "max_tokens": 256, "temperature": 0.0,
                    },
                    timeout=30,
                )
                text = r.json()["choices"][0]["message"]["content"].strip()
                if text:
                    ocr_parts.append(text)
                    logger.info(f"  OCR image ({img.width}x{img.height}): {text[:80]}")
            except Exception as e:
                logger.warning(f"  OCR image failed: {e}")

    return "\n".join(ocr_parts)


def _extract_elements_in_order(doc: Document) -> List[str]:
    """
    Извлекает все элементы документа (абзацы и таблицы) в порядке появления.

    DOCX хранит абзацы и таблицы как siblings в document.body.
    Обходим XML напрямую, чтобы сохранить правильный порядок.

    Args:
        doc: объект Document (python-docx)

    Returns:
        List[str] — элементы документа в порядке появления
    """
    from docx.oxml.ns import qn

    elements = []
    body = doc.element.body

    # Маппинг XML-элементов таблиц → python-docx Table objects
    table_elements = {}
    for table in doc.tables:
        table_elements[table._element] = table

    for child in body:
        tag = child.tag

        if tag == qn("w:p"):
            # Параграф — основной текст из run-элементов
            runs_text = []
            for r in child.findall(qn("w:r")):
                for t in r.findall(qn("w:t")):
                    if t.text:
                        runs_text.append(t.text)
            full_text = "".join(runs_text).strip()
            if full_text:
                elements.append(full_text)

            # Textbox-ы внутри параграфа (w:drawing → wps:txbx → w:txbxContent)
            ns_w = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
            ns_wps = 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape'
            for txbx_content in child.findall(f'.//{{{ns_wps}}}txbx/{{{ns_w}}}txbxContent'):
                txbx_runs = []
                for tp in txbx_content.findall(f'{{{ns_w}}}p'):
                    for tr in tp.findall(f'{{{ns_w}}}r'):
                        for tt in tr.findall(f'{{{ns_w}}}t'):
                            if tt.text:
                                txbx_runs.append(tt.text)
                txbx_text = "".join(txbx_runs).strip()
                if txbx_text:
                    elements.append(txbx_text)

        elif tag == qn("w:tbl"):
            # Таблица
            table = table_elements.get(child)
            if table:
                html = _table_to_html(table)
                if html.strip():
                    elements.append(html)

    return elements


def _split_into_pages(elements: List[str], total_pages: int) -> Dict[int, str]:
    """
    Разбивает список элементов на N приблизительно равных частей ("страниц").

    Args:
        elements: список текстовых элементов
        total_pages: количество "страниц"

    Returns:
        Dict[int, str] — номер страницы (1-based) → текст
    """
    if total_pages <= 1:
        return {1: "\n\n".join(elements)}

    # Делим по количеству элементов
    chunk_size = max(1, len(elements) // total_pages)
    pages = {}
    for i in range(total_pages):
        start = i * chunk_size
        if i == total_pages - 1:
            # Последняя страница забирает остаток
            end = len(elements)
        else:
            end = start + chunk_size
        pages[i + 1] = "\n\n".join(elements[start:end])

    return pages


def _table_to_html(table) -> str:
    """
    Конвертирует таблицу python-docx в HTML <table>.

    Args:
        table: объект docx.table.Table

    Returns:
        str — HTML-таблица
    """
    rows_html = []
    for row in table.rows:
        cells_html = []
        for cell in row.cells:
            cell_text = cell.text.strip().replace("\n", " ")
            cells_html.append(f"  <td>{cell_text}</td>")
        rows_html.append("<tr>\n" + "\n".join(cells_html) + "\n</tr>")

    if not rows_html:
        return ""

    return "<table>\n" + "\n".join(rows_html) + "\n</table>"
