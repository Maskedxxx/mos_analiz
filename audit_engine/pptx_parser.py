#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PPTX-парсер: извлечение текста из презентаций через python-pptx.

Альтернатива Paddle OCR для формата PPTX — читает текст напрямую из XML,
без конвертации в изображения. Работает мгновенно и без галлюцинаций.

Возвращает Dict[str, Any] в формате legacy-парсера (как PaddleParser/VisionParser).
"""

import json
import logging
from pathlib import Path
from typing import Any, Dict, List, Optional

from pptx import Presentation
from pptx.util import Emu

logger = logging.getLogger(__name__)


def parse_pptx(
    file_path: str,
    chunks_vision_path: str,
    chunk_filter: Optional[str] = None,
) -> Dict[str, Any]:
    """
    Парсинг PPTX-файла через python-pptx.

    Args:
        file_path: путь к PPTX-файлу
        chunks_vision_path: путь к chunks_vision.json (берём name + pages)
        chunk_filter: имя чанка (если указан — собираем только его)

    Returns:
        Dict[str, Any] — {"имя_файла": ..., "путь": ..., "чанк1": текст, ...}
    """
    file_path = Path(file_path)

    # --- Загрузка конфигурации чанков ---
    with open(chunks_vision_path, "r", encoding="utf-8") as f:
        chunks_config = json.load(f)
    chunks = chunks_config.get("chunks", [])

    # --- Открытие презентации ---
    prs = Presentation(str(file_path))
    total_slides = len(prs.slides)
    logger.info(f"PPTX: {file_path.name}, слайдов: {total_slides}")

    # --- Извлечение текста по слайдам ---
    slides_text: Dict[int, str] = {}  # 1-based index → текст слайда
    for idx, slide in enumerate(prs.slides, start=1):
        text = _extract_slide_text(slide)
        slides_text[idx] = text
        logger.debug(f"  Слайд {idx}: {len(text)} символов")

    # --- Сборка чанков ---
    result: Dict[str, Any] = {
        "имя_файла": file_path.name,
        "путь": str(file_path.resolve()),
    }

    for chunk in chunks:
        name = chunk["name"]
        if chunk_filter and name != chunk_filter:
            continue

        pages = chunk.get("pages", [])
        # Резолвим отрицательные индексы: -1 = последний слайд
        resolved_pages = []
        for p in pages:
            if p < 0:
                resolved_pages.append(total_slides + p + 1)
            else:
                resolved_pages.append(p)

        # Собираем текст из указанных слайдов
        parts = []
        for page_num in resolved_pages:
            if page_num < 1 or page_num > total_slides:
                logger.warning(f"  Чанк '{name}': слайд {page_num} вне диапазона (1-{total_slides})")
                continue
            slide_text = slides_text.get(page_num, "")
            if not slide_text.strip():
                continue
            if len(resolved_pages) > 1:
                parts.append(f"[СТРАНИЦА {page_num}]\n{slide_text}")
            else:
                parts.append(slide_text)

        result[name] = "\n\n".join(parts)

    return result


def _extract_slide_text(slide) -> str:
    """
    Извлекает весь текст со слайда: заголовки, текстовые блоки, таблицы.

    Shapes сортируются по позиции (сверху вниз, слева направо).
    Таблицы конвертируются в HTML <table>.

    Args:
        slide: объект pptx.slide.Slide

    Returns:
        str — текст слайда
    """
    # Собираем shapes с их позициями для сортировки
    items: List[tuple] = []  # (top, left, text)

    for shape in slide.shapes:
        top = shape.top if shape.top is not None else 0
        left = shape.left if shape.left is not None else 0

        if shape.has_table:
            html = _table_to_html(shape.table)
            if html.strip():
                items.append((top, left, html))

        elif shape.has_text_frame:
            text = _textframe_to_text(shape.text_frame)
            if text.strip():
                items.append((top, left, text))

    # Сортировка: сверху вниз, потом слева направо
    items.sort(key=lambda x: (x[0], x[1]))

    return "\n\n".join(item[2] for item in items)


def _textframe_to_text(text_frame) -> str:
    """
    Извлекает текст из TextFrame, сохраняя абзацы.

    Args:
        text_frame: объект pptx.text.text.TextFrame

    Returns:
        str — текст с переносами строк между абзацами
    """
    paragraphs = []
    for para in text_frame.paragraphs:
        # Собираем runs в один текст абзаца
        text = "".join(run.text for run in para.runs)
        if text.strip():
            paragraphs.append(text.strip())
    return "\n".join(paragraphs)


def _table_to_html(table) -> str:
    """
    Конвертирует таблицу PPTX в HTML <table>.

    Args:
        table: объект pptx.table.Table

    Returns:
        str — HTML-таблица
    """
    rows_html = []
    for row in table.rows:
        cells_html = []
        for cell in row.cells:
            # Текст ячейки: все абзацы через пробел
            cell_text = " ".join(
                "".join(run.text for run in para.runs)
                for para in cell.text_frame.paragraphs
            ).strip()
            cells_html.append(f"  <td>{cell_text}</td>")
        rows_html.append("<tr>\n" + "\n".join(cells_html) + "\n</tr>")

    if not rows_html:
        return ""

    return "<table>\n" + "\n".join(rows_html) + "\n</table>"
