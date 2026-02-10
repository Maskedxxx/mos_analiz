#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Парсер листа «выпадающий список» из XLSX-файла.

Извлекает справочник допустимых единиц измерения по категориям показателей.
Результат → dropdown_units.json
"""

import json
from pathlib import Path

import openpyxl


def _find_sheet(wb, keyword: str):
    """Поиск листа по подстроке в имени."""
    for name in wb.sheetnames:
        if keyword.lower() in name.lower():
            return wb[name]
    return None


def parse(xlsx_path: Path, output_dir: Path) -> dict:
    """
    Парсит лист 'выпадающий список' — справочник допустимых единиц измерения.

    Строка 1 — заголовки (названия категорий): A, B, C, D.
    Строки 2+ — значения единиц измерения по столбцам.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = _find_sheet(wb, "выпадающий список")

    if ws is None:
        wb.close()
        result = {"categories": {}}
        output_file = output_dir / "dropdown_units.json"
        with open(output_file, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        return result

    categories = {}

    # Строка 1 — заголовки категорий (столбцы A-D)
    columns = ["A", "B", "C", "D"]
    headers = []
    for col in columns:
        val = ws[f"{col}1"].value
        header = str(val).strip() if val else None
        headers.append(header)

    # Строки 2+ — значения
    for col, header in zip(columns, headers):
        if header is None:
            continue

        units = []
        for row in range(2, ws.max_row + 1):
            val = ws[f"{col}{row}"].value
            if val is not None:
                unit = str(val).strip()
                if unit:
                    units.append(unit)

        categories[header] = units

    wb.close()

    result = {"categories": categories}

    # Сохраняем результат
    output_file = output_dir / "dropdown_units.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    return result
