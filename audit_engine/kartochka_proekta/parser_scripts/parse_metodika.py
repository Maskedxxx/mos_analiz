#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Парсер листа «Методика расчета» из XLSX-файла.

Извлекает показатели с их единицами измерения, способом расчёта и источником данных.
Результат → metodika.json
"""

import json
import re
from pathlib import Path

import openpyxl


def _find_sheet(wb, keyword: str, exclude: str = "ШАБЛОН"):
    """Поиск листа по подстроке в имени, исключая шаблоны."""
    for name in wb.sheetnames:
        if keyword.lower() in name.lower() and exclude.upper() not in name.upper():
            return wb[name]
    return None


def _cell_str(ws, coord: str) -> str:
    """Получить строковое значение ячейки."""
    val = ws[coord].value
    if val is None:
        return ""
    return str(val).strip()


def _is_field_label(text: str, label: str) -> bool:
    """Проверяет, начинается ли текст с метки поля (напр. 'Единицы измерения:')."""
    return text.lower().startswith(label.lower())


def parse(xlsx_path: Path, output_dir: Path) -> dict:
    """
    Парсит лист 'Методика расчета' из XLSX-файла.

    Алгоритм: итерируем строки, ищем названия показателей (не являющиеся метками полей),
    затем читаем 3 строки ниже: единицы, способ расчёта, источник данных.
    """
    output_dir.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)
    ws = _find_sheet(wb, "Методика расчет")

    if ws is None:
        wb.close()
        result = {"meta": {"sheet": None}, "header": {}, "indicators": []}
        output_file = output_dir / "metodika.json"
        with open(output_file, "w", encoding="utf-8") as f:
            json.dump(result, f, ensure_ascii=False, indent=2)
        return result

    result = {
        "meta": {"sheet": ws.title},
        "header": {
            "org_name": _cell_str(ws, "B2"),
            "project_name": _cell_str(ws, "B4"),
        },
        "indicators": [],
    }

    # Метки полей — эти строки НЕ являются названиями показателей
    field_labels = ["единицы измерения", "способ расчет", "источник данных", "методика расчет"]
    # Заголовки секции — тоже пропускаем
    skip_keywords = ["карточка проекта", "методика расчет"]

    row = 1
    max_row = ws.max_row

    while row <= max_row:
        b_val = _cell_str(ws, f"B{row}")

        if not b_val:
            row += 1
            continue

        # Пропускаем заголовки и метки
        b_lower = b_val.lower()
        is_label = any(b_lower.startswith(lbl) for lbl in field_labels)
        is_skip = any(kw in b_lower for kw in skip_keywords)

        if is_label or is_skip:
            row += 1
            continue

        # Проверяем — это название показателя?
        # Следующая строка должна быть "Единицы измерения:"
        next_b = _cell_str(ws, f"B{row + 1}") if row + 1 <= max_row else ""
        if _is_field_label(next_b, "единицы измерения"):
            # Это название показателя
            indicator_name = re.sub(r':$', '', b_val).strip()
            unit = _cell_str(ws, f"F{row + 1}")
            calc_method = _cell_str(ws, f"F{row + 2}") if row + 2 <= max_row else ""
            data_source = _cell_str(ws, f"F{row + 3}") if row + 3 <= max_row else ""

            result["indicators"].append({
                "name": indicator_name,
                "unit": unit if unit else None,
                "calc_method": calc_method if calc_method else None,
                "data_source": data_source if data_source else None,
                "row_start": row,
            })

            row += 4  # Перепрыгиваем блок показателя
            continue

        row += 1

    wb.close()

    # Сохраняем результат
    output_file = output_dir / "metodika.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)

    return result
