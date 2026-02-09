#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Парсер для документа "График обхода ОМ" (.xlsx).

Извлекает 2 чанка:
- шапка: наименование компании + ссылка на приказ
- фио_должности: список ФИО и должностей из таблицы

Использование:
    from audit_engine.parsers.grafik_obhod import parse
    chunks = parse("path/to/grafik.xlsx")
"""

import re
from openpyxl import load_workbook


def parse(xlsx_path: str) -> dict:
    """
    Парсит документ "График обхода ОМ".

    Args:
        xlsx_path: путь к .xlsx файлу

    Returns:
        dict с чанками: шапка, фио_должности (БЕЗ префиксов)
    """
    wb = load_workbook(xlsx_path, data_only=True)

    # Ищем нужный лист: "График" или "Шаблон"
    sheet_name = None
    for name in wb.sheetnames:
        if name.lower() in ["график", "шаблон"]:
            sheet_name = name
            break

    # Fallback — первый лист
    if not sheet_name:
        sheet_name = wb.sheetnames[0]

    ws = wb[sheet_name]

    # Собираем данные
    company_name = ""
    prikaz_ref = ""
    employees = []

    # Проходим по первым 50 строкам
    for row_num in range(1, min(51, ws.max_row + 1)):
        row_values = []
        for col_num in range(1, min(35, ws.max_column + 1)):
            cell = ws.cell(row=row_num, column=col_num)
            if cell.value:
                row_values.append(str(cell.value).strip())

        # Ищем наименование компании (ООО/ЗАО/АО/ПАО/ИП)
        for val in row_values:
            if re.search(r'\b(ООО|ЗАО|АО|ПАО|ИП)\b', val) and not company_name:
                company_name = val

        # Ищем "Приложение №1 к Приказу..."
        for val in row_values:
            if "приложение" in val.lower() and "приказ" in val.lower():
                prikaz_ref = val.replace("\n", " ").strip()

        # Ищем строки с ФИО и должностью (первая ячейка = ФИО, вторая = должность)
        if len(row_values) >= 2:
            first_val = row_values[0]
            second_val = row_values[1]

            # Пропускаем заголовки и служебные строки
            if first_val.lower() not in ["фио", "график", ""] and \
               "график обхода" not in first_val.lower() and \
               "приложение" not in first_val.lower():

                # Паттерн ФИО: "Фамилия И.О." или "Фамилия Имя Отчество"
                if re.search(r'[А-ЯЁ][а-яё]+\s+[А-ЯЁ]\.?[А-ЯЁа-яё]?\.?', first_val) or \
                   re.search(r'[А-ЯЁ][а-яё]+\s+[А-ЯЁ][а-яё]+', first_val):
                    # Второй столбец — должность (не заголовок)
                    if second_val and second_val.lower() != "должность":
                        employees.append(f"{first_val} — {second_val}")

    # Формируем чанки
    chunks = {
        "шапка": f"{company_name}\n{prikaz_ref}".strip(),
        "фио_должности": "\n".join(dict.fromkeys(employees)),  # дедупликация
    }

    return chunks
