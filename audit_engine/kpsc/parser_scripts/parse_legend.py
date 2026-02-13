#!/usr/bin/env python3
"""
Парсер листа "Условные обозначения": связывает пиктограммы (drawing3.xml) с расшифровками в колонке B.
Выводит список entries: row, text, picture (rel id, target, anchor bbox).

Использует fuzzy-поиск листа через sheet_finder.
"""
import argparse
import json
import posixpath
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path
from typing import List, Dict, Any, Optional

from openpyxl import load_workbook

from audit_engine.kpsc.sheet_finder import find_sheet

NS = {
    "wb": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}


def read_sheet_drawing(xlsx: Path, sheet_name: str):
    with zipfile.ZipFile(xlsx) as z:
        wb_xml = ET.fromstring(z.read("xl/workbook.xml"))
        wb_rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        rel_map = {rel.attrib["Id"]: rel.attrib["Target"] for rel in wb_rels}
        sheet_target = None
        for sheet in wb_xml.find("wb:sheets", NS):
            if sheet.attrib["name"] == sheet_name:
                rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
                sheet_target = rel_map[rid]
                break
        rel_path = f"xl/worksheets/_rels/{posixpath.basename(sheet_target)}.rels"
        rels = ET.fromstring(z.read(rel_path))
        drawing_path = None
        for rel in rels:
            if rel.attrib.get("Type") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing":
                drawing_path = posixpath.normpath(posixpath.join("xl/worksheets", rel.attrib["Target"]))
        d_rels_path = posixpath.join(posixpath.dirname(drawing_path), "_rels", posixpath.basename(drawing_path) + ".rels")
        d_rels = ET.fromstring(z.read(d_rels_path)) if d_rels_path in z.namelist() else None
        rel_pic = {rel.attrib["Id"]: rel.attrib["Target"] for rel in d_rels} if d_rels is not None else {}
        drawing_root = ET.fromstring(z.read(drawing_path))
    return drawing_root, rel_pic


def parse_pictures(drawing_root: ET.Element, rel_pic: Dict[str, str]):
    pics = []
    for anc in drawing_root.findall("./", NS):
        pic_el = anc.find("xdr:pic", NS)
        if pic_el is None:
            continue
        pf = anc.find("xdr:from", NS)
        pt = anc.find("xdr:to", NS)
        fcol = int(pf.find("xdr:col", NS).text)
        frow = int(pf.find("xdr:row", NS).text)
        tcol = int(pt.find("xdr:col", NS).text)
        trow = int(pt.find("xdr:row", NS).text)
        blip = pic_el.find(".//a:blip", NS)
        rid = blip.attrib.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed") if blip is not None else None
        pics.append(
            {
                "row": frow + 1,
                "col": fcol + 1,
                "bbox": {
                    "from": {"row": frow + 1, "col": fcol + 1},
                    "to": {"row": trow + 1, "col": tcol + 1},
                },
                "rel_id": rid,
                "target": rel_pic.get(rid),
            }
        )
    return pics


def collect_text(ws):
    texts = []
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row=row, column=2).value  # column B
        if val not in (None, ""):
            texts.append({"row": row, "col": 2, "text": val})
    return texts


def match_pics(texts: List[Dict[str, Any]], pics: List[Dict[str, Any]]):
    matched = []
    # Уберём заголовок (обычно строка 2) и сопоставим по порядку: картинки и тексты отсортированы по строке.
    texts_order = [t for t in sorted(texts, key=lambda x: x["row"]) if t["text"] != "Расшифровка или пояснение"]
    pics_order = sorted(pics, key=lambda x: x["row"])
    n = min(len(texts_order), len(pics_order))
    for i, t in enumerate(texts_order):
        pic = pics_order[i] if i < len(pics_order) else None
        matched.append(
            {
                "row": t["row"],
                "text": t["text"],
                "picture": pic,
            }
        )
    return matched


def build_payload(xlsx: Path, sheet_name: Optional[str] = None):
    wb = load_workbook(xlsx, data_only=True)

    # Поиск листа через sheet_finder
    if sheet_name:
        ws = wb[sheet_name]
        actual_sheet = sheet_name
    else:
        ws = find_sheet(wb, keywords=["условн", "обозн"])
        if ws is None:
            # Лист не найден — возвращаем пустой результат (не у всех компаний он есть)
            return {
                "meta": {"workbook": str(xlsx), "sheet": None},
                "pictures": [],
                "entries": [],
            }
        actual_sheet = ws.title

    drawing_root, rel_pic = read_sheet_drawing(xlsx, actual_sheet)
    pics = parse_pictures(drawing_root, rel_pic)
    texts = collect_text(ws)
    entries = match_pics(texts, pics)
    return {
        "meta": {"workbook": str(xlsx), "sheet": actual_sheet},
        "pictures": pics,
        "entries": entries,
    }


def parse(xlsx_path: Path, output_dir: Path) -> dict:
    """Парсит лист 'Условные обозначения' и сохраняет в legend_v2.json."""
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = build_payload(xlsx_path)
    output_path = output_dir / "legend_v2.json"
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def main():
    ap = argparse.ArgumentParser(description="Парсер листа 'Условные обозначения'")
    ap.add_argument("-i", "--input", required=True)
    ap.add_argument("-s", "--sheet", default=None)
    ap.add_argument("-o", "--output")
    args = ap.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
