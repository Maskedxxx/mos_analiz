#!/usr/bin/env python3
"""
Парсер стартовой таблицы на листе "ПА-1" (до диаграммы).
Определяет границы по первой строке с заголовком (содержит «ИТОГО»),
захватывает все столбцы/строки с данными до первой полностью пустой строки.
Учитывает merged-ячейки, возвращает координаты, значения, диапазоны.
"""

import argparse
import json
from pathlib import Path
from typing import Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries


def find_header_row(ws) -> Optional[int]:
    """Ищем строку, где в заголовках встречается слово 'итого'."""
    for r in range(1, ws.max_row + 1):
        if any(isinstance(c.value, str) and "итого" in c.value.lower() for c in ws[r]):
            return r
    return None


def compute_col_bounds(ws, header_row: int) -> Tuple[int, int]:
    """
    Берём минимальный/максимальный столбцы с данными в строках заголовка
    и двух строках ниже (чтобы захватить пустой заголовок первого столбца,
    но заполненные значения 'Замер 1' и т.п.).
    """
    left = None
    right = 0
    last_row = min(ws.max_row, header_row + 2)
    for r in range(header_row, last_row + 1):
        for c in range(1, ws.max_column + 1):
            if ws.cell(row=r, column=c).value not in (None, ""):
                left = c if left is None else min(left, c)
                right = max(right, c)
    if left is None:
        left = 1
        right = 1
    return left, right


def find_bottom_row(ws, header_row: int, left_col: int, right_col: int) -> int:
    """Первая строка, полностью пустая в пределах таблицы, завершает данные."""
    r = header_row + 1
    while r <= ws.max_row:
        if all(ws.cell(row=r, column=c).value in (None, "") for c in range(left_col, right_col + 1)):
            return r - 1
        r += 1
    return ws.max_row


def build_merged_lookup(ws):
    lookup = {}
    for m in ws.merged_cells.ranges:
        coord = m.coord
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                lookup[(r, c)] = coord
    return lookup


def extract_table(ws, top_row: int, bottom_row: int, left_col: int, right_col: int):
    merged_lookup = build_merged_lookup(ws)
    rows = []
    for r in range(top_row, bottom_row + 1):
        row_cells = []
        for c in range(left_col, right_col + 1):
            coord = f"{get_column_letter(c)}{r}"
            merge_range = merged_lookup.get((r, c))
            if merge_range:
                min_col_m, min_row_m, max_col_m, max_row_m = range_boundaries(merge_range)
                anchor = (r == min_row_m and c == min_col_m)
                value = ws.cell(row=min_row_m, column=min_col_m).value
            else:
                anchor = True
                value = ws.cell(row=r, column=c).value
            row_cells.append(
                {
                    "coord": coord,
                    "row": r,
                    "col": c,
                    "value": value,
                    "merge_range": merge_range,
                    "merge_anchor": anchor if merge_range else False,
                }
            )
        rows.append({"row": r, "cells": row_cells})
    return rows


def build_payload(xlsx_path: Path, sheet: str):
    wb = load_workbook(xlsx_path, data_only=True)
    ws = wb[sheet]

    header_row = find_header_row(ws)
    if not header_row:
        raise SystemExit("Не найден заголовок с 'ИТОГО' на листе ПА-1")

    left_col, right_col = compute_col_bounds(ws, header_row)
    bottom_row = find_bottom_row(ws, header_row, left_col, right_col)
    rows = extract_table(ws, header_row, bottom_row, left_col, right_col)

    return {
        "meta": {"workbook": str(xlsx_path), "sheet": sheet, "header_row": header_row},
        "bounds": {
            "top_row": header_row,
            "bottom_row": bottom_row,
            "left_col": left_col,
            "right_col": right_col,
            "left_letter": get_column_letter(left_col),
            "right_letter": get_column_letter(right_col),
            "height": bottom_row - header_row + 1,
            "width": right_col - left_col + 1,
        },
        "rows": rows,
    }


def parse(xlsx_path: Path, output_dir: Path) -> dict:
    """Парсит стартовую таблицу на листе 'ПА-1' и сохраняет в pa1_table_v1.json."""
    payload = build_payload(xlsx_path, "ПА-1")
    output_path = output_dir / "pa1_table_v1.json"
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return payload


def main():
    ap = argparse.ArgumentParser(description="Парсер стартовой таблицы на листе 'ПА-1'")
    ap.add_argument("-i", "--input", required=True)
    ap.add_argument("-s", "--sheet", default="ПА-1")
    ap.add_argument("-o", "--output")
    args = ap.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
