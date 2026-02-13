#!/usr/bin/env python3
"""
Парсер таблицы "1. Определение показателей потока" на листе КПСЦ.
Выделяет границы таблицы (по заголовку 1.), учитывает merged-ячейки, возвращает
прямоугольную выборку ячеек с координатами и метаданными.

Использует fuzzy-поиск листа через sheet_finder.
"""

import argparse
import json
from pathlib import Path
from typing import Optional, Tuple

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from audit_engine.kpsc.sheet_finder import find_sheet


def find_section_anchor(ws, phrase: str) -> Optional[Tuple[int, int]]:
    """
    Ищем начало секции таблицы КПСЦ с показателями потока.

    Стратегия (по приоритету):
    1. "1. Определение показателей потока" — стандартный формат (biznes_otel)
    2. "Название этапа процесса" — плоская таблица (mapper, sodex)
    3. "Показатель" рядом с "Ед. измерения" — сводная таблица (rotosnab, ruslet)
    """
    # Стратегия 1: ищем "1. Определение показателей потока"
    phrase_low = phrase.lower()
    for row in ws.iter_rows():
        for cell in row:
            val = cell.value
            if isinstance(val, str) and phrase_low in val.lower():
                if val.strip().startswith("1"):
                    return cell.row, cell.column

    # Стратегия 2: ищем "Название этапа процесса" как заголовок таблицы
    for row in ws.iter_rows(max_row=min(15, ws.max_row)):
        for cell in row:
            val = cell.value
            if isinstance(val, str) and "название этапа" in val.lower():
                # Возвращаем row-1 как «anchor», потому что find_header_row() возьмёт anchor+1
                return max(1, cell.row - 1), cell.column

    # Стратегия 3: ищем строку с "Показатель" + "Ед." в той же строке
    for r in range(1, ws.max_row + 1):
        row_texts = []
        first_col = None
        for c in range(1, min(10, (ws.max_column or 10) + 1)):
            v = ws.cell(r, c).value
            if isinstance(v, str):
                row_texts.append(v.lower())
                if first_col is None:
                    first_col = c
        joined = " ".join(row_texts)
        if "показатель" in joined and "ед." in joined:
            return max(1, r - 1), first_col or 1

    return None


def find_header_row(ws, anchor_row: int) -> int:
    # ищем первую непустую строку ниже заголовка
    for r in range(anchor_row + 1, anchor_row + 5):
        if any(c.value not in (None, "") for c in ws[r]):
            return r
    return anchor_row + 1


def find_bottom_row(ws, header_row: int, section_col: int) -> int:
    """Ищем конец таблицы: первую строку, где в колонке section_col начинается следующая секция ("2.")."""
    r = header_row + 1
    while r <= ws.max_row:
        cell_val = ws.cell(row=r, column=section_col).value
        if isinstance(cell_val, str) and cell_val.strip().startswith("2."):
            return r - 1
        # стоп если встретили полностью пустую строку после данных
        row_vals = [c.value for c in ws[r]]
        if all(v in (None, "") for v in row_vals) and r > header_row + 2:
            return r - 1
        r += 1
    return ws.max_row


def compute_col_bounds(ws, header_row: int) -> Tuple[int, int]:
    """
    Границы по строке заголовков:
      left  — первый непустой столбец,
      right — последний столбец в заголовке, где текст содержит 'ед. измерения'
              (регистр не важен); если не найдено, берем правый непустой.
    """
    cells = list(ws[header_row])
    left = None
    right = None
    right_by_text = None
    for c in cells:
        val = c.value
        if val not in (None, ""):
            if left is None:
                left = c.column
            if isinstance(val, str) and "ед. измерения" in val.lower():
                right_by_text = c.column
            right = c.column
    if right_by_text:
        right = right_by_text
    if left is None:
        left = 1
    if right is None:
        right = left
    return left, right


def build_merged_lookup(ws):
    lookup = {}
    for merge in ws.merged_cells.ranges:
        coord = merge.coord
        min_row, min_col, max_row, max_col = merge.min_row, merge.min_col, merge.max_row, merge.max_col
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                lookup[(r, c)] = coord
    return lookup


def extract_table(ws, top_row: int, bottom_row: int, left_col: int, right_col: int):
    merged_lookup = build_merged_lookup(ws)
    rows = []
    for r in range(top_row, bottom_row + 1):
        row_cells = []
        for c in range(left_col, right_col + 1):
            cell = ws.cell(row=r, column=c)
            coord = f"{get_column_letter(c)}{r}"
            merge_range = merged_lookup.get((r, c))
            # Значение берём из верхнего-левого якоря merged, если оно есть
            if merge_range:
                min_col_m, min_row_m, max_col_m, max_row_m = range_boundaries(merge_range)
                anchor = (r == min_row_m and c == min_col_m)
                value = ws.cell(row=min_row_m, column=min_col_m).value
            else:
                anchor = True
                value = cell.value
            row_cells.append({
                "coord": coord,
                "col": c,
                "row": r,
                "value": value,
                "merge_range": merge_range,
                "merge_anchor": anchor if merge_range else False,
            })
        rows.append({"row": r, "cells": row_cells})
    return rows


def build_payload(xlsx_path: Path, sheet_name: Optional[str] = None):
    wb = load_workbook(xlsx_path, data_only=True)

    # Поиск листа через sheet_finder
    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = find_sheet(
            wb,
            keywords=["кпсц"],
            exclude_keywords=["спагетти", "укрупн", "оцифровк"],
            prefer_keywords=["тс", "текущ"],
        )
        if ws is None:
            raise ValueError(f"Лист КПСЦ не найден среди {wb.sheetnames}")

    actual_sheet = ws.title

    anchor = find_section_anchor(ws, "определение показателей потока")
    if not anchor:
        # Нет секции "1. Определение показателей потока" — возвращаем пустой результат
        return {
            "meta": {"workbook": str(xlsx_path), "sheet": actual_sheet, "section_title_cell": None},
            "bounds": None,
            "rows": [],
        }
    anchor_row, anchor_col = anchor

    header_row = find_header_row(ws, anchor_row)
    bottom_row = find_bottom_row(ws, header_row, anchor_col)
    left_col, right_col = compute_col_bounds(ws, header_row)

    table_rows = extract_table(ws, header_row, bottom_row, left_col, right_col)

    payload = {
        "meta": {
            "workbook": str(xlsx_path),
            "sheet": actual_sheet,
            "section_title_cell": f"{get_column_letter(anchor_col)}{anchor_row}",
        },
        "bounds": {
            "top_row": header_row,
            "bottom_row": bottom_row,
            "left_col": left_col,
            "right_col": right_col,
            "left_letter": get_column_letter(left_col),
            "right_letter": get_column_letter(right_col),
            "height": bottom_row - header_row + 1,
            "width": right_col - left_col + 1,
        },
        "rows": table_rows,
    }
    return payload


def parse(xlsx_path: Path, output_dir: Path) -> dict:
    """Парсит таблицу '1. Определение показателей потока' и сохраняет в kpsc_table1_v2.json."""
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = build_payload(xlsx_path)
    output_path = output_dir / "kpsc_table1_v2.json"
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return payload


def main():
    parser = argparse.ArgumentParser(description="Парсер таблицы '1. Определение показателей потока'")
    parser.add_argument("-i", "--input", required=True, help="XLSX файл")
    parser.add_argument("-s", "--sheet", default=None, help="Лист (default: auto)")
    parser.add_argument("-o", "--output", default=None, help="JSON файл вывода (stdout если не указан)")
    args = parser.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
