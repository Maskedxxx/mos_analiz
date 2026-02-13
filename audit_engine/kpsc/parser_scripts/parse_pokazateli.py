#!/usr/bin/env python3
"""
Парсер таблицы "Текущие показатели потока" на листе "Показатели".
Находит заголовок, снимает прямоугольник таблицы, учитывает merged-ячейки.

Использует fuzzy-поиск листа через sheet_finder.
"""
import argparse
import json
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from audit_engine.kpsc.sheet_finder import find_sheet

TITLE_PHRASE = "текущие показатели потока"


def find_title(ws) -> int:
    for r in range(1, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and TITLE_PHRASE in v.lower():
                return r
    raise SystemExit("Title 'Текущие показатели потока' not found")


def find_header_row(ws, title_row: int) -> int:
    for r in range(title_row + 1, title_row + 5):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in range(1, ws.max_column + 1)):
            return r
    return title_row + 1


def find_bottom_row(ws, header_row: int) -> int:
    r = header_row + 1
    while r <= ws.max_row:
        if all(ws.cell(row=r, column=c).value in (None, "") for c in range(1, ws.max_column + 1)):
            return r - 1
        r += 1
    return ws.max_row


def compute_col_bounds(ws, header_row: int) -> Tuple[int, int]:
    """
    Граница по строке заголовков: берём первую непустую ячейку и продолжаем вправо,
    пока идут непустые. Если встречаем пустую колонку после начала таблицы — там обрываем.
    Это отсечёт служебные столбцы справа (I,J...).
    """
    left = None
    right = None
    started = False
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v not in (None, ""):
            if left is None:
                left = c
            right = c
            started = True
        else:
            if started:
                break
    if left is None:
        left = 1
        right = 1
    return left, right


def build_merged_lookup(ws):
    lookup = {}
    for m in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = m.min_col, m.min_row, m.max_col, m.max_row
        for r in range(min_row, max_row + 1):
            for c in range(min_col, max_col + 1):
                lookup[(r, c)] = m.coord
    return lookup


def extract_table(ws, top_row: int, bottom_row: int, left_col: int, right_col: int):
    merged_lookup = build_merged_lookup(ws)
    rows = []
    for r in range(top_row, bottom_row + 1):
        row_cells = []
        for c in range(left_col, right_col + 1):
            coord = f"{get_column_letter(c)}{r}"
            merge_range = merged_lookup.get((r, c))
            if merge_range:
                min_col_m, min_row_m, max_col_m, max_row_m = range_boundaries(merge_range)
                anchor = (r == min_row_m and c == min_col_m)
                value = ws.cell(row=min_row_m, column=min_col_m).value
            else:
                anchor = True
                value = ws.cell(row=r, column=c).value
            row_cells.append(
                {
                    "coord": coord,
                    "row": r,
                    "col": c,
                    "value": value,
                    "merge_range": merge_range,
                    "merge_anchor": anchor if merge_range else False,
                }
            )
        rows.append({"row": r, "cells": row_cells})
    return rows


def build_payload(xlsx: Path, sheet_name: Optional[str] = None):
    wb = load_workbook(xlsx, data_only=True)

    # Поиск листа через sheet_finder
    if sheet_name:
        ws = wb[sheet_name]
        actual_sheet = sheet_name
    else:
        ws = find_sheet(wb, keywords=["показател"])
        if ws is None:
            return {
                "meta": {"workbook": str(xlsx), "sheet": None, "title_row": None},
                "bounds": None,
                "rows": [],
            }
        actual_sheet = ws.title

    # Ищем заголовок «Текущие показатели потока»
    title_row = None
    try:
        title_row = find_title(ws)
    except SystemExit:
        # Заголовок не найден — возвращаем пустой результат
        return {
            "meta": {"workbook": str(xlsx), "sheet": actual_sheet, "title_row": None},
            "bounds": None,
            "rows": [],
        }

    header_row = find_header_row(ws, title_row)
    bottom_row = find_bottom_row(ws, header_row)
    left_col, right_col = compute_col_bounds(ws, header_row)
    table_rows = extract_table(ws, header_row, bottom_row, left_col, right_col)
    return {
        "meta": {"workbook": str(xlsx), "sheet": actual_sheet, "title_row": title_row},
        "bounds": {
            "top_row": header_row,
            "bottom_row": bottom_row,
            "left_col": left_col,
            "right_col": right_col,
        },
        "rows": table_rows,
    }


def parse(xlsx_path: Path, output_dir: Path) -> dict:
    """Парсит 'Текущие показатели потока' и сохраняет в pokazateli_v3.json."""
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = build_payload(xlsx_path)
    output_path = output_dir / "pokazateli_v3.json"
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return payload


def main():
    ap = argparse.ArgumentParser(description="Парсер 'Текущие показатели потока'")
    ap.add_argument("-i", "--input", required=True)
    ap.add_argument("-s", "--sheet", default=None)
    ap.add_argument("-o", "--output")
    args = ap.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
