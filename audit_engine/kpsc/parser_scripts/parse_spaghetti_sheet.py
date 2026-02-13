#!/usr/bin/env python3
"""
Парсер листа "Диаграмма Спагетти" (без самой диаграммы).
Снимает:
  - верхние текстовые/служебные строки до таблицы;
  - таблицу путей перемещений начиная со строки с заголовком ("Путь", ...),
    с учётом merged-ячеек.
Диаграмма-графика игнорируется.

Использует fuzzy-поиск листа через sheet_finder.
"""

import argparse
import json
from pathlib import Path
from typing import Optional, Tuple, Dict, Any, List

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, range_boundaries

from audit_engine.kpsc.sheet_finder import find_sheet


HEADER_PHRASES = ["шаги процесса", "путь", "перемещени"]


def find_header_row(ws) -> Optional[int]:
    """Ищем строку-заголовок таблицы перемещений по ключевым словам."""
    for r in range(1, min(20, ws.max_row + 1)):
        row_vals = [c.value for c in ws[r]]
        for v in row_vals:
            if isinstance(v, str):
                v_low = v.lower()
                if any(phrase in v_low for phrase in HEADER_PHRASES):
                    return r
    return None


def compute_col_bounds(ws, header_row: int) -> Tuple[int, int]:
    left = None
    right = 0
    # левая граница — первый непустой столбец заголовка
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=c).value
        if v not in (None, ""):
            left = c
            break
    if left is None:
        left = 1
    # правая граница — максимальный столбец с данными начиная с строки заголовка
    for r in range(header_row, ws.max_row + 1):
        for c in range(left, ws.max_column + 1):
            if ws.cell(row=r, column=c).value not in (None, ""):
                right = max(right, c)
    if right < left:
        right = left
    return left, right


def find_bottom_row(ws, header_row: int, left_col: int, right_col: int) -> int:
    r = header_row + 1
    while r <= ws.max_row:
        if all(ws.cell(row=r, column=c).value in (None, "") for c in range(left_col, right_col + 1)):
            return r - 1
        r += 1
    return ws.max_row


def build_merged_lookup(ws):
    lookup = {}
    for m in ws.merged_cells.ranges:
        coord = m.coord
        for r in range(m.min_row, m.max_row + 1):
            for c in range(m.min_col, m.max_col + 1):
                lookup[(r, c)] = coord
    return lookup


def extract_table(ws, top_row: int, bottom_row: int, left_col: int, right_col: int):
    merged_lookup = build_merged_lookup(ws)
    rows = []
    for r in range(top_row, bottom_row + 1):
        row_cells = []
        for c in range(left_col, right_col + 1):
            coord = f"{get_column_letter(c)}{r}"
            merge_range = merged_lookup.get((r, c))
            if merge_range:
                min_col_m, min_row_m, max_col_m, max_row_m = range_boundaries(merge_range)
                anchor = (r == min_row_m and c == min_col_m)
                value = ws.cell(row=min_row_m, column=min_col_m).value
            else:
                anchor = True
                value = ws.cell(row=r, column=c).value
            row_cells.append(
                {
                    "coord": coord,
                    "row": r,
                    "col": c,
                    "value": value,
                    "merge_range": merge_range,
                    "merge_anchor": anchor if merge_range else False,
                }
            )
        rows.append({"row": r, "cells": row_cells})
    return rows


def extract_pre_table(ws, header_row: int):
    cells = []
    for r in range(1, header_row):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if v not in (None, ""):
                cells.append(
                    {
                        "coord": f"{get_column_letter(c)}{r}",
                        "row": r,
                        "col": c,
                        "value": v,
                    }
                )
    return cells


def build_payload(xlsx_path: Path, sheet_name: Optional[str] = None):
    wb = load_workbook(xlsx_path, data_only=True)

    # Поиск листа через sheet_finder
    if sheet_name:
        ws = wb[sheet_name]
        actual_sheet = sheet_name
    else:
        ws = find_sheet(
            wb,
            keywords=["спагетти"],
            exclude_keywords=["пробл", "улучш", "перечень"],
        )
        if ws is None:
            return {
                "meta": {"workbook": str(xlsx_path), "sheet": None, "header_row": None},
                "bounds": None,
                "pre_table_cells": [],
                "rows": [],
            }
        actual_sheet = ws.title

    header_row = find_header_row(ws)
    if not header_row:
        return {
            "meta": {"workbook": str(xlsx_path), "sheet": actual_sheet, "header_row": None},
            "bounds": None,
            "pre_table_cells": [],
            "rows": [],
        }

    left_col, right_col = compute_col_bounds(ws, header_row)
    bottom_row = find_bottom_row(ws, header_row, left_col, right_col)

    pre_table = extract_pre_table(ws, header_row)
    table_rows = extract_table(ws, header_row, bottom_row, left_col, right_col)

    return {
        "meta": {"workbook": str(xlsx_path), "sheet": actual_sheet, "header_row": header_row},
        "bounds": {
            "top_row": header_row,
            "bottom_row": bottom_row,
            "left_col": left_col,
            "right_col": right_col,
            "left_letter": get_column_letter(left_col),
            "right_letter": get_column_letter(right_col),
            "height": bottom_row - header_row + 1,
            "width": right_col - left_col + 1,
        },
        "pre_table_cells": pre_table,
        "rows": table_rows,
    }


def parse(xlsx_path: Path, output_dir: Path) -> dict:
    """Парсит лист 'Диаграмма Спагетти' и сохраняет в spaghetti_sheet_v2.json."""
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = build_payload(xlsx_path)
    output_path = output_dir / "spaghetti_sheet_v2.json"
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return payload


def main():
    ap = argparse.ArgumentParser(description="Парсер листа 'Диаграмма Спагетти' (без диаграммы)")
    ap.add_argument("-i", "--input", required=True)
    ap.add_argument("-s", "--sheet", default=None)
    ap.add_argument("-o", "--output")
    args = ap.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
