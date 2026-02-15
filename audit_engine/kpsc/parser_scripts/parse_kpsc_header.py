#!/usr/bin/env python3
"""
Парсер верхнего блока листа КПСЦ.

Извлекает ключевые поля (title, flow_name, responsible, date_developed и т.д.)
через динамический поиск лейблов в ячейках, а не через захардкоженные координаты.
Использует fuzzy-поиск листа через sheet_finder.
"""
import argparse
import json
from pathlib import Path
from typing import Dict, Any, List, Optional
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from audit_engine.kpsc.sheet_finder import find_sheet

# Максимум строк для сканирования заголовочного блока
HEADER_SCAN_MAX_ROW = 15


def collect_cells(ws, max_row: int = HEADER_SCAN_MAX_ROW, max_col: Optional[int] = None):
    """Собираем все непустые ячейки в заголовочном регионе."""
    if max_col is None:
        max_col = ws.max_column or 50
    cells = []
    comments = []
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            val = cell.value
            if val not in (None, ""):
                cells.append({
                    "coord": f"{get_column_letter(c)}{r}",
                    "row": r,
                    "col": c,
                    "value": val,
                })
            if cell.comment:
                comments.append({
                    "coord": f"{get_column_letter(c)}{r}",
                    "row": r,
                    "col": c,
                    "author": cell.comment.author,
                    "text": cell.comment.text,
                })
    return cells, comments


def collect_merged(ws, max_row: int = HEADER_SCAN_MAX_ROW, max_col: Optional[int] = None):
    """Собираем merged-диапазоны в заголовочном регионе."""
    if max_col is None:
        max_col = ws.max_column or 50
    merged = []
    for m in ws.merged_cells.ranges:
        if m.min_row <= max_row and m.min_col <= max_col:
            merged.append({
                "coord": m.coord,
                "min_row": m.min_row,
                "max_row": m.max_row,
                "min_col": m.min_col,
                "max_col": m.max_col,
            })
    return merged


def _find_label_value(ws, label_keywords: List[str], max_row: int = HEADER_SCAN_MAX_ROW) -> Optional[Any]:
    """
    Динамический поиск значения по лейблу.

    Ищем ячейку, содержащую одно из label_keywords, затем берём значение
    из ближайшей непустой ячейки справа в той же строке.

    Если лейбл содержит значение inline (например "Наименование потока: Производство..."),
    извлекаем часть после двоеточия.
    """
    for r in range(1, max_row + 1):
        for c in range(1, min(5, (ws.max_column or 5) + 1)):
            v = ws.cell(row=r, column=c).value
            if not isinstance(v, str):
                continue
            v_lower = v.strip().lower()

            for kw in label_keywords:
                if kw not in v_lower:
                    continue

                # Случай 1: inline-значение после двоеточия
                if ":" in v:
                    parts = v.split(":", 1)
                    inline_val = parts[1].strip()
                    if inline_val:
                        return inline_val

                # Случай 2: значение в соседней ячейке справа
                for vc in range(c + 1, min(c + 4, (ws.max_column or c) + 1)):
                    val = ws.cell(row=r, column=vc).value
                    if val not in (None, ""):
                        return val

                # Не нашли значение, но лейбл есть — вернём None (не продолжаем поиск)
                return None
    return None


def _find_title(ws, max_row: int = HEADER_SCAN_MAX_ROW) -> Optional[str]:
    """
    Извлекает заголовок карты потока.

    Ищем в первых строках длинную строку, содержащую ключевые слова
    типа "карта потока", "КПСЦ", "текущее состояние".
    """
    title_keywords = ["карта потока", "кпсц", "текущее состояние", "наименование потока"]
    for r in range(1, min(4, max_row + 1)):
        for c in range(1, 4):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and len(v.strip()) > 15:
                v_lower = v.strip().lower()
                if any(tk in v_lower for tk in title_keywords):
                    return v.strip()
    # Fallback: просто берём первую длинную строку
    for r in range(1, 3):
        for c in range(1, 4):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and len(v.strip()) > 15:
                return v.strip()
    return None


def _find_organization(ws, max_row: int = 3) -> Optional[str]:
    """
    Ищем название организации (ООО/АО/ПАО + наименование) во всех ячейках первых строк.

    Некоторые компании (biznes_otel) размещают ООО в merged-ячейках далеко справа (col 45+),
    поэтому сканируем всю ширину строки.
    """
    import re
    org_pattern = re.compile(r'(ООО|АО|ПАО|ОАО|ЗАО)\s*[«"\'"].+?[»"\'\"]', re.IGNORECASE)
    max_col = ws.max_column or 50
    for r in range(1, max_row + 1):
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str):
                m = org_pattern.search(v)
                if m:
                    return m.group(0)
    return None


def extract_fields(ws) -> Dict[str, Any]:
    """
    Динамическое извлечение полей заголовка КПСЦ.

    Вместо захардкоженных координат ищем лейблы по ключевым словам
    и берём значения из соседних ячеек.
    """
    return {
        "title": _find_title(ws),
        "organization": _find_organization(ws),
        "flow_name": _find_label_value(ws, ["поток:", "наименование потока"]),
        "responsible": _find_label_value(ws, ["ответственн"]),
        "date_developed": _find_label_value(ws, ["дата разработ"]),
        "date_implementation": _find_label_value(ws, ["дата реализ", "дата достиж"]),
        "compiled_by": _find_label_value(ws, ["составил", "разработал"]),
        "takt_time": _find_label_value(ws, ["такт", "время такта", "takt"]),
    }


def build_payload(xlsx: Path, sheet_name: Optional[str] = None):
    """Строит payload из данных header-блока КПСЦ."""
    wb = load_workbook(xlsx, data_only=True)

    # Поиск листа
    if sheet_name:
        ws = wb[sheet_name]
    else:
        ws = find_sheet(
            wb,
            keywords=["кпсц"],
            exclude_keywords=["спагетти", "укрупн", "оцифровк"],
            prefer_keywords=["тс", "текущ"],
        )
        if ws is None:
            raise ValueError(f"Лист КПСЦ не найден среди {wb.sheetnames}")

    actual_sheet = ws.title
    max_col = ws.max_column or 50

    # Собираем сырые данные
    cells, _comments = collect_cells(ws, max_col=max_col)

    # Извлекаем поля через динамический поиск
    fields = extract_fields(ws)

    return {
        "meta": {
            "workbook": str(xlsx),
            "sheet": actual_sheet,
            "region": f"A1:{get_column_letter(max_col)}{HEADER_SCAN_MAX_ROW}",
        },
        "fields": fields,
    }


def parse(xlsx_path: Path, output_dir: Path) -> dict:
    """Парсит верхний блок КПСЦ и сохраняет результат в output_dir/kpsc_header_v2.json."""
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = build_payload(xlsx_path)
    output_path = output_dir / "kpsc_header_v2.json"
    output_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
    return payload


def main():
    ap = argparse.ArgumentParser(description="Парсер верхнего блока КПСЦ")
    ap.add_argument("-i", "--input", required=True)
    ap.add_argument("-s", "--sheet", default=None)
    ap.add_argument("-o", "--output")
    args = ap.parse_args()

    payload = build_payload(Path(args.input), args.sheet)
    data = json.dumps(payload, ensure_ascii=False, indent=2, default=str)
    if args.output:
        Path(args.output).write_text(data, encoding="utf-8")
        print(f"Wrote {args.output}")
    else:
        print(data)


if __name__ == "__main__":
    main()
