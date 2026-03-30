#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Парсер XLSX для 2.6 План-график мероприятий.

Извлекает данные из фиксированных ячеек и строк таблицы.
"""

from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter, column_index_from_string


def parse_plan_grafik(file_path: str) -> Dict[str, Any]:
    """
    Парсит XLSX план-графика.

    Args:
        file_path: путь к XLSX

    Returns:
        Dict со всеми извлечёнными данными
    """
    wb = load_workbook(file_path, data_only=False)
    sheets = wb.sheetnames

    result = {
        "filename": Path(file_path).name,
        "sheets": sheets,
        "approval": {},
        "dates": {},
        "title": "",
        "headers_row16": {},
        "headers_row17": {},
        "data_rows": [],
        "status_columns": {},
        "signature": "",
        "formulas": {},
    }

    if "План мероприятий" not in sheets:
        result["error"] = "Лист 'План мероприятий' не найден"
        return result

    ws = wb["План мероприятий"]

    # --- Заголовок документа ---
    result["title"] = _cell_str(ws, "B", 8)

    # --- Блок УТВЕРЖДАЮ ---
    dm_col = column_index_from_string("DM")
    result["approval"] = {
        "marker": _cell_str_by_idx(ws, dm_col, 8),      # УТВЕРЖДАЮ
        "position": _cell_str_by_idx(ws, dm_col, 9),     # Генеральный директор
        "company": _cell_str_by_idx(ws, dm_col, 10),     # ООО "___"
        "fio": _cell_str_by_idx(ws, dm_col, 11),         # ФИО
        "date_line": _cell_str_by_idx(ws, dm_col, 12),   # Дата подписания
    }

    # --- Даты мероприятий ---
    i12 = ws.cell(row=12, column=column_index_from_string("I")).value
    i13 = ws.cell(row=13, column=column_index_from_string("I")).value
    result["dates"] = {
        "start": str(i12) if i12 else "",
        "end": str(i13) if i13 else "",
        "start_raw": i12,
        "end_raw": i13,
    }

    # --- Заголовки таблицы (строка 16) ---
    for col in range(1, 160):
        val = ws.cell(row=16, column=col).value
        if val and str(val).strip():
            letter = get_column_letter(col)
            result["headers_row16"][letter] = str(val).strip()

    # --- Подзаголовки (строка 17) ---
    for col in range(1, 160):
        val = ws.cell(row=17, column=col).value
        if val and str(val).strip():
            letter = get_column_letter(col)
            result["headers_row17"][letter] = str(val).strip()

    # --- Данные (строки 18+) ---
    max_row = ws.max_row
    for row in range(18, min(max_row + 1, 500)):
        a_val = ws.cell(row=row, column=1).value  # A = № проблемы
        j_val = ws.cell(row=row, column=10).value  # J = План/Факт
        i_val = ws.cell(row=row, column=9).value   # I = Ответственный
        k_val = ws.cell(row=row, column=11).value  # K = Начало
        l_val = ws.cell(row=row, column=12).value  # L = Окончание
        c_val = ws.cell(row=row, column=3).value   # C = Проблема
        d_val = ws.cell(row=row, column=4).value   # D = Мероприятие

        # Пропускаем полностью пустые строки
        if all(v is None for v in [a_val, j_val, i_val, k_val, l_val, c_val, d_val]):
            continue

        row_data = {
            "row": row,
            "problem_num": str(a_val).strip() if a_val else "",
            "problem": str(c_val).strip() if c_val else "",
            "measure": str(d_val).strip() if d_val else "",
            "responsible": str(i_val).strip() if i_val else "",
            "plan_fact": str(j_val).strip() if j_val else "",
            "start_date": k_val,
            "end_date": l_val,
        }
        result["data_rows"].append(row_data)

    # --- Столбцы статуса/отклонений ---
    dj_col = column_index_from_string("DJ")
    dk_col = column_index_from_string("DK")
    dl_col = column_index_from_string("DL")
    dm_col_idx = column_index_from_string("DM")

    result["status_columns"] = {
        "status_header": _cell_str_by_idx(ws, dj_col, 16),
        "comments_header": _cell_str_by_idx(ws, dm_col_idx, 16),
        "status_present": bool(ws.cell(row=16, column=dj_col).value),
        "comments_present": bool(ws.cell(row=16, column=dm_col_idx).value),
    }

    # --- Формулы (проверка целостности) ---
    formulas = {}

    # DK17 = длительность
    dk17 = ws.cell(row=17, column=dk_col)
    formulas["DK17"] = {
        "value": str(dk17.value) if dk17.value else "",
        "is_formula": str(dk17.value).startswith("=") if dk17.value else False,
    }

    # DK18 = отклонение начала (первая строка данных)
    dk18 = ws.cell(row=18, column=dk_col)
    formulas["DK18"] = {
        "value": str(dk18.value) if dk18.value else "",
        "is_formula": str(dk18.value).startswith("=") if dk18.value else False,
    }

    # DL18 = отклонение окончания
    dl18 = ws.cell(row=18, column=dl_col)
    formulas["DL18"] = {
        "value": str(dl18.value) if dl18.value else "",
        "is_formula": str(dl18.value).startswith("=") if dl18.value else False,
    }

    # DK12 — проверяем на #REF
    dk12 = ws.cell(row=12, column=dk_col)
    formulas["DK12"] = {
        "value": str(dk12.value) if dk12.value else "",
        "is_formula": str(dk12.value).startswith("=") if dk12.value else False,
        "has_error": "#REF" in str(dk12.value) if dk12.value else False,
    }

    # Проверяем несколько ячеек Ганта на наличие формул
    n12 = ws.cell(row=12, column=column_index_from_string("N"))
    formulas["N12"] = {
        "value": str(n12.value) if n12.value else "",
        "is_formula": str(n12.value).startswith("=") if n12.value else False,
    }

    result["formulas"] = formulas

    # --- Подпись (последние строки) ---
    for row in range(max(1, max_row - 5), max_row + 1):
        for col in range(1, 160):
            val = ws.cell(row=row, column=col).value
            if val and str(val).strip() and ("подпись" in str(val).lower() or "___" in str(val) or "202" in str(val)):
                result["signature"] += f"{get_column_letter(col)}{row}: {str(val).strip()}\n"

    return result


def _cell_str(ws, col_letter: str, row: int) -> str:
    """Извлекает строковое значение ячейки по букве столбца и номеру строки."""
    col_idx = column_index_from_string(col_letter)
    val = ws.cell(row=row, column=col_idx).value
    return str(val).strip() if val else ""


def _cell_str_by_idx(ws, col_idx: int, row: int) -> str:
    """Извлекает строковое значение ячейки по индексу столбца и номеру строки."""
    val = ws.cell(row=row, column=col_idx).value
    return str(val).strip() if val else ""
